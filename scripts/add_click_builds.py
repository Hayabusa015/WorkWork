#!/usr/bin/env python3
"""Add native click-by-click orbital-diagram builds to a SHULL slide deck (.pptx).

pptxgenjs (templates/slide/build_deck.js) cannot write animations, so this is a POST-PROCESSOR.
It never touches the masters or the deck builder. It reads a companion spec, draws the orbital
diagram on the named slides as native shapes, and writes a p:timing tree so that each click
does exactly one thing:

    click 1      boxes + sublevel labels + title fly in (from bottom)       [start: "empty"]
    click 2..n   one electron arrow per click, in Hund / Pauli order        (fade or appear)
    ion steps    one arrow per click leaves (exit) or arrives (entrance)

    python3 scripts/add_click_builds.py deck.pptx deck.builds.json [out.pptx]
    python3 scripts/add_click_builds.py deck.pptx deck.builds.json out.pptx --preview K [--slide N]
        --preview K  writes a STATIC copy showing the state after K clicks (no animation). For
                     looking at intermediate states in a raster; never a deliverable.
    python3 scripts/add_click_builds.py deck.pptx deck.builds.json out.pptx --static
        --static     writes the PRINT / HANDOUT variant: no p:timing, and each build slide shows only its
                     coherent FINAL state (see "Static final-state mode" below). Opt-in; without the flag the
                     output is unchanged. Use it for any export that is not slideshow mode (PDF, print,
                     handout, Normal view, Google Slides import), where every shape shows at once.

Spec (JSON). One object per build slide; see templates/slide/README.md, "Click builds".
    {"course": "chemistry",
     "builds": [{
        "slide": 8,                        # 1-based
        "layout": "08_DIAGRAM_ANNOTATION", # asserted, so a renumbered deck cannot be built wrongly
        "region": "diagram",               # a slot on that layout: coordinates come from build.js SLOTS
        "title": "OXYGEN  Z = 8",          # ^{..} superscript and _{..} subscript allowed
        "z": 8,                            # optional; asserted against the sublevel totals
        "core": {"label": "[Ar]", "electrons": 18},   # optional noble-gas core, counted in Z
        "sublevels": [{"label": "1s", "boxes": 1, "electrons": 2}, ...],  # ground state, Aufbau order
        "start": "empty" | "filled",       # empty: fill by clicks. filled: atom flies in whole
        "effect": "fade" | "appear",       # electron arrows (default fade)
        "captions": {"hund": "...", "pauli": "..."},   # fill builds; shown once, at the step the rule starts
        "steps": [{"highlight": "valence", "caption": "..."},      # ion builds, "start": "filled"
                  {"remove": 2, "caption": "..."}, {"add": 1}],
        "start_text": "Fe: [Ar] 4s^{2} 3d^{6}",   # optional state line, arrives with click 1
        "sum": "auto" | "text" | false,    # final line; "auto" = 2 + 2 + 4 = 8
        "result": "Fe^{2+}",               # optional ion result, appears with the last click
        "name": "Na",                      # optional; prefixes this build's lines in the notes
        "box_max": 0.5                     # optional cap on the box size, inches (matching two atoms)
     }]}

Ion-build extras (added for S2.4):
  {"highlight": "valence", "valence": 2}   a step with no electron change: the arrows of the outermost
        level (highest n among the drawn sublevels) are recoloured at that click. Drawn as a second set of
        arrows laid over the first (course Primary fill, Asphalt outline), entering by fade, and leaving
        together with their arrow when it is removed. "valence" is optional; when given it is asserted
        against the count found. One accent only, from the course tokens.
  {"remove": 1, "text": "Fe^{2+}: charge 2+, 24 electrons"}   a step may carry "text": a state line
        that fades in at that click. With "start_text" the lines form a record under the diagram, one
        row each, so the final state shows the whole record and no two texts ever share a box. Do not
        combine with "sum" or "result".
  Captions may wrap to two lines (asserted <= 2); each caption keeps its own row so the final state
        never overprints.
  "band": [1, 2]   draw this build in the first of two equal horizontal bands of the region, so two
        atoms (Na, then Cl) share one slot without any shape or text overlapping the other's. Compact
        metrics apply in a band: no more than one state row, and no captions, will fit (asserted by the
        minimum box size). Several objects may name the same slide; their clicks run in spec order as
        one sequence and the notes line lists them all.

Added for S2.3 (every key below is opt-in; a spec that does not use them builds byte-identically):
  top level   "alt_text": true       give each drawn diagram a computed description (descr) from its data;
                                     each build then needs "element". "picture_alt": [{"slide": N, "descr": "..."}]
                                     sets the description of the static pictures on a slide.
  {"remove": 1, "from": "4s"}  {"add": 1, "to": "3d"}   a step names its sublevel instead of the default
        (valence-first removal, first-open-sublevel addition). Two such steps are a "move".
  "predicted": true   sublevels are the Aufbau-predicted state of Z (asserted), which differs from the measured
        ground state; after the steps the final state is asserted to EQUAL the measured ground state
        (the S2.3 checker's exception table), and the electron total is asserted unchanged.
  "story": true, "start_caption": "..."   caption rows and state rows are laid out in the order they
        arrive (caption, state, caption, state), each its own row, so the final state reads as a record.
  "late_counts": true   fill builds only: the sublevel label shows its name from click 1 and the raised
        electron count fades in WITH the last arrow of that sublevel (a second text box, adjacent).
  "type": "text_line"   a build of native text boxes, one per configuration segment (see process_text_line):
        {"longhand": "1s2 2s2 ...", "core": "Ar", "z": 32, "title": ..., "captions": {"highlight", "replace"},
         "result": "auto"}: click 1 the longhand and its sum appear; click 2 the segments the noble gas
        covers are highlighted; click 3 they leave and [Ar] appears in their place; click 4 the finished
        shorthand line appears. Everything is computed from the checker's tables and asserted.

Static final-state mode (--static; every key is opt-in, the animated output is unaffected):
  Same builds, same asserted chemistry, no animation. Per build the static slide shows
    arrows        only the arrows of the final state (removed arrows absent, added arrows present)
    highlights    none (the valence overlay is a teaching step, never a final state)
    labels        every sublevel label carries its FINAL count, e.g. 4s^{0} 3d^{5} for Fe3+ (ion builds)
    state lines   the "start_text" (before) and the LAST step "text" (after); intermediate species lines
                  (e.g. Fe2+ between Fe and Fe3+) are animated-only
    captions      kept, each on its own row (fill builds: Hund and Pauli; ion builds: the step captions)
    text_line     (Ge) the longhand line, then the shorthand line ([Ar] chip + the remaining segments),
                  one caption saying what the noble-gas symbol stands for, and the check sum. The click-2
                  caption ("these 18 electrons...") describes a highlight that is not shown, so it is not used.
  {"static_title": "Fe^{3+}  Z = 26  [Ar] CORE"}   optional: the diagram title in the static output, when the
        final state is a different species from the title that opens the animation.
  {"static_caption": "..."}   text_line only, optional: replaces the computed noble-gas caption.
  Notes: the static output carries "STATIC FINAL STATE" instead of the "CLICK BUILD" line.

Everything chemical is computed and asserted here, not typed: arrows are generated in fill order
(within a sublevel one up arrow into each box left to right, THEN the down arrows left to right),
sublevels must be in Aufbau order, no box may hold more than 2 arrows and two must be opposite,
and the electron totals must match. Ions remove valence-first (highest n, then highest l, so
4s before 3d), reverse of fill order within a sublevel, and add to the first open sublevel.

The animation XML is written from the ECMA-376 / PowerPoint conventions and is validated by this
script (ids, targets, click count, final state). It has been checked by LibreOffice import. It has
NOT been played in PowerPoint, Keynote or Google Slides from this environment.
"""
import copy
import importlib.util
import json
import os
import re
import subprocess
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SLIDE_DIR = os.path.join(REPO, "templates", "slide")
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
IN = 914400
FLOOR_PT = 16
ORDER_L = {"s": 0, "p": 1, "d": 2, "f": 3}


# --------------------------------------------------------------------------- tokens / geometry
def load_tokens(course):
    """Slots, fonts and colours straight from the template (build.js SLOTS, tokens.generated.js).
    Nothing is retyped here."""
    js = ("const b=require('./build');const K=require('./tokens.generated');"
          "console.log(JSON.stringify({slots:b.SLOTS,fonts:K.fonts,ground:K.ground,"
          "courses:K.courses,floor:K.floor}))")
    r = subprocess.run(["node", "-e", js], cwd=SLIDE_DIR, capture_output=True, text=True,
                       env=dict(os.environ, SHULL_COURSE=course))
    if r.returncode:
        sys.exit("add_click_builds: could not read tokens via node: " + r.stderr.strip())
    return json.loads(r.stdout.strip().splitlines()[-1])


def load_checker():
    """Aufbau order, caps and exceptions live in ONE place: the S2.3 figure checker."""
    p = os.path.join(REPO, "scripts", "build_diagram_u02_s02.3_configurations.py")
    spec = importlib.util.spec_from_file_location("s023_checker", p)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


# --------------------------------------------------------------------------- chemistry
class Arrow:
    def __init__(self, sub, box, spin, order):
        self.sub, self.box, self.spin, self.order = sub, box, spin, order
        self.spid = None


def fill_arrows(sub, n):
    """Arrows of one sublevel in fill order: ups left to right, then downs left to right."""
    b = sub["boxes"]
    seq = [(i, "u") for i in range(min(n, b))] + [(i, "d") for i in range(max(0, n - b))]
    return [Arrow(sub["label"], box, spin, k) for k, (box, spin) in enumerate(seq)]


def check_boxes(arrows, sublevels):
    """Pauli and Hund, asserted on a set of arrows."""
    cells = {}
    for a in arrows:
        cells.setdefault((a.sub, a.box), []).append(a.spin)
    for (sub, box), spins in cells.items():
        assert len(spins) <= 2, f"{sub} box {box + 1}: more than 2 arrows"
        assert len(set(spins)) == len(spins), f"{sub} box {box + 1}: two arrows with the same spin (Pauli)"
    for s in sublevels:
        nb = s["boxes"]
        for box in range(nb):
            if "d" in cells.get((s["label"], box), []):
                assert "u" in cells.get((s["label"], box), []), f"{s['label']}: down arrow without its up arrow"
        if any("d" in cells.get((s["label"], b), []) for b in range(nb)):
            assert all("u" in cells.get((s["label"], b), []) for b in range(nb)), \
                f"{s['label']}: a pair before every box has a single (Hund)"


def validate_atom(b, chk):
    subs = b["sublevels"]
    idx = []
    for s in subs:
        lab = s["label"]
        assert re.fullmatch(r"\d[spdf]", lab), f"bad sublevel label {lab}"
        assert lab in chk.FILL_ORDER, f"{lab} is not an Aufbau sublevel"
        assert s["boxes"] == 2 * ORDER_L[lab[1]] + 1, f"{lab}: {s['boxes']} boxes, expected {2 * ORDER_L[lab[1]] + 1}"
        assert 0 <= s["electrons"] <= chk.CAP[lab[1]], f"{lab}: {s['electrons']} electrons exceeds the cap"
        idx.append(chk.FILL_ORDER.index(lab))
    assert idx == sorted(idx) and len(set(idx)) == len(idx), "sublevels are not in Aufbau order"
    core = b.get("core") or {"label": "", "electrons": 0}
    total = core["electrons"] + sum(s["electrons"] for s in subs)
    if b.get("z") is not None:
        assert total == b["z"], f"slide {b['slide']}: electrons total {total}, Z = {b['z']}"
        if b.get("core"):
            sym = core["label"].strip("[]")
            assert chk.NOBLE_Z[sym] == core["electrons"], "core electron count does not match its noble gas"
        if b.get("predicted"):
            assert b["z"] in chk.EXCEPTION, f"slide {b['slide']}: \"predicted\" is for a taught exception, Z={b['z']} is not one"
            gs = aufbau_state(b["z"], chk)
            assert gs != chk.ground_state(b["z"]), "the predicted state equals the measured one: nothing to show"
        else:
            gs = chk.ground_state(b["z"])
        for s in subs:
            assert gs.get(s["label"], 0) == s["electrons"], \
                f"slide {b['slide']}: {s['label']} has {s['electrons']}, {'Aufbau' if b.get('predicted') else 'ground'} state of Z={b['z']} has {gs.get(s['label'], 0)}"
    return total


def aufbau_state(Z, chk):
    """What the filling rules alone give for Z: the checker's fill with the exception table ignored."""
    left, out = Z, {}
    for sub in chk.FILL_ORDER:
        if left <= 0:
            break
        take = min(chk.CAP[sub[1]], left)
        out[sub] = take
        left -= take
    return out


def plan_steps(b):
    """Return (initial_arrows, steps, all_arrows). A step is
    {"kind": "add"|"remove", "arrow": Arrow, "caption": str|None, "label": str}."""
    subs = b["sublevels"]
    by_label = {s["label"]: s for s in subs}
    caps = b.get("captions", {})
    steps_spec = b.get("steps")
    if not steps_spec:
        steps, seen_hund, seen_pauli, initial = [], False, False, []
        for s in subs:
            for a in fill_arrows(s, s["electrons"]):
                cap = None
                if a.spin == "u" and s["boxes"] > 1 and not seen_hund and caps.get("hund"):
                    cap, seen_hund = caps["hund"], True
                if a.spin == "d" and not seen_pauli and caps.get("pauli"):
                    cap, seen_pauli = caps["pauli"], True
                steps.append({"kind": "add", "arrow": a, "caption": cap,
                              "label": f"{s['label']} {'up' if a.spin == 'u' else 'down'} arrow"})
        arrows = [st["arrow"] for st in steps]
        if b.get("start", "empty") == "filled":
            return arrows, [], arrows
        return [], steps, arrows
    assert b.get("start") == "filled", "ion steps need \"start\": \"filled\""
    count = {s["label"]: s["electrons"] for s in subs}
    initial = [a for s in subs for a in fill_arrows(s, s["electrons"])]
    pool = {s["label"]: fill_arrows(s, s["boxes"] * 2) for s in subs}   # every possible arrow, in fill order
    present = {a.sub: [pool[a.sub][i] for i in range(count[a.sub])] for a in initial}
    for s in subs:
        present.setdefault(s["label"], [])
    steps = []
    for st in steps_spec:
        if "highlight" in st:
            assert st["highlight"] == "valence", f"unknown highlight {st['highlight']!r}"
            live = [s for s in subs if present[s["label"]]]
            assert live, "highlight on an empty atom"
            top = max(int(s["label"][0]) for s in live)
            arrs = [a for s in live if int(s["label"][0]) == top for a in present[s["label"]]]
            if st.get("valence") is not None:
                assert len(arrs) == st["valence"], \
                    f"slide {b['slide']}: {len(arrs)} arrows in n = {top}, spec says {st['valence']} valence"
            steps.append({"kind": "highlight", "arrow": None, "arrows": arrs, "caption": st.get("caption"),
                          "text": st.get("text"),
                          "label": "highlight the valence arrows (" +
                                   ", ".join(s["label"] for s in live if int(s["label"][0]) == top) + ")"})
            continue
        kind = "remove" if "remove" in st else "add"
        k = st[kind]
        for j in range(k):
            if kind == "remove":
                if st.get("from"):                               # a named sublevel (a "move")
                    assert st["from"] in by_label, f"remove from {st['from']!r}: not a drawn sublevel"
                    tgt = by_label[st["from"]]
                    assert present[tgt["label"]], f"remove from {tgt['label']}: it holds no electrons"
                else:
                    cands = [s for s in subs if present[s["label"]]]
                    assert cands, "ion step removes from an empty atom"
                    tgt = max(cands, key=lambda s: (int(s["label"][0]), ORDER_L[s["label"][1]]))
                a = present[tgt["label"]].pop()
            else:
                if st.get("to"):
                    assert st["to"] in by_label, f"add to {st['to']!r}: not a drawn sublevel"
                    tgt = by_label[st["to"]]
                    assert len(present[tgt["label"]]) < tgt["boxes"] * 2, f"add to {tgt['label']}: it is full"
                else:
                    cands = [s for s in subs if len(present[s["label"]]) < s["boxes"] * 2]
                    assert cands, "no open sublevel to add to"
                    tgt = cands[0]
                a = pool[tgt["label"]][len(present[tgt["label"]])]
                present[tgt["label"]].append(a)
            steps.append({"kind": kind, "arrow": a, "caption": st.get("caption") if j == 0 else None,
                          "text": st.get("text") if j == k - 1 else None,
                          "label": f"{'remove' if kind == 'remove' else 'add'} {tgt['label']} "
                                   f"{'up' if a.spin == 'u' else 'down'} arrow"})
    allarrows = {}
    for s in subs:
        for a in pool[s["label"]]:
            allarrows[(a.sub, a.box, a.spin)] = a
    initial_objs = [pool[a.sub][a.order] for a in initial]
    used = initial_objs + [st["arrow"] for st in steps if st["kind"] == "add"]
    hl = [a for st in steps if st["kind"] == "highlight" for a in st["arrows"]]
    uniq = []
    for a in used:
        if a not in uniq:
            uniq.append(a)
    # final-state check against the electron arithmetic
    final = {s["label"]: len(present[s["label"]]) for s in subs}
    net = sum(1 if st["kind"] == "add" else -1 for st in steps if st["kind"] != "highlight")
    assert sum(final.values()) == sum(count.values()) + net
    check_boxes([a for s in subs for a in present[s["label"]]], subs)
    return initial_objs, steps, uniq


# --------------------------------------------------------------------------- drawing
def rgb(hex6):
    return RGBColor.from_string(hex6.upper())


def add_runs(par, text, font, size, color, bold):
    """^{..} superscript, _{..} subscript runs; the run keeps font, size, colour."""
    pos = 0
    for m in re.finditer(r"([\^_])\{([^{}]*)\}", text):
        pieces = [(text[pos:m.start()], None), (m.group(2), m.group(1))]
        pos = m.end()
        for t, kind in pieces:
            _run(par, t, font, size, color, bold, kind)
    _run(par, text[pos:], font, size, color, bold, None)


def _run(par, t, font, size, color, bold, kind):
    if not t:
        return
    r = par.add_run()
    r.text = t
    f = r.font
    f.name, f.size, f.bold = font, Pt(size), bold
    f.color.rgb = rgb(color)
    if kind:
        r._r.get_or_add_rPr().set("baseline", "30000" if kind == "^" else "-25000")


def textbox(shapes, x, y, w, h, text, font, size, color, bold=False, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, name="BUILD text", wrap=True, fill=None, outline=None):
    assert size >= FLOOR_PT, f"{name}: {size} pt is below the {FLOOR_PT} pt floor"
    tb = shapes.add_textbox(Emu(int(x * IN)), Emu(int(y * IN)), Emu(int(w * IN)), Emu(int(h * IN)))
    tb.name = name
    if fill:
        tb.fill.solid(); tb.fill.fore_color.rgb = rgb(fill)
    if outline:
        tb.line.color.rgb = rgb(outline); tb.line.width = Pt(1.5)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line_ in enumerate(text.split("\n")):          # "\n" starts a new paragraph
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        add_runs(p, line_, font, size, color, bold)
    # no autofit: text is never shrunk to fit
    return tb


def plain(t):
    return re.sub(r"([\^_])\{([^{}]*)\}", r"\2", t)


def est_lines(text, width_in, size_pt, bold=False):
    cw = size_pt * (0.58 if bold else 0.54) / 72.0          # Archivo is wide; conservative
    per = max(1, int(width_in / cw))
    words, lines, cur = plain(text).split(), 1, 0
    for w in words:
        if cur and cur + 1 + len(w) > per:
            lines, cur = lines + 1, len(w)
        else:
            cur = cur + (1 if cur else 0) + len(w)
    return lines


_PIL = {}


def text_w(text, size, bold=True, sup=0.62):
    """Width in inches of text with ^{..}/_{..} runs, from the real Archivo metrics (brand/fonts).
    A raised or lowered run is measured at `sup` of the size (PowerPoint and LibreOffice differ a little)."""
    from PIL import ImageFont
    key = bool(bold)
    if key not in _PIL:
        _PIL[key] = ImageFont.truetype(os.path.join(REPO, "brand", "fonts",
                                                    "Archivo-Bold.ttf" if bold else "Archivo-Regular.ttf"), 200)
    f, w, pos = _PIL[key], 0.0, 0
    for m in re.finditer(r"([\^_])\{([^{}]*)\}", text):
        w += f.getlength(text[pos:m.start()]) * size / 200 / 72
        w += f.getlength(m.group(2)) * size * sup / 200 / 72
        pos = m.end()
    w += f.getlength(text[pos:]) * size / 200 / 72
    return w


def wrap_count(text, width_in, size, bold=True, safety=1.06):
    """Lines a word-wrapped text needs in width_in, measured (not guessed); 6 % safety on the width."""
    lines, cur = 1, ""
    for w in plain(text).split():
        trial = (cur + " " + w).strip()
        if cur and text_w(trial, size, bold) * safety > width_in:
            lines, cur = lines + 1, w
        else:
            cur = trial
    return lines


NUMW = {0: "no", 1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six", 7: "seven"}


def state_words(subs, arrows):
    """'1s: 1 box, one pair. 2p: 3 boxes, one pair and two single up arrows.' computed from the arrows."""
    out = []
    for s in subs:
        cells = {}
        for a in arrows:
            if a.sub == s["label"]:
                cells.setdefault(a.box, []).append(a.spin)
        pairs = sum(1 for v in cells.values() if len(v) == 2)
        ups = sum(1 for v in cells.values() if v == ["u"])
        downs = sum(1 for v in cells.values() if v == ["d"])
        empty = s["boxes"] - len(cells)
        bits = []
        if pairs:
            bits.append(f"{NUMW[pairs]} pair" + ("s" if pairs > 1 else ""))
        if ups:
            bits.append(f"{NUMW[ups]} single up arrow" + ("s" if ups > 1 else ""))
        if downs:
            bits.append(f"{NUMW[downs]} single down arrow" + ("s" if downs > 1 else ""))
        if empty:
            bits.append(f"{NUMW[empty]} empty box" + ("es" if empty > 1 else ""))
        out.append(f"{s['label']}: {NUMW[s['boxes']]} box{'es' if s['boxes'] > 1 else ''}, "
                   + (" and ".join(bits) if len(bits) < 3 else ", ".join(bits[:-1]) + " and " + bits[-1]) + ".")
    return " ".join(out)


def describe_diagram(b, initial, final, arrows):
    """Alt text for a drawn orbital diagram, computed from the data that drew it."""
    el = b.get("element")
    assert el, f"slide {b['slide']}: alt_text needs an \"element\" on every build"
    subs = b["sublevels"]
    core = b.get("core")
    head = f"Orbital diagram of {el}" + (f", {core['label']} core not drawn" if core else "") + ". "
    if b.get("steps") and b.get("start") == "filled":
        start = state_words(subs, initial)
        end = state_words(subs, final)
        if start == end:
            return head + end
        return head + "As first drawn: " + start + " After the steps: " + end
    return head + state_words(subs, final)


def draw(slide, b, region, T, initial, steps, arrows):
    """Draw every shape. Returns dict of shape handles (python-pptx shapes)."""
    ink, white = T["ground"]["asphalt"], T["ground"]["white"]
    font = T["fonts"]["body"]
    deep = T["courses"][b["_course"]]["primaryDeep"]
    x0, y0, W, H = region["x"], region["y"], region["w"], region["h"]
    compact = b.get("band") is not None
    pad = 0.02 if compact else 0.12
    subs = b["sublevels"]
    nb = sum(s["boxes"] for s in subs)
    inner, outer = 0.05, 0.28
    title_h, label_h, line = 0.30, 0.30, 0.28
    srow = 0.30 if compact else 0.32
    gap = 0.02 if compact else 0.04
    caps = [st["caption"] for st in steps if st["caption"]]
    story = bool(b.get("story"))
    rows = []
    if story:
        # caption / state rows in arrival order, each its own row (see the docstring)
        assert b.get("steps") and not b.get("result") and b.get("sum", "auto") in ("auto", False, None), \
            "story rows are for step builds without a result or custom sum"
        assert not compact, "story rows do not fit a band"
        if b.get("start_caption"):
            rows.append((None, "cap", b["start_caption"]))
        if b.get("start_text"):
            rows.append((None, "state", b["start_text"]))
        for st in steps:
            if st["caption"]:
                rows.append((st, "cap", st["caption"]))
            if st.get("text"):
                rows.append((st, "state", st["text"]))
        story_rows = []
        for st_, kind, t in rows:
            if kind == "cap":
                nl = est_lines(t, W - 2 * pad, FLOOR_PT, True)
                assert nl <= 2, f"caption does not fit two lines here: {t!r}"
                story_rows.append((st_, kind, t, nl * line))
            else:
                assert est_lines(t, W - 2 * pad, 18, True) == 1, f"state text does not fit one line: {t!r}"
                story_rows.append((st_, kind, t, srow))
        caps = []
    else:
        assert not b.get("start_caption"), "start_caption needs \"story\": true"
    # Captions accumulate, each in its own row (one or two lines), so a static export (PDF, handout) of
    # the final state never has two captions drawn on top of each other.
    cap_lines = [est_lines(c, W - 2 * pad, FLOOR_PT, True) for c in caps]
    for c, n in zip(caps, cap_lines):
        assert n <= 2, f"caption does not fit two lines here: {c!r}"
    cap_h = sum(n * line for n in cap_lines)
    states = bool(b.get("start_text")) or any(st.get("text") for st in steps)
    if story:
        sum_h = sum(r[3] for r in story_rows)
    elif states:
        assert not b.get("result") and b.get("sum", "auto") in ("auto", False, None), \
            "start_text / step text cannot be combined with result or a custom sum"
        texts = [t for t in [b.get("start_text")] + [st.get("text") for st in steps] if t]
        for t in texts:
            assert est_lines(t, W - 2 * pad, 18, True) == 1, f"state text does not fit one line: {t!r}"
        sum_h = srow * len(texts)
    else:
        sum_h = 0.32 if (b.get("sum", "auto") or b.get("result")) else 0
    gaps = gap * 4
    # compact (band) builds put the title in a column to the left of the boxes, not above them
    tw = 0.95 if compact else 0.0
    if compact:
        title_h = 0.0
    avail_h = H - 2 * pad - title_h - label_h - cap_h - sum_h - gaps
    avail_w = W - 2 * pad - tw - (nb - len(subs)) * inner - (len(subs) - 1) * outer
    size = min(b.get("box_max", 1.0), avail_w / nb, avail_h)
    assert size >= 0.4, f"slide {b['slide']}: boxes would be {size:.2f} in; the region is too small"
    row_w = nb * size + (nb - len(subs)) * inner + (len(subs) - 1) * outer
    # vertical stack, centred in the region
    stack = title_h + size + label_h + cap_h + sum_h + gaps
    y = y0 + (H - stack) / 2
    xs = x0 + pad + tw + (W - 2 * pad - tw - row_w) / 2

    out = {"backdrop": None, "group": None, "arrows": {}, "hl": {}, "captions": [], "sum": None,
           "states": [], "extra": [], "counts": {}}
    bd = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(x0 * IN)), Emu(int(y0 * IN)),
                                Emu(int(W * IN)), Emu(int(H * IN)))
    bd.name = "BUILD backdrop"
    bd.fill.solid(); bd.fill.fore_color.rgb = rgb(white); bd.line.fill.background()
    bd.shadow.inherit = False
    out["backdrop"] = bd

    grp = slide.shapes.add_group_shape()
    grp.name = "BUILD boxes and labels"
    out["group"] = grp
    gs = grp.shapes
    if compact:
        th = max(size, 0.31 * (b["title"].count("\n") + 1))        # 18 pt lines, never clipped
        textbox(gs, x0 + pad, y + size / 2 - th / 2, tw, th, b["title"], font, 18, ink, True, name="BUILD title")
    else:
        textbox(gs, x0 + pad, y, W - 2 * pad, title_h, b["title"], font, 18, ink, True, name="BUILD title")
    by = y + title_h + (0 if compact else gap)
    box_xy = {}
    late = bool(b.get("late_counts"))
    late_boxes = []
    if late:
        assert not b.get("steps") and b.get("start", "empty") == "empty", \
            "late_counts is for fill builds that start empty"
    x = xs
    for s in subs:
        gw = s["boxes"] * size + (s["boxes"] - 1) * inner
        for i in range(s["boxes"]):
            bx = x + i * (size + inner)
            r = gs.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(bx * IN)), Emu(int(by * IN)),
                             Emu(int(size * IN)), Emu(int(size * IN)))
            r.name = f"BUILD box {s['label']} {i + 1}"
            r.fill.solid(); r.fill.fore_color.rgb = rgb(white)
            r.line.color.rgb = rgb(ink); r.line.width = Pt(2.25)
            r.shadow.inherit = False
            box_xy[(s["label"], i)] = (bx, by)
        if late:
            # name in the group (flies in at click 1), raised count a separate box that fades in with the
            # last arrow of this sublevel. Together they read, and are centred, like "2p^{4}".
            nm, cnt = s["label"], str(s["electrons"])
            wn, wc = text_w(nm, 18), text_w(f"^{{{cnt}}}", 18)
            cx0 = x + gw / 2 - (wn + wc) / 2
            textbox(gs, cx0, by + size + gap, wn + 0.02, label_h, nm, font, 18, ink, True,
                    align=PP_ALIGN.LEFT, name=f"BUILD label {nm}", wrap=False)
            late_boxes.append((s["label"], cx0 + wn - 0.04, by + size + gap, wc, label_h, cnt))
        else:
            textbox(gs, x - 0.1, by + size + gap, gw + 0.2, label_h,
                    (f"{s['label']}^{{{b['_final_counts'][s['label']]}}}" if b.get("_final_counts") is not None
                     else s["label"]) if b.get("steps") else f"{s['label']}^{{{s['electrons']}}}", font, 18, ink, True,
                    name=f"BUILD label {s['label']}")
        x += gw + outer
    # arrows (separate top-level shapes: each is animated on its own)
    aw, ah = size * 0.25, size * 0.62
    for a in arrows:
        bx, byy = box_xy[(a.sub, a.box)]
        cx = bx + size * (0.30 if a.spin == "u" else 0.70)
        shp = slide.shapes.add_shape(MSO_SHAPE.UP_ARROW if a.spin == "u" else MSO_SHAPE.DOWN_ARROW,
                                     Emu(int((cx - aw / 2) * IN)), Emu(int((byy + (size - ah) / 2) * IN)),
                                     Emu(int(aw * IN)), Emu(int(ah * IN)))
        shp.name = f"BUILD arrow {a.sub} box{a.box + 1} {'up' if a.spin == 'u' else 'down'}"
        shp.adjustments[0] = 0.34
        shp.adjustments[1] = 0.70
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(ink); shp.line.fill.background()
        shp.shadow.inherit = False
        a.spid = shp.shape_id
        a.geom = (cx - aw / 2, byy + (size - ah) / 2, aw, ah)
        out["arrows"][id(a)] = shp
    # valence highlight: a second arrow laid over the first (Primary fill, Asphalt outline)
    accent = T["courses"][b["_course"]]["primary"]
    for st in steps:
        for a in st.get("arrows") or []:
            if id(a) in out["hl"]:
                continue
            x_, y_, w_, h_ = a.geom
            hs = slide.shapes.add_shape(MSO_SHAPE.UP_ARROW if a.spin == "u" else MSO_SHAPE.DOWN_ARROW,
                                        Emu(int(x_ * IN)), Emu(int(y_ * IN)), Emu(int(w_ * IN)), Emu(int(h_ * IN)))
            hs.name = f"BUILD valence {a.sub} box{a.box + 1} {'up' if a.spin == 'u' else 'down'}"
            hs.adjustments[0] = 0.34
            hs.adjustments[1] = 0.70
            hs.fill.solid(); hs.fill.fore_color.rgb = rgb(accent)
            hs.line.color.rgb = rgb(ink); hs.line.width = Pt(1.5)
            hs.shadow.inherit = False
            out["hl"][id(a)] = hs
    for lab, bx_, by_, bw_, bh_, cnt in late_boxes:
        tbx = textbox(slide.shapes, bx_, by_, bw_ + 0.12, bh_, f"^{{{cnt}}}", font, 18, ink, True,
                      align=PP_ALIGN.LEFT, name=f"BUILD count {lab}", wrap=False)
        out["counts"][lab] = tbx
        out["extra"].append(tbx)
    cy = by + size + gap + label_h + gap
    row = 0
    if story:
        yy = cy
        for st_, kind, t, h in story_rows:
            if kind == "cap":
                tb = textbox(slide.shapes, x0 + pad, yy, W - 2 * pad, h, t, font, FLOOR_PT, deep, True,
                             name="BUILD caption")
                out["captions"].append((st_, tb))
            else:
                tb = textbox(slide.shapes, x0 + pad, yy, W - 2 * pad, h, t, font, 18, ink, True,
                             name="BUILD state")
                out["states"].append((st_, tb))
            yy += h
        return out
    for st in steps:
        if st["caption"]:
            nl = cap_lines[len(out["captions"])]
            tb = textbox(slide.shapes, x0 + pad, cy + row * line, W - 2 * pad, nl * line, st["caption"], font,
                         FLOOR_PT, deep, True, name="BUILD caption")
            out["captions"].append((st, tb))
            row += nl
    ty = cy + cap_h + gap
    s_text = None
    core = b.get("core")
    if states:
        for st_, t in [(None, b.get("start_text"))] + [(st, st.get("text")) for st in steps]:
            if t:
                out["states"].append((st_, textbox(slide.shapes, x0 + pad, ty + srow * len(out["states"]),
                                                   W - 2 * pad, srow, t, font, 18, ink, True,
                                                   name="BUILD state")))
        s_text = None
    elif b.get("result"):
        s_text = b["result"]
    elif b.get("sum", "auto"):
        if b.get("sum", "auto") == "auto":
            parts = ([str(core["electrons"])] if core else []) + [str(s["electrons"]) for s in subs]
            tot = (core["electrons"] if core else 0) + sum(s["electrons"] for s in subs)
            s_text = (core["label"] + " " if core else "") + " + ".join(parts) + f" = {tot}"
        else:
            s_text = b["sum"]
    if s_text:
        out["sum"] = textbox(slide.shapes, x0 + pad, ty, W - 2 * pad, sum_h, s_text, font, 18, ink,
                             True, name="BUILD sum")
    return out


# --------------------------------------------------------------------------- timeline
class Timeline:
    """clicks: list of clicks; click: list of effects (spid, cls, effect, is_group, has_text)."""
    def __init__(self):
        self.clicks = []

    def click(self, *effects):
        self.clicks.append(list(effects))


def eff(spid, cls, effect, grp=False, text=False, fill=False):
    return dict(spid=spid, cls=cls, effect=effect, grp=grp, text=text, fill=fill)


def build_timeline(b, shp, initial, steps):
    tl = Timeline()
    fx = b.get("effect", "fade")
    box = eff(shp["group"].shape_id, "entr", "fly", grp=True)
    pre = []
    start_state = [eff(tb.shape_id, "entr", "fade", text=True) for st_, tb in shp["states"] if st_ is None]
    start_state += [eff(tb.shape_id, "entr", "fade", text=True) for st_, tb in shp["captions"] if st_ is None]
    # late counts: the raised count of a sublevel arrives WITH the last arrow of that sublevel
    last_of = {}
    for i_, st_ in enumerate(steps):
        if st_["arrow"] is not None:
            last_of[st_["arrow"].sub] = i_
    desc = []
    if initial:                                    # filled start: the atom arrives whole
        tl.click(*pre, box, *[eff(shp["arrows"][id(a)].shape_id, "entr", "fly", fill=True) for a in initial],
                 *start_state)
        desc.append("boxes, labels and every electron arrow fly in")
    else:
        tl.click(*pre, box, *start_state)
        desc.append("boxes and labels fly in")
    present_hl = []                               # valence overlays currently showing
    last = len(steps) - 1
    for i, st in enumerate(steps):
        a = st["arrow"]
        effects = []
        if st["kind"] == "highlight":
            for ha in st["arrows"]:
                effects.append(eff(shp["hl"][id(ha)].shape_id, "entr", "fade", fill=True))
                present_hl.append(ha)
        else:
            sp = shp["arrows"][id(a)]
            effects.append(eff(sp.shape_id, "entr" if st["kind"] == "add" else "exit", fx, fill=True))
            if st["kind"] == "remove" and a in present_hl:      # its valence overlay leaves with it
                effects.append(eff(shp["hl"][id(a)].shape_id, "exit", "fade", fill=True))
                present_hl.remove(a)
        for lab, tb in shp["counts"].items():
            if last_of.get(lab) == i:
                effects.append(eff(tb.shape_id, "entr", "fade", text=True))
        for s2, tb in shp["captions"]:
            if s2 is st:
                effects.append(eff(tb.shape_id, "entr", "fade", text=True))
        for s2, tb in shp["states"]:
            if s2 is st:
                effects.append(eff(tb.shape_id, "entr", "fade", text=True))
        if i == last and shp["sum"] is not None:
            effects.append(eff(shp["sum"].shape_id, "entr", "fade", text=True))
        tl.click(*effects)
        desc.append(st["label"])
    if not steps and shp["sum"] is not None:
        tl.clicks[-1].append(eff(shp["sum"].shape_id, "entr", "fade", text=True))
    return tl, desc


def simulate(tl, static_visible, n):
    """Visible shape ids after n clicks. Shapes with any entrance effect start hidden."""
    entr = {e["spid"] for c in tl.clicks for e in c if e["cls"] == "entr"}
    vis = set(static_visible) - entr
    for c in tl.clicks[:n]:
        for e in c:
            (vis.add if e["cls"] == "entr" else vis.discard)(e["spid"])
    return vis


def _xml_effect(ids, e, node):
    """One PowerPoint effect as an XML string. ids is a counter list."""
    def nid():
        ids[0] += 1
        return ids[0]
    spid, ent = e["spid"], e["cls"] == "entr"
    preset = {"fly": 2, "fade": 10, "appear": 1}[e["effect"]]
    sub = 4 if e["effect"] == "fly" else 0
    grp = "" if e["grp"] else ' grpId="0"'
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    vis = lambda v, delay=0: (
        f'<p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{delay}"/>'
        f'</p:stCondLst></p:cTn>{tgt}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
        f'</p:cBhvr><p:to><p:strVal val="{v}"/></p:to></p:set>')
    head = (f'<p:par><p:cTn id="{nid()}" presetID="{preset}" presetClass="{e["cls"]}" presetSubtype="{sub}" '
            f'fill="hold"{grp} nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>')
    tail = '</p:childTnLst></p:cTn></p:par>'
    if e["effect"] == "appear":
        return head + vis("visible" if ent else "hidden") + tail
    if e["effect"] == "fade":
        fade = (f'<p:animEffect transition="{"in" if ent else "out"}" filter="fade"><p:cBhvr>'
                f'<p:cTn id="{{FID}}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>')
        if ent:
            body = vis("visible")
            fid = nid()
            return head + body + fade.replace("{FID}", str(fid)) + tail
        fid = nid()
        return head + fade.replace("{FID}", str(fid)) + vis("hidden", 399) + tail
    # fly in from bottom
    s1 = vis("visible")
    def anim(attr, a, z):
        return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{nid()}" dur="500" '
                f'fill="hold"/>{tgt}<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr>'
                f'<p:tavLst><p:tav tm="0"><p:val><p:strVal val="{a}"/></p:val></p:tav><p:tav tm="100000">'
                f'<p:val><p:strVal val="{z}"/></p:val></p:tav></p:tavLst></p:anim>')
    return head + s1 + anim("ppt_x", "#ppt_x", "#ppt_x") + anim("ppt_y", "1+#ppt_h/2", "#ppt_y") + tail


def timing_xml(tl):
    ids = [2]                                    # 1 = root, 2 = main sequence
    clicks = []
    for c in tl.clicks:
        outer, inner = None, None
        outer = (ids.__setitem__(0, ids[0] + 1), ids[0])[1]
        inner = (ids.__setitem__(0, ids[0] + 1), ids[0])[1]
        effs = "".join(_xml_effect(ids, e, "clickEffect" if i == 0 else "withEffect") for i, e in enumerate(c))
        clicks.append(
            f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
            f'<p:childTnLst><p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            f'<p:childTnLst>{effs}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    bld, seen = [], set()
    for c in tl.clicks:
        for e in c:
            if e["grp"] or e["spid"] in seen:
                continue
            seen.add(e["spid"])
            bld.append(f'<p:bldP spid="{e["spid"]}" grpId="0"' + (' animBg="1"' if e["fill"] else "") + "/>")
    xml = (f'<p:timing xmlns:p="{NS_P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" '
           f'nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" '
           f'nodeType="mainSeq"><p:childTnLst>{"".join(clicks)}</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
           f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
           + (f'<p:bldLst>{"".join(bld)}</p:bldLst>' if bld else "") + '</p:timing>')
    return xml



# --------------------------------------------------------------------------- text-line build
def set_descr(shape, text):
    """Alt text: the descr attribute of the shape's cNvPr (works for shapes, groups and pictures)."""
    el = shape._element
    nv = next(el.iter(f"{{{NS_P}}}cNvPr"))
    nv.set("descr", text)


def process_text_line(slide, b, region, T, chk):
    """A build of native text boxes, one per configuration segment. Returns (timeline, desc, shp).

        click 1   title, every segment of the longhand line, and the sum of its superscripts appear
        click 2   the segments the noble gas covers are highlighted (Primary fill, Asphalt outline, as the
                  valence highlight in the ion builds) with the first caption
        click 3   the highlight and those segments leave (exit fade); [Ar] appears in the highlight style,
                  right-aligned in the space they occupied so it sits against the remaining segments
        click 4   the finished shorthand line appears as the result row
    Effects are the validated ones: fade in, fade out. Nothing moves. Chemistry is computed from the
    checker: ground state, previous noble gas, the prefix of segments that equals that gas."""
    ink, white = T["ground"]["asphalt"], T["ground"]["white"]
    font = T["fonts"]["body"]
    deep = T["courses"][b["_course"]]["primaryDeep"]
    accent = T["courses"][b["_course"]]["primary"]
    x0, y0, W, H = region["x"], region["y"], region["w"], region["h"]
    pad, line, gap = 0.12, 0.28, 0.04
    cfg = b["longhand"]
    core_sym, parts = chk.parse(cfg)
    assert core_sym is None, "longhand must not use a noble-gas core"
    z = b["z"]
    assert chk.electrons(cfg) == z, f"longhand totals {chk.electrons(cfg)}, Z = {z}"
    assert chk.check_config(cfg, z) == [], f"longhand breaks a rule: {chk.check_config(cfg, z)}"
    core = b["core"]
    want = max(nz for _, nz in chk.NOBLE if nz < z)
    assert chk.NOBLE_Z[core] == want, f"{core} is not the noble gas BEFORE Z = {z} (that is {want} electrons)"
    run, k = 0, None
    for i, (_, n) in enumerate(parts):
        run += n
        if run == want:
            k = i + 1
            break
    assert k, f"no prefix of the longhand sums to {core}'s {want} electrons"
    assert dict(parts[:k]) == chk.ground_state(want), f"the first {k} segments are not {core}'s configuration"
    tail = parts[k:]
    short = f"[{core}] " + " ".join(f"{a}{n}" for a, n in tail)
    assert short == chk.shorthand_string(z), f"shorthand {short!r} != checker {chk.shorthand_string(z)!r}"
    seg_txt = [f"{a}^{{{n}}}" for a, n in parts]
    ar_txt = f"[{core}]"
    res_txt = b.get("result_text") or chk.mark(short)
    assert plain(res_txt) == short
    sum_txt = "+".join(str(n) for _, n in parts) + f" = {z}"
    caps = [b["captions"]["highlight"], b["captions"]["replace"]]
    inner = W - 2 * pad
    for padx, g in ((0.07, 0.09), (0.05, 0.07)):
        ws = [text_w(t, 18) + 2 * padx for t in seg_txt]
        row_w = sum(ws) + g * (len(ws) - 1)
        if row_w <= inner:
            break
    assert row_w <= inner, f"the longhand line needs {row_w:.2f} in at 18 pt; the region has {inner:.2f}. Split the slide."
    n_cap = [wrap_count(c, inner, FLOOR_PT) for c in caps]
    assert all(n <= 2 for n in n_cap), f"a caption needs more than two lines: {n_cap}"
    title_h, chip_h, sum_h, res_h = 0.30, 0.44, 0.30, 0.32
    cap_h = sum(n * line for n in n_cap)
    stack = title_h + chip_h + sum_h + cap_h + res_h + 5 * gap
    assert stack <= H - 2 * pad, f"text-line build needs {stack:.2f} in, the region has {H - 2 * pad:.2f}"
    y = y0 + (H - stack) / 2
    out = {"backdrop": None, "group": None, "arrows": {}, "hl": {}, "captions": [], "sum": None,
           "states": [], "extra": [], "counts": {}}
    bd = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(x0 * IN)), Emu(int(y0 * IN)),
                                Emu(int(W * IN)), Emu(int(H * IN)))
    bd.name = "BUILD backdrop"
    bd.fill.solid(); bd.fill.fore_color.rgb = rgb(white); bd.line.fill.background()
    bd.shadow.inherit = False
    out["backdrop"] = bd
    sh = slide.shapes
    title = textbox(sh, x0 + pad, y, inner, title_h, b["title"], font, 18, ink, True, name="BUILD title")
    cy = y + title_h + gap
    xs = x0 + pad + (inner - row_w) / 2
    xl = [xs]
    for w_ in ws[:-1]:
        xl.append(xl[-1] + w_ + g)
    # the highlight block sits BEHIND the segments of the core: a text-free rectangle, added first
    hx0, hx1 = xl[0] - 0.04, xl[k - 1] + ws[k - 1] + 0.04
    hl = sh.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(hx0 * IN)), Emu(int((cy - 0.02) * IN)),
                      Emu(int((hx1 - hx0) * IN)), Emu(int((chip_h + 0.04) * IN)))
    hl.name = "BUILD highlight core"
    hl.fill.solid(); hl.fill.fore_color.rgb = rgb(accent)
    hl.line.color.rgb = rgb(ink); hl.line.width = Pt(1.5)
    hl.shadow.inherit = False
    chips = []
    for i, t in enumerate(seg_txt):
        chips.append(textbox(sh, xl[i], cy, ws[i], chip_h, t, font, 18, ink, True,
                             name=f"BUILD segment {parts[i][0]}", wrap=False))
    sm = textbox(sh, x0 + pad, cy + chip_h + gap, inner, sum_h, sum_txt, font, 18, ink, True, name="BUILD sum")
    yy = cy + chip_h + gap + sum_h + gap
    cap_boxes = []
    for c, n in zip(caps, n_cap):
        cap_boxes.append(textbox(sh, x0 + pad, yy, inner, n * line - 0.02, c, font, FLOOR_PT, deep, True,
                                 name="BUILD caption"))
        yy += n * line
    res = textbox(sh, x0 + pad, yy + gap, inner, res_h, res_txt, font, 18, ink, True, name="BUILD result")
    # [Ar]: in the highlight style, right-aligned in the block's span so it sits against the tail
    war = text_w(ar_txt, 18) + 2 * 0.07
    assert war <= hx1 - hx0, "the noble-gas symbol is wider than the block it replaces"
    ar_x = xl[k] - g - war
    ar = textbox(sh, ar_x, cy + 0.02, war, chip_h - 0.04, ar_txt, font, 18, ink, True, name=f"BUILD {core} bracket",
                 wrap=False, fill=accent, outline=ink)
    out["captions"] = [(None, c) for c in cap_boxes]
    out["extra"] = [title, hl] + chips + [sm, ar, res]
    tl = Timeline()
    f_ = lambda shape, cls, text=False, fill=False: eff(shape.shape_id, cls, "fade", text=text, fill=fill)
    tl.click(f_(title, "entr", True), *[f_(c, "entr", True) for c in chips], f_(sm, "entr", True))
    tl.click(f_(hl, "entr", fill=True), f_(cap_boxes[0], "entr", True))
    tl.click(f_(hl, "exit", fill=True), *[f_(c, "exit", True) for c in chips[:k]],
             f_(ar, "entr", True, True), f_(cap_boxes[1], "entr", True))
    tl.click(f_(res, "entr", True))
    desc = [f"the longhand line, one text box per segment, and its sum ({sum_txt}) appear",
            f"the first {k} segments (the {core} part, {want} electrons) are highlighted: {caps[0]}",
            f"the highlighted segments leave and {ar_txt} appears in their place: {caps[1]}",
            f"the finished line appears: {short}"]
    # replay: what is showing after each click
    static = [bd.shape_id]
    every = [x.shape_id for x in all_shapes(out)]
    vis = lambda n: simulate(tl, static + every[1:], n)
    ids = lambda L: {x.shape_id for x in L}
    v1, v2, v3, v4 = vis(1), vis(2), vis(3), vis(4)
    assert ids(chips) <= v1 and title.shape_id in v1 and sm.shape_id in v1
    assert not (ids([hl, ar, res] + cap_boxes) & v1), "something arrives before its click"
    assert hl.shape_id in v2 and cap_boxes[0].shape_id in v2 and ids(chips) <= v2
    assert not (ids(chips[:k]) | {hl.shape_id}) & v3 and ar.shape_id in v3 and ids(chips[k:]) <= v3
    assert ids(chips[k:]) | {ar.shape_id, res.shape_id, title.shape_id, sm.shape_id} | ids(cap_boxes) <= v4
    assert not (ids(chips[:k]) | {hl.shape_id}) & v4
    # what remains reads as the shorthand: the bracket, then the tail, left to right
    shown = sorted([(ar.left, ar_txt)] + [(chips[i].left, plain(seg_txt[i])) for i in range(k, len(seg_txt))])
    assert " ".join(t for _, t in shown) == plain(res_txt), "the line that remains does not equal the result row"
    if b.get("alt_text"):
        set_descr(bd, f"Text build for {b.get('element', 'the element')}, Z = {z}: the longhand "
                  f"{cfg}, whose first {k} segments ({want} electrons) equal {core}, then {core} in brackets "
                  f"standing in for them, then the finished shorthand {short}.")
    return tl, desc, out, ids(chips[:k]) | {hl.shape_id}


def process_text_line_static(slide, b, region, T, chk):
    """The print / handout state of a text-line build: no animation, only the coherent final state.

        title
        Longhand:   the full configuration as one line of text
        Shorthand:  [Ar] in the highlight style, then the segments the noble gas does not cover
        caption     what the noble-gas symbol stands for (computed; "static_caption" overrides)
        check       core electrons + the remaining superscripts = Z
    Nothing is hidden and nothing overlaps: the click-2 caption, the highlight block and the covered
    segments of the animated build are not drawn. Chemistry is computed and asserted as in the animated build."""
    ink, white = T["ground"]["asphalt"], T["ground"]["white"]
    font = T["fonts"]["body"]
    deep = T["courses"][b["_course"]]["primaryDeep"]
    accent = T["courses"][b["_course"]]["primary"]
    x0, y0, W, H = region["x"], region["y"], region["w"], region["h"]
    pad, line, gap = 0.12, 0.28, 0.04
    cfg = b["longhand"]
    core_sym, parts = chk.parse(cfg)
    assert core_sym is None, "longhand must not use a noble-gas core"
    z = b["z"]
    assert chk.electrons(cfg) == z and chk.check_config(cfg, z) == [], "longhand fails the checker"
    core = b["core"]
    want = max(nz for _, nz in chk.NOBLE if nz < z)
    assert chk.NOBLE_Z[core] == want, f"{core} is not the noble gas BEFORE Z = {z}"
    run, k = 0, None
    for i, (_, n) in enumerate(parts):
        run += n
        if run == want:
            k = i + 1
            break
    assert k and dict(parts[:k]) == chk.ground_state(want), "the prefix is not the noble gas"
    tail = parts[k:]
    short = f"[{core}] " + " ".join(f"{a}{n}" for a, n in tail)
    assert short == chk.shorthand_string(z), f"shorthand {short!r} != checker {chk.shorthand_string(z)!r}"
    inner = W - 2 * pad
    long_txt = " ".join(f"{a}^{{{n}}}" for a, n in parts)
    ar_txt = f"[{core}]"
    seg_txt = [f"{a}^{{{n}}}" for a, n in tail]
    cap = b.get("static_caption") or (f"{ar_txt} stands for the first {want} electrons: the noble gas "
                                      f"before {b.get('element') or chk.SYMBOL[z]}")
    chk_txt = f"{want} + " + " + ".join(str(n) for _, n in tail) + f" = {z} electrons"
    assert want + sum(n for _, n in tail) == z
    assert text_w(long_txt, 18) * 1.04 <= inner, f"the longhand row needs {text_w(long_txt, 18):.2f} in; the region has {inner:.2f}"
    war = text_w(ar_txt, 18) + 2 * 0.07
    wseg = [text_w(t, 18) + 2 * 0.05 for t in seg_txt]
    gsp = 0.07
    row_w = war + sum(wseg) + gsp * len(seg_txt)
    assert row_w <= inner, f"the shorthand row needs {row_w:.2f} in; the region has {inner:.2f}"
    n_cap = wrap_count(cap, inner, FLOOR_PT)
    assert n_cap <= 2, f"the caption needs {n_cap} lines"
    title_h, lab_h, row_h, chip_h, sum_h = 0.30, 0.26, 0.34, 0.44, 0.30
    stack = title_h + 2 * lab_h + row_h + chip_h + n_cap * line + sum_h + 6 * gap
    assert stack <= H - 2 * pad, f"static text-line build needs {stack:.2f} in, the region has {H - 2 * pad:.2f}"
    y = y0 + (H - stack) / 2
    out = {"backdrop": None, "group": None, "arrows": {}, "hl": {}, "captions": [], "sum": None,
           "states": [], "extra": [], "counts": {}}
    bd = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(x0 * IN)), Emu(int(y0 * IN)),
                                Emu(int(W * IN)), Emu(int(H * IN)))
    bd.name = "BUILD backdrop"
    bd.fill.solid(); bd.fill.fore_color.rgb = rgb(white); bd.line.fill.background()
    bd.shadow.inherit = False
    out["backdrop"] = bd
    sh = slide.shapes
    textbox(sh, x0 + pad, y, inner, title_h, b["title"], font, 18, ink, True, name="BUILD title")
    yy = y + title_h + gap
    textbox(sh, x0 + pad, yy, inner, lab_h, "Longhand", font, FLOOR_PT, deep, True, name="BUILD caption")
    yy += lab_h + gap
    textbox(sh, x0 + pad, yy, inner, row_h, long_txt, font, 18, ink, True, name="BUILD longhand")
    yy += row_h + gap
    textbox(sh, x0 + pad, yy, inner, lab_h, "Shorthand", font, FLOOR_PT, deep, True, name="BUILD caption")
    yy += lab_h + gap
    x = x0 + pad + (inner - row_w) / 2
    textbox(sh, x, yy + 0.02, war, chip_h - 0.04, ar_txt, font, 18, ink, True, name=f"BUILD {core} bracket",
            wrap=False, fill=accent, outline=ink)
    x += war + gsp
    for (a, n), t, w_ in zip(tail, seg_txt, wseg):
        textbox(sh, x, yy, w_, chip_h, t, font, 18, ink, True, name=f"BUILD segment {a}", wrap=False)
        x += w_ + gsp
    yy += chip_h + gap
    textbox(sh, x0 + pad, yy, inner, n_cap * line - 0.02, cap, font, FLOOR_PT, deep, True, name="BUILD caption")
    yy += n_cap * line + gap
    textbox(sh, x0 + pad, yy, inner, sum_h, chk_txt, font, 18, ink, True, name="BUILD sum")
    if b.get("alt_text"):
        set_descr(bd, f"Text build for {b.get('element', 'the element')}, Z = {z}: the longhand {cfg}, then the "
                  f"shorthand {short}, where {ar_txt} stands for the first {want} electrons.")
    return out


def validate_timing(slide, expected_clicks):
    """Structural validation of the p:timing that was written."""
    el = slide._element
    t = el.find(f"{{{NS_P}}}timing")
    assert t is not None, "no p:timing"
    kids = [c.tag.split("}")[1] for c in el]
    assert kids.index("timing") > kids.index("clrMapOvr"), "p:timing must follow p:clrMapOvr"
    ids = [int(c.get("id")) for c in t.iter(f"{{{NS_P}}}cTn")]
    assert ids == list(range(1, len(ids) + 1)), "cTn ids are not unique, ascending and gap-free"
    shape_ids = [int(c.get("id")) for c in el.iter(f"{{{NS_P}}}cNvPr")]
    assert len(shape_ids) == len(set(shape_ids)), "duplicate shape ids on the slide"
    for tg in t.iter(f"{{{NS_P}}}spTgt"):
        assert int(tg.get("spid")) in shape_ids, f"animation targets missing shape {tg.get('spid')}"
    for bp in t.iter(f"{{{NS_P}}}bldP"):
        assert int(bp.get("spid")) in shape_ids, "bldP targets a missing shape"
    seq = t.find(f".//{{{NS_P}}}cTn[@nodeType='mainSeq']")
    clicks = seq.find(f"{{{NS_P}}}childTnLst")
    assert len(clicks) == expected_clicks, f"{len(clicks)} click pars, expected {expected_clicks}"
    n_click = len(t.findall(f".//{{{NS_P}}}cTn[@nodeType='clickEffect']"))
    assert n_click == expected_clicks, f"{n_click} clickEffect nodes, expected {expected_clicks}"
    for c in clicks:
        first = c.findall(f".//{{{NS_P}}}cTn[@presetID]")
        assert first[0].get("nodeType") == "clickEffect"
        assert all(f.get("nodeType") == "withEffect" for f in first[1:])
    return len(ids)


# --------------------------------------------------------------------------- slide editing
def remove_pictures_in(slide, region):
    removed = []
    x0, y0, x1, y1 = (int(v * IN) for v in (region["x"], region["y"],
                                            region["x"] + region["w"], region["y"] + region["h"]))
    for sh in list(slide.shapes):
        if sh.shape_type == 13:                                # PICTURE
            cx, cy = sh.left + sh.width // 2, sh.top + sh.height // 2
            if x0 <= cx <= x1 and y0 <= cy <= y1:
                sh._element.getparent().remove(sh._element)
                removed.append(sh.name)
    return removed


def append_note(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = (tf.text.rstrip() + "\n" if tf.text.strip() else "") + text


def all_shapes(shp):
    """Every shape a build drew, backdrop first."""
    return ([shp["backdrop"]] + ([shp["group"]] if shp["group"] is not None else [])
            + list(shp["arrows"].values()) + list(shp["hl"].values())
            + [tb for _, tb in shp["captions"]] + [tb for _, tb in shp["states"]]
            + ([shp["sum"]] if shp["sum"] is not None else []) + list(shp.get("extra", [])))


def all_shape_ids(shp):
    return [sh.shape_id for sh in all_shapes(shp)]


def drop(shape):
    shape._element.getparent().remove(shape._element)


def static_spec(b):
    """The spec of a build as the static (print / handout) copy draws it: a copy, never the original.
    Ion steps keep their captions and the LAST state line; the species passed through on the way are animated-only."""
    import copy as _copy
    b = _copy.deepcopy(b)
    if b.get("static_title"):
        b["title"] = b["static_title"]
    texts = [st for st in b.get("steps") or [] if st.get("text")]
    if b.get("steps") and not b.get("story"):
        for st in texts[:-1]:
            del st["text"]
    return b


def process(deck, specfile, out, preview=None, only=None, static_mode=False):
    spec = json.load(open(specfile, encoding="utf-8"))
    course = spec["course"]
    T = load_tokens(course)
    chk = load_checker()
    prs = Presentation(deck)
    report = []
    per_slide = {}                       # slide number -> [(timeline, desc)], in spec order
    for b in spec["builds"]:
        if only and b["slide"] != only:
            continue
        b["_course"] = course
        slide = prs.slides[b["slide"] - 1]
        layout = slide.slide_layout.name
        assert layout == b["layout"], f"slide {b['slide']}: layout is {layout}, spec says {b['layout']}"
        slot = T["slots"][layout][b["region"]]
        region = dict(slot)
        if b.get("band"):
            i_, n_ = b["band"]
            assert 1 <= i_ <= n_, f"bad band {b['band']}"
            region["h"] = region["h"] / n_
            region["y"] = region["y"] + region["h"] * (i_ - 1)
        b["alt_text"] = bool(spec.get("alt_text"))
        earlier = per_slide.setdefault(b["slide"], [])
        if static_mode:
            b = static_spec(b)
        if b.get("type") == "text_line" and static_mode:
            removed = remove_pictures_in(slide, slot)
            shp = process_text_line_static(slide, b, region, T, chk)
            earlier.append((None, ["the longhand line, then the shorthand line with the noble-gas symbol"],
                            b["region"], b.get("name"), b, shp, None, None, 0, removed, 0))
            continue
        if b.get("type") == "text_line":
            removed = remove_pictures_in(slide, slot)
            tl, desc, shp, _gone = process_text_line(slide, b, region, T, chk)
            n = len(tl.clicks)
            earlier.append((tl, desc, b["region"], b.get("name"), b, shp, [shp["backdrop"].shape_id],
                            all_shape_ids(shp), n, removed, 0))
            continue
        validate_atom(b, chk)
        initial, steps, arrows = plan_steps(b)
        if static_mode and b.get("steps"):
            fc = {s_["label"]: s_["electrons"] for s_ in b["sublevels"]}
            for st in steps:
                if st["kind"] != "highlight":
                    fc[st["arrow"].sub] += 1 if st["kind"] == "add" else -1
            b["_final_counts"] = fc
        removed = remove_pictures_in(slide, slot)
        shp = draw(slide, b, region, T, initial, steps, arrows)
        tl, desc = build_timeline(b, shp, initial, steps)
        n = len(tl.clicks)
        static = [shp["backdrop"].shape_id]
        every = all_shape_ids(shp)
        # final-state assertion: replay the clicks and compare with the target
        final_vis = simulate(tl, static + every[1:], n)
        if b.get("steps"):
            want, want_hl = set(initial), set()
            for st in steps:
                if st["kind"] == "highlight":
                    want_hl |= set(st["arrows"])
                else:
                    (want.add if st["kind"] == "add" else want.discard)(st["arrow"])
            want_hl &= want                            # an overlay leaves with its arrow
        else:
            want, want_hl = set(arrows), set()
        check_boxes(list(want), b["sublevels"])
        got = {a for a in arrows if shp["arrows"][id(a)].shape_id in final_vis}
        assert got == want, f"slide {b['slide']}: final state differs from the computed target"
        if b.get("predicted"):
            # a "move": the sublevels are what the rules predict; the steps must land on the measured state
            fin = {}
            for a in want:
                fin[a.sub] = fin.get(a.sub, 0) + 1
            for sub_ in b["sublevels"]:
                fin.setdefault(sub_["label"], 0)
            exp = chk.ground_state(b["z"])
            assert all(fin[k_] == exp.get(k_, 0) for k_ in fin), \
                f"slide {b['slide']}: the steps end at {fin}, the measured ground state of Z={b['z']} is {exp}"
            assert sum(fin.values()) == sum(sub_["electrons"] for sub_ in b["sublevels"]), "a move must not change the count"
        got_hl = {a for a in arrows if id(a) in shp["hl"] and shp["hl"][id(a)].shape_id in final_vis}
        assert got_hl == want_hl, f"slide {b['slide']}: final valence highlights differ from the target"
        if shp["states"]:
            shown = [tb for _, tb in shp["states"] if tb.shape_id in final_vis]
            assert len(shown) == len(shp["states"]), f"slide {b['slide']}: a state line is not showing at the end"
        total_final = len(got)
        if b["alt_text"]:
            set_descr(shp["group"], describe_diagram(b, initial, [a for a in arrows if a in want], arrows))
            for a in arrows:
                set_descr(shp["arrows"][id(a)], f"{'Up' if a.spin == 'u' else 'Down'} arrow in {a.sub} box {a.box + 1}")
        earlier.append((tl, desc, b["region"], b.get("name"), b, shp, static, every, n, removed, total_final))
    for slide_no, items in per_slide.items():
        slide = prs.slides[slide_no - 1]
        if static_mode:
            for tl, desc, _, name, b, shp, static, every, n, _, _ in items:
                if tl is not None:                       # a computed final state: drop everything else
                    vis = simulate(tl, static + every[1:], n)
                    for shape in all_shapes(shp):
                        if shape.shape_id not in vis or shape in shp["hl"].values():
                            drop(shape)
                what = ("final state of the click build" + (f" ({name})" if name else ""))
                append_note(slide, f"STATIC FINAL STATE: print and handout copy, no animation. This slide shows the {what}; "
                                   "the animated deck shows it one click at a time.")
                report.append((slide_no, "static final state", n))
            continue
        if preview is not None:
            left = preview
            for tl, desc, _, _, b, shp, static, every, n, _, _ in items:
                k = max(0, min(left, n))
                left -= k
                vis = simulate(tl, static + every[1:], k)
                for shape in all_shapes(shp):
                    if shape.shape_id not in vis:
                        drop(shape)
                report.append((slide_no, f"preview after {k} of {n} clicks"))
            continue
        merged = Timeline()
        notes = []
        for tl, desc, _, name, *_ in items:
            merged.clicks += tl.clicks
            notes += [f"{name}: {d}" if name else d for d in desc]
        el = slide._element
        tm = etree.fromstring(timing_xml(merged))
        anchor = el.find(f"{{{NS_P}}}clrMapOvr")
        tr = el.find(f"{{{NS_P}}}transition")
        (tr if tr is not None else anchor).addnext(tm)
        n = len(merged.clicks)
        nodes = validate_timing(slide, n)
        append_note(slide, f"CLICK BUILD: {n} clicks: " + "; ".join(f"{i + 1} {d}" for i, d in enumerate(notes)) + ".")
        for tl, desc, _, _, b, _, _, _, nn, removed, total_final in items:
            report.append((slide_no, nn, nodes, removed, total_final))
    if preview is None:
        for pa in spec.get("picture_alt", []):
            pics = [sh for sh in prs.slides[pa["slide"] - 1].shapes if sh.shape_type == 13]
            assert pics, f"picture_alt: slide {pa['slide']} has no picture"
            assert len(pics) == 1 or pa.get("index") is not None, \
                f"picture_alt: slide {pa['slide']} has {len(pics)} pictures; give an \"index\""
            set_descr(pics[pa.get("index", 0)], pa["descr"])
            report.append((pa["slide"], "picture alt text set"))
    prs.save(out)
    return report


def main(argv):
    pos = [a for a in argv if not a.startswith("--")]
    flags = argv
    static_mode = "--static" in flags
    preview = only = None
    if "--preview" in flags:
        preview = int(flags[flags.index("--preview") + 1]); pos.remove(str(preview))
    if "--slide" in flags:
        only = int(flags[flags.index("--slide") + 1]); pos.remove(str(only))
    if len(pos) < 2:
        sys.exit(__doc__)
    deck, specfile = pos[0], pos[1]
    assert not (static_mode and preview is not None), "--static and --preview are separate modes"
    out = pos[2] if len(pos) > 2 else re.sub(r"\.pptx$", ".built.pptx", deck)
    rep = process(deck, specfile, out, preview, only, static_mode)
    for r in rep:
        print("slide", *r)
    print("wrote", out)


if __name__ == "__main__":
    main(sys.argv[1:])
