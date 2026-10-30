#!/usr/bin/env python3
"""Build the Physics S2.3 "Feed the Monkey" Gizmo activity (student sheet and key), one page each.

    python3 templates/activity/build_gizmo_monkey.py OUTDIR

A short completion activity for a day the teacher is out: students run the Explore Learning
Gizmo, copy what the screen shows, and circle or write one short answer per part. Nothing is
calculated, so nothing here depends on a number the Gizmo might round differently.
Colour comes from brand/tokens.json via _shull_docx; none is typed here.
"""
import os, subprocess, sys, tempfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from _shull_name import document_name                                   # noqa: E402
from _shull_docx import (T, Palette, borders, para, fix_widths, gap, one_cell,   # noqa: E402
                         page_setup, running_footer, trim_tail, schema_order, tight_cells)

KEY_INK = T["keyInk"]["pen"]["hex"]
TITLE = "FEED THE MONKEY"
SUB = "PROJECTILE MOTION GIZMO  ·  ABOUT 20 MINUTES"

SETUP = [
    "Open the Gizmo: Explore Learning › Feed the Monkey (Projectile Motion).",
    "Set Atmosphere to None. Check Show velocity components.",
    "Set v initial to 17 m/s and the angle to 45°. Press play.",
]
TABLE_ROWS = ["Just after launch", "At the highest point", "Just before landing"]
ANGLES = ["30°", "45°", "60°", "75°"]


def new_doc(pal, footer):
    doc = Document()
    sec = page_setup(doc, margin_lr=0.55, margin_tb=0.5, base_pt=10)
    running_footer(sec, footer, pal)
    return doc


def header(doc, pal, key=False):
    t = doc.add_table(rows=1, cols=2)
    fix_widths(t, [5.6, 1.8])
    a, b = t.rows[0].cells
    for c in (a, b):
        borders(c, pal.display, sz=18, edges=("bottom",))
    para(a, "SHULL SCIENCE  ·  JAMES A. GARFIELD LOCAL SCHOOLS", 7.5, bold=True, color=pal.accent,
         caps_track=True, first=True)
    para(a, TITLE, 16, bold=True, color=pal.ink)
    para(a, SUB + ("  ·  TEACHER KEY" if key else ""), 8.5, bold=True,
         color=KEY_INK if key else pal.label, caps_track=True)
    p = para(b, "U02 / S02.3", 9, bold=True, color=pal.accent, caps_track=True, first=True)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def fields_row(doc, pal):
    gap(doc, 4)
    t = doc.add_table(rows=1, cols=4)
    fix_widths(t, [3.0, 1.7, 1.2, 1.5])
    for c, lab in zip(t.rows[0].cells, ("NAME", "DATE", "PERIOD", "DONE")):
        borders(c, pal.hair, sz=6)
        para(c, lab, 8.5, bold=True, color=pal.label, caps_track=True, first=True)
        c.paragraphs[0].paragraph_format.space_after = Pt(9)


def part(doc, pal, n, title, lines):
    gap(doc, 6)
    c = one_cell(doc, 7.4)
    borders(c, pal.hair, sz=6)
    para(c, f"PART {n}  ·  {title}", 8.5, bold=True, color=pal.accent, caps_track=True, first=True)
    for ln in lines:
        para(c, ln, 10)
    return c


def grid(doc, pal, head, rows, widths, key_rows=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    fix_widths(t, widths)
    tight_cells(t, top=50, bottom=50, left=70, right=70)
    for c, h in zip(t.rows[0].cells, head):
        borders(c, pal.display, sz=8)
        para(c, h, 8, bold=True, color=pal.label, caps_track=True, first=True)
    for i, r in enumerate(rows, 1):
        for j, c in enumerate(t.rows[i].cells):
            borders(c, pal.hair, sz=6)
            txt = r[j]
            if j > 0 and key_rows:
                para(c, key_rows[i - 1][j - 1], 10, color=KEY_INK, first=True)
            else:
                para(c, txt, 10, first=True)
            c.paragraphs[0].paragraph_format.space_after = Pt(6)
    return t


def body(doc, pal, key=False):
    # setup
    gap(doc, 6)
    c = one_cell(doc, 7.4); borders(c, pal.display, sz=18, edges=("left",))
    para(c, "SET UP", 8.5, bold=True, color=pal.accent, caps_track=True, first=True)
    for i, s in enumerate(SETUP, 1):
        para(c, f"{i}.  {s}", 10)
    para(c, "This is a completion activity. Fill in what you see and answer each part. "
            "Close counts.", 9, color=pal.label)

    # part 1
    c = part(doc, pal, 1, "WATCH ONE SHOT", [
        "Pause the Gizmo near each moment and copy the two velocity readouts."])
    kr = [["≈ 12 m/s", "≈ 12 m/s"], ["≈ 12 m/s", "0 m/s"], ["≈ 12 m/s", "≈ −12 m/s"]] if key else None
    grid(doc, pal, ["MOMENT", "vx (m/s)", "vy (m/s)"], [[r, "", ""] for r in TABLE_ROWS],
         [3.0, 2.2, 2.2], kr)
    q = one_cell(doc, 7.4)
    para(q, "Which one stayed the same the whole flight?   vx    vy     (circle)", 10, first=True)
    if key:
        para(q, "vx. Nothing pushes sideways, so it never changes.", 9, color=KEY_INK)
    para(q, "What was vy at the highest point?   ______________", 10)
    if key:
        para(q, "Zero. The arrow stops rising and has not started falling.", 9, color=KEY_INK)

    # part 2
    part(doc, pal, 2, "CHANGE THE ANGLE", [
        "Keep v initial at 17 m/s. Launch at each angle and copy the time when it lands."])
    kr = [["≈ 1.7 s"], ["≈ 2.4 s"], ["≈ 3.0 s"], ["≈ 3.3 s"]] if key else None
    grid(doc, pal, ["ANGLE", "TIME IN THE AIR (s)"], [[a, ""] for a in ANGLES], [3.0, 4.4], kr)
    q = one_cell(doc, 7.4)
    para(q, "A bigger angle keeps the shot in the air:   shorter    longer     (circle)", 10, first=True)
    if key:
        para(q, "Longer. The extra upward speed takes longer to use up. Accept the trend; "
                "the exact times depend on where the Gizmo stops the run.", 9, color=KEY_INK)

    # part 3
    part(doc, pal, 3, "FEED THE MONKEY", [
        "Check Show target line and turn the cannon until the line points at the monkey. Launch. "
        "Then raise v initial and launch again."])
    q = one_cell(doc, 7.4)
    para(q, "Did the shot still reach the monkey the second time?   yes    no     (circle)", 10, first=True)
    if key:
        para(q, "Yes. Aimed along the line, it reaches the monkey at any speed fast enough to get there.",
             9, color=KEY_INK)
    para(q, "One sentence: why does aiming straight at the monkey work?", 10)
    if key:
        para(q, "Sample: the monkey and the shot both fall the same distance in the same time, so they "
                "stay on the line. Accept any answer that says gravity pulls both down equally.",
             9, color=KEY_INK)
    else:
        for _ in range(1):
            p = para(q, " ", 10)
            from _shull_docx import rule_lines
        rule_lines(q, 2, pal.hair)

    # wrap-up
    gap(doc, 6)
    c = one_cell(doc, 7.4); borders(c, pal.display, sz=18, edges=("left",))
    para(c, "WRAP-UP", 8.5, bold=True, color=pal.accent, caps_track=True, first=True)
    para(c, "During a projectile’s flight, the speed that stays constant is the ________________ "
            "velocity.", 10)
    if key:
        para(c, "horizontal (vx)", 9, color=KEY_INK)


def build(out, key=False):
    pal = Palette("physics")
    doc = new_doc(pal, "SHULL SCIENCE / PHYSICS / UNIT 02 / S02.3 / FEED THE MONKEY"
                  + (" / TEACHER KEY" if key else ""))
    header(doc, pal, key)
    if not key:
        fields_row(doc, pal)
    else:
        gap(doc, 4)
        c = one_cell(doc, 7.4); borders(c, pal.hair)
        para(c, "Grade for completion only: every blank filled, every question circled or answered. "
                "Values are for Atmosphere: None, 17 m/s, 45°; the Gizmo’s own readouts rule. Part 2 "
                "times assume it ends the run at launch height.", 8.5, color=pal.label, first=True)
    body(doc, pal, key)
    trim_tail(doc)
    schema_order(doc)
    doc.save(out)


def pages(path):
    with tempfile.TemporaryDirectory() as td:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", td, path],
                       check=True, capture_output=True)
        pdf = os.path.join(td, os.path.splitext(os.path.basename(path))[0] + ".pdf")
        out = subprocess.run(["pdfinfo", pdf], check=True, capture_output=True, text=True).stdout
    return int([l for l in out.splitlines() if l.startswith("Pages")][0].split()[1])


def main():
    outdir = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "out")
    os.makedirs(outdir, exist_ok=True)
    bad = 0
    for key in (False, True):
        p = os.path.join(outdir, document_name("Activity", ["2.3"], 2, key=key))
        build(p, key)
        n = pages(p)
        print("ok" if n == 1 else "FAIL", n, "page", p)
        bad += n != 1
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
