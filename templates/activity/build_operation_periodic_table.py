#!/usr/bin/env python3
"""Build CHEM U03 Operation Periodic Table: the student handout (with cut-out cards) and the
separate teacher key, from one spec so the two cannot drift.

    python3 templates/activity/build_operation_periodic_table.py specs/chem_u03_operation_periodic_table.json OUTDIR

Teams get 25 element cards under letter code names (A-Y, none on its own symbol) and place
them on a blank 6 x 5 grid using the trends. Hydrogen is in on purpose: it is the card that
fits no family cleanly. Germanium is held out, so one cell stays empty and each team predicts what belongs
there - Mendeleev's eka-silicon. The key maps every code name back to its element.

The drawing primitives come from templates/_shull_docx.py, so this sheet shares its header,
rules and footer with the worksheet builder rather than keeping copies of them.
"""
import json, os, sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from _shull_docx import (   # noqa: E402
    Palette, known_sections, unit_title, borders, para, run, fix_widths, no_split, page_setup,
    running_footer, cell_margins, gap, one_cell, shade, rule_lines,
)
from _shull_name import document_name   # noqa: E402
sys.path.insert(0, HERE)
from _print_parts import TEXT_W, label, page_break, header, fields, boxed, plain   # noqa: E402


def refuse(msg):
    print("build_operation_periodic_table: " + msg, file=sys.stderr)
    sys.exit(1)


def fmt_mp(v):
    # A true minus, not a hyphen: "-189" reads as a dash on a printed card.
    return f"{v} °C".replace("-", "−") if isinstance(v, int) else v


def grid_table(doc, pal, spec, fill=None):
    """The blank 6 x 5 placement grid. No group or period labels: finding those is the task."""
    cols, rows = len(spec["grid"]["columns"]), len(spec["grid"]["rows"])
    w = round(TEXT_W / cols, 3)
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fix_widths(t, [w] * cols)
    for r in range(rows):
        no_split(t.rows[r], 0.62)
        for c in range(cols):
            cell = t.rows[r].cells[c]
            borders(cell, pal.ink, sz=6)
            if fill:
                txt = fill.get((r, c))
                if txt:
                    p = para(cell, txt[0], 10, bold=True, color=txt[2], first=True)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p = para(cell, txt[1], 8.5, color=pal.ink)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return t


def answer_lines(doc, pal, prompts, n=2):
    c = one_cell(doc)
    borders(c, pal.hair, sz=4)
    cell_margins(c, top=60, bottom=60, left=120, right=120)
    label(c, "AFTER THE REVEAL — ANSWER ON YOUR OWN", pal, first=True)
    for i, q in enumerate(prompts, 1):
        p = para(c, "", 9.5)
        run(p, f"{i}.  ", 9.5, bold=True, color=pal.ink)
        run(p, q, 9.5, color=pal.ink)
        rule_lines(c, n, pal.hair)


def card_rows(card):
    """(label, value, unit) in the order the rounds use them: ion sorts, radius orders."""
    return [("Ion it forms", card["ion"], ""),
            ("Atomic radius", str(card["radius"]), "pm"),
            ("Ionization energy", str(card["ie"]), "kJ/mol"),
            ("Electronegativity", card["en"], ""),
            ("Melting point", *(fmt_mp(card["mp"]).rsplit(" ", 1) if isinstance(card["mp"], int)
                               else (card["mp"], "")))]


CARD_ACROSS, CARD_ROWS, CARD_H, CARD_GUTTER = 3, 5, 1.78, 0.15


def card_face(cell, card, pal, inner_w):
    """One card. The letter is the largest thing on it, the five numbers line up in a right-hand
    column so a team can compare cards held side by side, and hairlines separate the rows.
    No fills: the card is cut from cardstock and the ink budget still applies."""
    head = cell.add_table(rows=1, cols=2)
    fix_widths(head, [0.55, inner_w - 0.55])
    lc, ac = head.rows[0].cells
    for c in (lc, ac):
        cell_margins(c, top=0, bottom=30, left=0, right=0)
    borders(lc, pal.hair, sz=4, edges=("bottom",))
    borders(ac, pal.hair, sz=4, edges=("bottom",))
    para(lc, card["code"], 28, bold=True, color=pal.accent, first=True)
    from docx.enum.table import WD_ALIGN_VERTICAL
    ac.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    p = para(ac, card["appearance"], 9, color=pal.ink, first=True)
    p.runs[0].italic = True
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT     # ends on the same edge as the value column
    rows = card_rows(card)
    t = cell.add_table(rows=len(rows), cols=2)
    fix_widths(t, [1.05, inner_w - 1.05])
    for r, (lab, val, unit) in enumerate(rows):
        lcell, vcell = t.rows[r].cells
        for c in (lcell, vcell):
            cell_margins(c, top=22, bottom=22, left=0, right=0)
            if r < len(rows) - 1:
                borders(c, pal.hair, sz=2, edges=("bottom",))
        from docx.enum.table import WD_ALIGN_VERTICAL as V
        lcell.vertical_alignment = vcell.vertical_alignment = V.CENTER
        para(lcell, lab, 8.5, color=pal.label, first=True)
        q = para(vcell, "", 10.5, first=True)
        q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run(q, val, 10.5 if len(val) < 14 else 9, bold=True, color=pal.ink)
        if unit:
            run(q, " " + unit, 8, color=pal.label)


def tidy_cell(cell):
    """python-docx puts an empty paragraph before the first nested table (the cell's own) and
    one after every nested table. Each is a full 10pt line, and three of them made a card a
    third taller than its content. Keep only the last, which Word requires, at 1pt."""
    from docx.oxml.ns import qn
    paras = [el for el in cell._tc if el.tag == qn("w:p")]
    for el in paras[:-1]:
        if not el.xpath(".//w:t"):
            cell._tc.remove(el)
    last = cell.paragraphs[-1]
    last.paragraph_format.space_after = Pt(0)
    last.paragraph_format.space_before = Pt(0)
    last.paragraph_format.line_spacing = Pt(1)


def cards_pages(doc, pal, spec):
    cards = spec["cards"] + [None]          # the last slot is the gap card
    per_page = CARD_ACROSS * CARD_ROWS
    cw = round((TEXT_W - CARD_GUTTER * (CARD_ACROSS - 1)) / CARD_ACROSS, 3)
    inner = round(cw - 0.20, 3)
    for start in range(0, len(cards), per_page):
        page_break(doc)
        p = para(doc, "", 8)
        run(p, "ELEMENT CARDS  ·  CUT ON THE DASHED LINES  ·  LETTERS ARE CODE NAMES, NOT SYMBOLS", 8,
            bold=True, color=pal.accent)
        gap(doc, 4)
        chunk = cards[start:start + per_page]
        for r in range(0, len(chunk), CARD_ACROSS):
            t = doc.add_table(rows=1, cols=CARD_ACROSS * 2 - 1)
            widths = []
            for i in range(CARD_ACROSS):
                widths.append(cw)
                if i < CARD_ACROSS - 1:
                    widths.append(CARD_GUTTER)
            fix_widths(t, widths)
            no_split(t.rows[0], CARD_H)
            for i, card in enumerate(chunk[r:r + CARD_ACROSS]):
                cell = t.rows[0].cells[i * 2]
                borders(cell, pal.label, sz=6, val="dashed")
                cell_margins(cell, top=60, bottom=40, left=100, right=100)
                if card is None:
                    para(cell, "?", 26, bold=True, color=pal.accent, first=True)
                    para(cell, "THE MISSING ELEMENT", 9, bold=True, color=pal.ink)
                    para(cell, "Put this card in the gap your team leaves. Write your predictions "
                               "on page 2.", 9.5, color=pal.ink)
                    continue
                card_face(cell, card, pal, inner)
                tidy_cell(cell)
            if r + CARD_ACROSS < len(chunk):
                gap(doc, 5)         # none after the last row: it would spill onto a blank page


def build_student(spec, pal, code, out):
    doc = Document()
    sec = page_setup(doc)
    running_footer(sec, f"SHULL SCIENCE   {code}   ACTIVITY", pal, with_page_numbers=True)
    header(doc, spec, pal, "Operation Periodic Table", "ACTIVITY  ·  ~45 MIN", code)
    gap(doc, 5)
    fields(doc, pal, [("TEAM MEMBERS", 4.6), ("PERIOD", 1.2), ("TEAM SCORE         / 30", 1.7)])
    gap(doc, 6)
    boxed(doc, pal, "YOUR MISSION", [
        "Your team has 25 cards. Each one is a real element under a letter code name. Arrange them into a "
        "table that shows the patterns you learned this week. When everyone is done, your teacher "
        "reveals the real table. The team whose table is closest to it wins.",
        "One element is missing on purpose. Leave a gap for it and predict its properties. In 1871 "
        "Dmitri Mendeleev did exactly this for an element nobody had found yet.",
    ])
    gap(doc, 6)
    plain(doc, pal, "READING A CARD", [
        ("Ion it forms:  ", "the charge the element takes when it reacts. Your best clue to its family."),
        ("Atomic radius:  ", "size of the atom in picometers (pm). Use it to put a family in order."),
        ("Ionization energy and electronegativity:  ", "check your order. They run opposite to radius."),
        ("Melting point and appearance:  ", "extra evidence. Not every family follows a clean pattern here."),
    ])
    gap(doc, 6)
    plain(doc, pal, "FOUR ROUNDS", [
        ("Sort into families (10 min).  ", "Use the ion each element forms. You should end up with five families "
                                         "and one card that does not fit cleanly anywhere. Decide where it goes "
                                         "and be ready to defend it."),
        ("Order each family (10 min).  ", "Smallest atom at the top. Check the order against ionization energy."),
        ("Line up the rows (10 min).  ", "Put the families side by side so atoms shrink from left to right "
                                         "along each row, the way they do across a real period. Two clues: the "
                                         "top row holds just two cards, one at each end. The missing element "
                                         "is in the middle of its family, not at either end."),
        ("Record and predict (10 min).  ", "Copy your table onto the grid on page 2 and fill in your predictions."),
    ], numbered=True)
    gap(doc, 6)
    plain(doc, pal, "RULES", [
        "No periodic table and no phones until the reveal.",
        "Every card goes on the grid. Every team member has to be able to explain one placement.",
    ])

    page_break(doc)
    p = para(doc, "", 7.5)
    run(p, "YOUR TABLE", 7.5, bold=True, color=pal.accent)
    para(doc, "Write each letter in the cell where your team placed it. Write ? in the gap. "
              "Some cells stay empty.", 9.5)
    gap(doc, 4)
    grid_table(doc, pal, spec)
    gap(doc, 8)
    t = doc.add_table(rows=6, cols=4)
    fix_widths(t, [2.2, 1.9, 1.9, 1.5])
    heads = ("THE MISSING ELEMENT", "YOUR PREDICTION", "ACTUAL (AT THE REVEAL)", "IN RANGE?")
    for c, hd in zip(t.rows[0].cells, heads):
        borders(c, pal.ink, sz=6, edges=("bottom",))
        para(c, hd, 6.5, bold=True, color=pal.label, caps_track=True, first=True)
    for r, prop in enumerate(("Melting point (°C)", "Atomic radius (pm)", "Ionization energy (kJ/mol)",
                              "Electronegativity", "Ion it forms"), 1):
        no_split(t.rows[r], 0.36)
        for c in t.rows[r].cells:
            borders(c, pal.hair, sz=4, edges=("bottom",))
        para(t.rows[r].cells[0], prop, 9.5, color=pal.ink, first=True)
    gap(doc, 8)
    plain(doc, pal, "SCORING", [
        "1 point for each card in exactly the right cell (25).  1 point for a gap in the right cell.  "
        "1 point for each prediction inside your teacher's range (4, ion not scored).",
    ])
    gap(doc, 8)
    answer_lines(doc, pal, [
        "Which card property helped your team most when ordering a family? Why that one?",
        "Find one family where a property did not follow the trend. Name the property and the family.",
        "What is the real element in the gap? How close was your team's ionization energy prediction?",
    ])

    cards_pages(doc, pal, spec)
    doc.save(out)


def build_key(spec, pal, code, out):
    doc = Document()
    sec = page_setup(doc)
    running_footer(sec, f"SHULL SCIENCE   {code}   ACTIVITY KEY", pal, with_page_numbers=True)
    header(doc, spec, pal, "Operation Periodic Table — Teacher Key", "ACTIVITY  ·  ANSWER KEY", code)
    gap(doc, 6)
    m = spec["missing"]
    plain(doc, pal, "BEFORE CLASS", [
        "Print pages 3 and 4 of the student file on cardstock, one set per team (about six), and cut them "
        "once. Keep each set in an envelope; they last all day and next year.",
        "Print pages 1 and 2 once per team. Hide wall periodic tables or plan to turn them.",
        "Germanium is held out on purpose. Its gap is group 14, period 4.",
    ])
    gap(doc, 6)
    plain(doc, pal, "THE REVEAL AND SCORING", [
        "Project the answer grid below (letters only first, then elements). Teams score their own grid "
        "with a different-colored pen.",
        f"Prediction ranges, 1 point each: melting point {m['ranges']['mp'][0]}–{m['ranges']['mp'][1]} °C "
        f"(actual {m['mp']}) · radius {m['ranges']['radius'][0]}–{m['ranges']['radius'][1]} pm (actual "
        f"{m['radius']}) · ionization energy {m['ranges']['ie'][0]}–{m['ranges']['ie'][1]} kJ/mol (actual "
        f"{m['ie']}) · electronegativity {m['ranges']['en'][0]}–{m['ranges']['en'][1]} (actual {m['en']}). "
        "Ion: none, shares electrons (not scored).",
        "Ties: closest ionization energy prediction wins.",
        "Hydrogen: score its point for the group 1, period 1 cell. A team that put it above fluorine and "
        "can say why (it also forms 1−, it is a nonmetal gas) has a good argument; discuss it rather than "
        "argue the point.",
    ])
    gap(doc, 6)

    para(doc, "ANSWER GRID", 7.5, bold=True, color=pal.accent, caps_track=True)
    para(doc, "Columns are groups 1, 2, 14, 17, 18. Rows are periods 1–6.", 9.5, color=pal.ink)
    gap(doc, 3)
    cols, rows = spec["grid"]["columns"], spec["grid"]["rows"]
    fill = {}
    for cd in spec["cards"]:
        fill[(rows.index(cd["period"]), cols.index(cd["group"]))] = (cd["code"], f"{cd['symbol']} · {cd['element']}", pal.accent)
    fill[(rows.index(m["period"]), cols.index(m["group"]))] = ("?", f"{m['symbol']} · {m['element']}", pal.ink)
    grid_table(doc, pal, spec, fill)

    page_break(doc)
    plain(doc, pal, "WHERE THE DATA BREAKS THE TREND — GOOD DISCUSSION, NOT A MISTAKE", [
        ("Group 14 electronegativity  ", "goes 2.55, 1.90, (2.01), 1.96, 2.33. These are the values most "
         "tables print; for Sn and Pb they belong to the 4+ state (2+ values are near 1.8–1.9), so tables "
         "disagree. Teams that order group 14 by electronegativity get Sn and Pb wrong. Radius gives the right order."),
        ("Lead vs. tin ionization energy:  ", "Pb 716 is slightly above Sn 709. Again, radius sorts it."),
        ("Group 2 melting points  ", "jump around (Be 1287, Mg 650, Ca 842, Sr 777, Ba 727). Melting point is "
         "weak evidence for order."),
        ("Tin and lead form 2+ ions,  ", "so some teams put them with group 2. The 4+ option, the higher "
         "electronegativity and the lower melting points separate them."),
        ("Radii are calculated values  ", "(one consistent source), so they will not match the Gizmo's numbers. "
         "The trend is the same."),
        ("Hydrogen  ", "fits no family cleanly: it forms 1+ like group 1, but it is a nonmetal gas, forms 1− "
         "like group 17, and its ionization energy (1312) is far above lithium's. Most tables put it over "
         "lithium for its single valence electron. Expect the best argument of the day here."),
        ("The two round-3 clues  ", "pin the layout. Without them, hydrogen over fluorine, or the gap at the "
         "bottom of group 14, also keeps every row shrinking left to right."),
        ("Letters are code names.  ", "Some match other elements' symbols (card C is not carbon). The card "
         "header says so; repeat it out loud."),
    ])
    gap(doc, 6)
    plain(doc, pal, "AFTER THE REVEAL — MODEL ANSWERS", [
        ("1.  ", "Usually atomic radius, because it changes steadily down every family. Accept ion charge if "
         "the team explains it was for sorting, not ordering."),
        ("2.  ", "Any one of the breaks above, named with its family. Most common: group 14 electronegativity "
         "or group 2 melting point."),
        ("3.  ", f"Germanium. Actual ionization energy {m['ie']} kJ/mol; the average of Si (787) and Sn (709) "
         "is 748, so a team that averaged its neighbors was within 14."),
    ])
    gap(doc, 6)
    plain(doc, pal, "THE STORY TO TELL AT THE REVEAL", [
        "In 1871 Mendeleev left a gap below silicon and predicted an element he called eka-silicon. Clemens "
        "Winkler found it in 1886 and named it germanium. Its properties matched the prediction closely. "
        "Your students just repeated that prediction with better data.",
    ])
    gap(doc, 6)
    para(doc, "LETTERS BY GROUP", 7.5, bold=True, color=pal.accent, caps_track=True)
    t = doc.add_table(rows=1, cols=5)
    fix_widths(t, [1.5] * 5)
    for c, g in zip(t.rows[0].cells, cols):
        borders(c, pal.hair, sz=4)
        cell_margins(c, top=40, bottom=40, left=80, right=60)
        para(c, f"GROUP {g}", 6.5, bold=True, color=pal.label, caps_track=True, first=True)
        for cd in sorted([x for x in spec["cards"] if x["group"] == g], key=lambda x: x["period"]):
            para(c, f"{cd['code']} = {cd['symbol']}", 9.5, color=pal.ink)
        if g == m["group"]:
            para(c, f"? = {m['symbol']} (held out)", 9.5, color=pal.ink)
    gap(doc, 6)
    para(doc, "Sources: " + spec["sources"], 8, color=pal.label)
    doc.save(out)


def main():
    if len(sys.argv) < 3:
        refuse("usage: build_operation_periodic_table.py <spec.json> <outdir>")
    spec = json.load(open(sys.argv[1]))
    valid = known_sections(spec["course"])
    bad = [s for s in spec["sections"] if s not in valid]
    if bad:
        refuse(f"section(s) {bad} not in courses/{spec['course']}/DECISIONS.md. Stop and ask.")
    if len(spec["cards"]) != 25:
        refuse(f"{len(spec['cards'])} cards; the grid and the scoring assume 25")
    for c in spec["cards"]:
        if c["code"] == c["symbol"]:
            refuse(f"card {c['code']} carries its own element's symbol")
    pal = Palette(spec["course"])
    code = f"U{int(spec['unit']):02d} / S{spec['sections'][0]}–S{spec['sections'][-1]}"
    outdir = sys.argv[2]
    for key, fn in ((False, build_student), (True, build_key)):
        name = document_name(spec["docType"], spec["sections"], spec["unit"], key=key,
                             descriptor=spec["descriptor"])
        fn(spec, pal, code, os.path.join(outdir, name))
        print("wrote", os.path.join(outdir, name))


if __name__ == "__main__":
    main()
