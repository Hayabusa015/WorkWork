"""Print parts shared by the activity and quiz builders: the outlined header, the ruled name
fields, and the two kinds of boxed block. Kept here once so the two builders cannot drift.
The header follows build_worksheet_docx.py's (outlined, accent bar on the left, no solid slab).
"""
import os, sys
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from _shull_docx import (   # noqa: E402
    unit_title, borders, para, run, fix_widths, cell_margins, one_cell, shade,
)

TEXT_W = 7.5


def label(cell, text, pal, size=7.5, first=False):
    return para(cell, text, size, bold=True, color=pal.accent, caps_track=True, first=first)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def header(doc, spec, pal, title, kind, code):
    t = doc.add_table(rows=1, cols=2)
    fix_widths(t, [5.55, 1.95])
    h, hr = t.rows[0].cells
    for cc in (h, hr):
        borders(cc, pal.hair, sz=4, edges=("top", "bottom"))
    borders(h, pal.display, sz=30, edges=("left",))
    cell_margins(h, top=80, bottom=80, left=160, right=80)
    cell_margins(hr, top=80, bottom=80, left=80, right=60)
    para(h, f"UNIT {int(spec['unit'])} — {unit_title(spec['course'], spec['unit']).upper()}", 7.5,
         bold=True, color=pal.accent, caps_track=True, first=True)
    para(h, title, 15, bold=True, color=pal.ink)
    p = para(hr, code, 11, bold=True, color=pal.accent, caps_track=True, first=True)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = para(hr, kind, 6.5, color=pal.label, caps_track=True)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def fields(doc, pal, labels_widths):
    t = doc.add_table(rows=1, cols=len(labels_widths))
    fix_widths(t, [w for _, w in labels_widths])
    for cc, (lab, _) in zip(t.rows[0].cells, labels_widths):
        borders(cc, pal.ink, sz=4, edges=("bottom",))
        cell_margins(cc, top=0, bottom=40, left=0, right=120)
        para(cc, lab, 6.5, bold=True, color=pal.label, caps_track=True, first=True)


def boxed(doc, pal, heading, lines, size=9.5):
    c = one_cell(doc)
    borders(c, pal.hair, sz=4)
    borders(c, pal.display, sz=24, edges=("left",))
    shade(c, pal.surface)
    cell_margins(c, top=70, bottom=70, left=130, right=130)
    label(c, heading, pal, first=True)
    for ln in lines:
        para(c, ln, size)
    return c


def plain(doc, pal, heading, lines, size=9.5, numbered=False):
    c = one_cell(doc)
    borders(c, pal.hair, sz=4)
    cell_margins(c, top=60, bottom=60, left=120, right=120)
    label(c, heading, pal, first=True)
    for i, ln in enumerate(lines, 1):
        p = para(c, "", size)
        if numbered:
            # Just the number. No box, no circle. SHULL-CHG-0015.
            run(p, f"{i}.  ", size, bold=True, color=pal.ink)
        if isinstance(ln, tuple):
            run(p, ln[0], size, bold=True, color=pal.ink)
            run(p, ln[1], size, color=pal.ink)
        else:
            run(p, ln, size, color=pal.ink)
    return c
