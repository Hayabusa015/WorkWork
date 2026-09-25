#!/usr/bin/env python3
"""Build a guided/Cornell notes packet as .docx from the same JSON spec. Option A.

Identical structure and content to build_notes.py - the difference is the output.
A .docx can be edited in Word and typed into by a student; it cannot hold the
layout as exactly as the PDF, because Word reflows.

    python3 templates/notes/build_notes_docx.py specs/<spec>.json out.docx

Colour comes from brand/tokens.json. No hex is typed in this file.
"""
import json, os, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Every primitive on this page - how a work box is drawn, how a fraction stacks, how a
# diagram block holds together across a page break - is shared with the worksheet
# builder. It lives in one file so the two cannot drift.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import recall                      # noqa: E402
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from _shull_docx import (          # noqa: E402
    T, G, FONT, FLOOR, COURSE_CODE, Palette, hexof, debullet, known_sections,
    unit_title,
    borders, para, check_item, rule_lines, wrap_to, fix_widths, one_cell, no_split, gap,
    stacked_frac, equation_bar, work_box, given_need, diagram_block,
    page_setup, running_footer, trim_tail,
)

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))


# SHULL-CHG-0016. Matthew's physics packet puts a bordered box under every worked
# example - a 1x1 table, ~1in tall, hRule "atLeast" so it grows but never shrinks, and
# cantSplit so it never breaks across a page. Students work the problem inside it.
# His rule: "anytime problems need solved in guided notes leave a box for them to do it."
WORK_BOX_MIN_IN = 1.4          # his was 0.98in; a little more room for kinematics
GIVEN_LABEL_IN = 0.72

# The Cornell split, decided once. His own packets ran a 1.88in (Geology) / 2.00in
# (Physics) cue column; he asked for it narrower and condensed - the cue is a prompt,
# not a second body column, and every inch it gives back goes to the notes side where
# the students actually write. Everything on the notes side is measured off NOTES_INNER
# so nothing has to be re-derived when this moves again.
TEXT_W_IN = 7.50
CUE_W_IN = 1.28
NOTES_W_IN = TEXT_W_IN - CUE_W_IN          # 6.22
CELL_MAR_IN = 0.56                          # default tcMar, both sides, both nestings
NOTES_INNER_IN = round(NOTES_W_IN - CELL_MAR_IN, 2)   # 5.66


# SHULL-CHG-0017. A fraction is stacked - numerator over denominator with a horizontal
# bar. Never "a/b" inline in the text. Built as a two-row table rather than Office Math
# (OMML): OMML is valid and Word renders it, but LibreOffice will not import it from

def page_break(doc):
    """A 1pt paragraph that starts a new page. Breaking on a pinned paragraph rather
    than on the section's own table keeps the break out of the table, where Word and
    LibreOffice disagree about whether pageBreakBefore in a cell counts."""
    p = gap(doc, 1)
    p._p.get_or_add_pPr().append(OxmlElement("w:pageBreakBefore"))


def keep_next(cell):
    """Hold a cell's paragraphs to whatever follows, so a section head is never left
    alone at the bottom of a page with its first row on the next."""
    for p in cell.paragraphs:
        p.paragraph_format.keep_with_next = True


def zpad(sec):
    """'2.1' -> '02.1'. SHULL-CHG-0024: both halves of a code zero-padded, everywhere."""
    major, minor = str(sec).split(".")
    return f"{int(major):02d}.{minor}"


def main():
    # SHULL-CHG-0025. "Always produce two files: student, blanks empty; key, everything
    # filled, generated from the same source structure so they can't drift." recall.py
    # already read a notes line as {"text", "key"} in its offences() check - the check
    # was written for a shape the renderer never actually implemented. This finishes it:
    # --key switches every blank-bearing line and every problem's answer from hidden to
    # shown, off the same spec, rather than a second hand-authored document.
    args = [a for a in sys.argv[1:] if a != "--key"]
    KEY = "--key" in sys.argv[1:]

    spec = json.load(open(args[0] if args
                          else os.path.join(HERE, "specs", "geo_u01_s01.2-s01.4.json")))
    course = spec["course"]
    pal = Palette(course)
    accent, display, ink, hair = pal.accent, pal.display, pal.ink, pal.hair
    label, footer, white = pal.label, pal.footer, pal.white

    valid = known_sections(course)
    bad = [s for s in spec["sections"] if s not in valid]
    if bad:
        print(f"build_notes_docx: section(s) {', '.join(bad)} are not in "
              f"courses/{course}/DECISIONS.md.", file=sys.stderr)
        return 1

    # Guided notes are strictly recall: the cue column names what to record, it does
    # not ask for it. recall.py carries the rule and both renderers check it.
    if recall.check(spec, "build_notes_docx"):
        return 1

    # "Anytime problems need solved in guided notes leave a box for them to do it."
    # A rule that only lives in a document is not a mechanism, so this refuses to build.
    PROBLEM_WORDS = ("EXAMPLE", "PRACTICE", "PROBLEM", "SOLVE", "CALCULATE", "YOUR TURN")
    missing = []
    for sec in spec["sectionsContent"]:
        for row in sec["rows"]:
            if row.get("problem") or row.get("noWorkBox"):
                continue
            hay = " ".join([row.get("cueLabel", ""), row.get("notesLabel", "")]).upper()
            if any(w in hay for w in PROBLEM_WORDS):
                missing.append(f"{sec['code']} · {row.get('notesLabel') or row.get('cueLabel')}")
    if missing:
        print("build_notes_docx: these rows look like problems to solve and have no work box:",
              file=sys.stderr)
        for m in missing:
            print("   " + m, file=sys.stderr)
        print('\nAdd a "problem" block, or "noWorkBox": true if it genuinely is not one.\n'
              'Students need somewhere to work it. SHULL-CHG-0016.', file=sys.stderr)
        return 1

    # Course profiles. The three courses do not want the same document, and pretending
    # they do is how a Geology packet ends up with a kinematics equation bar.
    eqs = spec.get("equations") or []
    has_problem = any(r.get("problem") for sec in spec["sectionsContent"] for r in sec["rows"])
    if course == "geology":
        if eqs or has_problem:
            print("build_notes_docx: Geology has no math. Remove the equation bar and the "
                  "problem blocks — a Geology packet labels diagrams instead. SHULL-CHG-0017.",
                  file=sys.stderr)
            return 1
    else:
        if has_problem and not eqs:
            print(f"build_notes_docx: {course} packet has problems to solve and no equation "
                  f"bar. Students need the equations at the top of the page so they know what "
                  f"they may reference. Add \"equations\", or \"noEquationBar\": true.",
                  file=sys.stderr)
            if not spec.get("noEquationBar"):
                return 1

    code = COURSE_CODE[course]
    unit = f"U{int(spec['unit']):02d}"
    secs = [zpad(x) for x in spec["sections"]]
    span = f"S{secs[0]}-S{secs[-1]}" if len(secs) > 1 else f"S{secs[0]}"
    # SHULL-CHG-0026. "Flow": each section after the first starts a page and runs
    # continuously - a row may break between prompts, never through one - and the
    # close gets its own page. Opt-in per spec
    # so a packet already in circulation does not reflow under anyone. fit_notes.py
    # then grows the ruled lines into whatever space the breaks free up.
    FLOW = bool(spec.get("flow"))
    suffix = "_Key" if KEY else ""
    out = args[1] if len(args) > 1 else \
        os.path.join(HERE, f"SHULL_{code}_Guided_Notes_{unit}_{span}{suffix}.docx")

    doc = Document()
    s = page_setup(doc)

    # Brand bar. Outlined, not filled: SHULL_DESIGN_SYSTEM section 8 - "no full-page
    # colour banners, no shaded section backgrounds, no solid-fill headers." The first
    # build of this template ignored that and measured 3.2x the ink of Matthew's own
    # packet, which had no cell fills anywhere. A heavy accent rule carries the same
    # hierarchy for a rule's worth of toner.
    c = one_cell(doc); borders(c, display, sz=18, edges=("bottom",))
    para(c, "SHULL SCIENCE  ·  JAMES A. GARFIELD LOCAL SCHOOLS", 7.5,
         bold=True, color=accent, caps_track=True, first=True)
    para(c, unit_title(course, spec["unit"]).upper(), 15, bold=True, color=ink)
    para(c, spec["kicker"], 7.5, color=label, caps_track=True)

    c = one_cell(doc); borders(c, hair)
    para(c, spec["fields"], 9, color=label, first=True)

    if eqs:
        gap(doc, 4)
        equation_bar(doc, eqs, pal,
                     spec.get("equationLabel", "EQUATIONS YOU MAY USE"))

    gap(doc, 4)
    t = doc.add_table(rows=1, cols=2)
    fix_widths(t, [3.75, 3.75])
    for i, (lab, items) in enumerate([("UNIT LEARNING TARGETS", spec["unitTargets"]),
                                      ("KEY TERMS", spec["keyTerms"])]):
        cell = t.rows[0].cells[i]; borders(cell, hair)
        para(cell, lab, 7.5, bold=True, color=accent, caps_track=True, first=True)
        for x in items:
            para(cell, "•  " + debullet(x), 9.5)

    gap(doc, 4)
    c = one_cell(doc); borders(c, display, sz=18, edges=("left",))
    para(c, "HOW THESE NOTES WORK", 7.5, bold=True, color=accent, caps_track=True, first=True)
    for x in spec["howItWorks"]:
        para(c, "•  " + debullet(x), 9.5)

    gap(doc, 2)
    c = one_cell(doc); borders(c, white)
    para(c, "SECTIONS IN THIS UNIT", 7.5, bold=True, color=accent, caps_track=True, first=True)
    for x in spec["sectionList"]:
        check_item(c, debullet(x), pal, 9.5)

    for si, sec in enumerate(spec["sectionsContent"]):
        if FLOW and si > 0:
            page_break(doc)
        else:
            gap(doc, 6)
        t = doc.add_table(rows=1, cols=2)
        fix_widths(t, [5.83, 1.67])
        a, b = t.rows[0].cells
        # Section head: ruled above and below, not filled.
        for cell in (a, b):
            borders(cell, ink, sz=12, edges=("top",))
            borders(cell, display, sz=18, edges=("bottom",))
        para(a, sec["title"], 12.5, bold=True, color=ink, first=True)
        p = para(b, sec["code"], 8.5, bold=True, color=accent, caps_track=True, first=True)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        c = one_cell(doc); borders(c, hair)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        r = p.add_run("LEARNING TARGET   "); r.font.name = FONT; r.font.size = Pt(7.5)
        r.bold = True; r.font.color.rgb = RGBColor.from_string(hexof(accent))
        r2 = p.add_run(sec["learningTarget"]); r2.font.name = FONT; r2.font.size = Pt(9.5)
        if FLOW:
            for cell in (a, b, c):
                keep_next(cell)

        # Under flow every prompt - its label, its ruled lines, and any must-write
        # line under it - is its own table row that cannot split, with the cue column
        # merged down beside them. A page can then break between prompts and never
        # through one, and the section runs on without the gap an unsplittable Cornell
        # row left when it jumped whole: "keep each section together, if a section ends
        # just head to a new page." Row-level, because keep-with-next on paragraphs
        # inside a cell is ignored by LibreOffice and read by Word as "keep this whole
        # row with the next", which chains every row into one unbreakable block.
        rows = sec["rows"]

        def blocks_of(row):
            if not FLOW:
                return [list(range(len(row["notes"])))]
            out = []
            for i, n in enumerate(row["notes"]):
                text = n if isinstance(n, str) else n.get("text", "")
                if out and text.startswith(("*", "✎")):
                    out[-1].append(i)
                else:
                    out.append([i])
            return out or [[]]

        plans = [blocks_of(r) for r in rows]
        t = doc.add_table(rows=sum(len(bl) for bl in plans), cols=2)
        fix_widths(t, [CUE_W_IN, NOTES_W_IN])
        ti = 0
        for ri, row in enumerate(rows):
            blocks = plans[ri]
            for bi, blk in enumerate(blocks):
                tr = t.rows[ti]; ti += 1
                if FLOW:
                    no_split(tr)
                cue, notes = tr.cells
                # No fill on the cue column. Matthew's packet separated the columns with
                # a rule alone, which is the Cornell convention and costs nothing to print.
                edges = ("left", "right") + (("top",) if bi == 0 else ()) \
                    + (("bottom",) if bi == len(blocks) - 1 else ())
                borders(cue, hair, edges=edges); borders(notes, hair, edges=edges)
                if len(blocks) > 1:
                    cue._tc.get_or_add_tcPr().vMerge_val = "restart" if bi == 0 else "continue"
                if bi == 0:
                    # Condensed: the cue label keeps caps and colour but loses its letter
                    # tracking - tracking is what made "DISTANCE VS. DISPLACEMENT" wrap -
                    # and the prompts drop half a point so a cue fits on two lines.
                    para(cue, row["cueLabel"], 7, bold=True, color=accent, first=True)
                    for q in row["cues"]:
                        para(cue, q, 8.5)
                    para(notes, row["notesLabel"], 7.5, bold=True, color=accent,
                         caps_track=True, first=True)
                    if row.get("diagram"):
                        diagram_block(notes, row["diagram"], pal, NOTES_INNER_IN, HERE)

                    prob = row.get("problem")
                    if prob:
                        para(notes, prob.get("label", "EXAMPLE"), 7.5, bold=True,
                             color=accent, caps_track=True)
                        if prob.get("statement"):
                            para(notes, prob["statement"], 9.5)
                        if prob.get("given") or prob.get("need"):
                            given_need(notes, prob.get("given", ""), prob.get("need", ""),
                                       pal, NOTES_INNER_IN)
                        work_box(notes, prob.get("workLabel", "WORK — SHOW EVERY STEP"), pal,
                                 float(prob.get("workHeightIn", WORK_BOX_MIN_IN)),
                                 NOTES_INNER_IN)
                        # SHULL-CHG-0025: an answer is what the key is for. Printing it
                        # unconditionally would leak it onto the student page.
                        if KEY and prob.get("answer"):
                            para(notes, prob["answer"], 9.5, bold=True, color=display)

                for k, idx in enumerate(blk):
                    n = row["notes"][idx]
                    # A notes line is a plain string, or {"text", "key"} - recall.py's
                    # offences() already reads it this shape.
                    text, keytext = (n, None) if isinstance(n, str) \
                        else (n.get("text", ""), n.get("key"))
                    mw = text.startswith(("*", "✎"))
                    body = text.lstrip("*✎").strip()
                    p = para(notes, body, 9.5, bold=mw, first=(bi > 0 and k == 0))
                    if mw:
                        pPr = p._p.get_or_add_pPr()
                        bd = OxmlElement("w:pBdr"); x = OxmlElement("w:left")
                        x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "18")
                        x.set(qn("w:space"), "6"); x.set(qn("w:color"), hexof(display))
                        bd.append(x); pPr.append(bd)
                        continue
                    lines = (n.get("lines") if isinstance(n, dict) else None) \
                        or recall.ruled_lines(body)
                    if not KEY:
                        rule_lines(notes, lines, hair)
                    elif keytext is not None:
                        # The answer goes on the student's own ruled lines, so the key
                        # is the student copy filled in - same pages, same breaks.
                        rule_lines(notes, lines, hair, color=display,
                                   written=wrap_to(keytext, NOTES_INNER_IN - 0.05, 9.5, True))
                    elif "___" in body or body.endswith(":"):
                        print(f"build_notes_docx: --key requested but {sec['code']} has a "
                              f"blank line with no \"key\": {body!r}", file=sys.stderr)
                        return 1


        gap(doc, 2)
        c = one_cell(doc); borders(c, accent)
        if FLOW:
            c._tc.getparent().get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        para(c, "SECTION SUMMARY — close your notes before you write this", 7.5,
             bold=True, color=accent, caps_track=True, first=True)
        para(c, sec["summaryPrompt"], 9.5)
        rule_lines(c, sec.get("summaryLines", 4), hair)
        para(c, "SELF-CHECK", 7.5, bold=True, color=accent, caps_track=True)
        for x in sec.get("selfCheck", []):
            check_item(c, x, pal)

    close = spec["close"]
    if FLOW:
        page_break(doc)
    else:
        gap(doc, 6)
    c = one_cell(doc); borders(c, ink, sz=12, edges=("top",))
    borders(c, display, sz=18, edges=("bottom",))
    para(c, close["banner"], 9, bold=True, color=ink, caps_track=True, first=True)
    c = one_cell(doc); borders(c, hair)
    para(c, "SECTION CHECKLIST", 7.5, bold=True, color=accent, caps_track=True, first=True)
    for x in close["checklist"]:
        check_item(c, x, pal)
    para(c, "UNIT BIG PICTURE", 7.5, bold=True, color=accent, caps_track=True)
    para(c, close["bigPicture"], 9.5)
    rule_lines(c, close.get("bigPictureLines", 3), hair)
    para(c, close["fuzzyLabel"], 8.5, bold=True, color=accent, caps_track=True)
    rule_lines(c, close.get("fuzzyLines", 3), hair)

    footer_txt = f"SHULL SCIENCE          {unit} · {span}"
    if KEY:
        footer_txt += "  ·  KEY"
    running_footer(s, footer_txt, pal)

    trim_tail(doc)
    doc.save(out)
    print(f"wrote {out}  —  {code} {unit} {span}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
