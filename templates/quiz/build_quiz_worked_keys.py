#!/usr/bin/env python3
"""Worked answer keys for the Physics S2.2 quiz: every problem solved step by step.

    python3 templates/quiz/build_quiz_worked_keys.py OUTDIR

Writes `2.2 Quiz A Key.docx`, `B`, `C` (and nothing else). Each key is two pages: page 1 the
three direction questions and Q4, page 2 Q5 and the grading notes. The solutions are drawn in the
key-handwriting face (standards/ANSWER_KEYS.md, SHULL-CHG-0043/0044): pencil for the math, pen for
written lines, a box round each final answer, arrow ink for notes, one struck-through slip where a
student genuinely makes it. Every division is a stacked fraction. The page header, the point
rubric and the grading notes are typed in the brand face.

Every number comes from vectors_22.py (solved in code); the steps below only format it. No sketch is
drawn: ANSWER_KEYS section 6 omits a sketch the scene builder cannot model, and it models only a
line of motion, not a two-dimensional sum.
"""
import math, os, sys, tempfile
from docx.shared import Pt, Inches
from docx.enum.text import WD_BREAK

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(REPO, "scripts"))
sys.path.insert(0, os.path.dirname(HERE))
import vectors_22 as V                                              # noqa: E402
import build_quiz_docx as Q                                         # noqa: E402
from key_solution import solution_image                             # noqa: E402
from _shull_name import document_name                               # noqa: E402
from _shull_docx import (Palette, borders, para, fix_widths, gap, one_cell,  # noqa: E402
                         trim_tail, schema_order, tight_cells)

MINUS = "−"
CONVENTION = ("East is +x and north is +y. Angles run counterclockwise from east. "
              "Keep the extra digits until the last step, then round to two significant figures.")


def n(x, nd=2):
    return f"{x:.{nd}f}".replace("-", MINUS)


def g(x, unit=""):
    """A given as the problem wrote it: 4.0 m/s stays 4.0, 28 m stays 28."""
    return f"{abs(x):.1f}" if unit == "m/s" else f"{abs(x):g}"


def quadrant(t):
    t %= 360
    return ("I", 0, 90) if t < 90 else ("II", 90, 180) if t < 180 else ("III", 180, 270) if t < 270 else ("IV", 270, 360)


def ref_formula(t):
    """The angle from the nearer horizontal axis, as arithmetic, and its value."""
    t %= 360
    if t < 90:   return f"{t:.0f}° is already measured from east", t
    if t < 180:  return f"180° − {t:.0f}° = {180 - t:.0f}°", 180 - t
    if t < 270:  return f"{t:.0f}° − 180° = {t - 180:.0f}°", t - 180
    return f"360° − {t:.0f}° = {360 - t:.0f}°", 360 - t


def say(s, ink="pen"):
    return {"say": s, "ink": ink}


def row(*items, note=None, ink="pencil"):
    d = {"row": list(items), "ink": ink}
    if note:
        d["note"] = note
    return d


def ans(label, value, note=None):
    d = {"answer": [label], "value": value}
    if note:
        d["note"] = note
    return d


# ---- Q1-3: restate a direction ------------------------------------------------------------

def restate_steps(it):
    t = round(it["theta"]) % 360
    e = it["eq"]
    q, lo, hi = quadrant(t)
    refline, a = ref_formula(t)
    steps = []
    key = it["given_key"]
    if key == "neg":
        steps += [row(f"{it['given']} + 360° = {t}°", note="negative = clockwise"), ans("0–360:", e["deg"])]
    elif key == "c1":
        formula, _ = V.standard_formula(*V.VERSIONS_Q2[it["letter"]])
        steps += [say("Start from the axis after “of”, then turn."),
                  row(formula), ans("0–360:", e["deg"])]
    steps += [say(f"{t}° is in quadrant {q}, between {lo}° and {hi}°.")]
    if key == "deg":
        steps += [row(refline), ans("compass:", e["c1"])]
    elif key == "neg":
        steps += [row(refline), ans("compass:", e["c1"])]
    else:
        steps += [row(f"90° − {a:.0f}° = {90 - a:.0f}°", note="from the other axis"), ans("other axis:", e["c2"])]
    if key == "deg":
        steps += [row(f"90° − {a:.0f}° = {90 - a:.0f}°", note="from the other axis"), ans("other axis:", e["c2"]),
                  row(f"{t}° − 360° = {n(t - 360, 0)}°", note="clockwise: subtract 360"), ans("negative:", e["neg"])]
    elif key == "neg":
        steps += [row(f"90° − {a:.0f}° = {90 - a:.0f}°", note="from the other axis"), ans("other axis:", e["c2"])]
    else:
        steps += [row(f"{t}° − 360° = {n(t - 360, 0)}°", note="clockwise: subtract 360"), ans("negative:", e["neg"])]
    steps += [say("Any two forms earn full credit.", "arrow")]
    return steps


# ---- Q4: perpendicular vectors ----------------------------------------------------------

def q4_steps(v4, s4):
    x, y, unit = v4["x"], v4["y"], v4["unit"]
    xs, ys = g(x, unit), g(y, unit)
    R2 = x * x + y * y
    ref = math.degrees(math.atan(abs(y) / abs(x)))
    ns, ew = ("north" if y > 0 else "south"), ("east" if x > 0 else "west")
    xw = "east" if x > 0 else "west"
    yw = "north" if y > 0 else "south"
    steps = [
        row(f"Rx = {'+' if x > 0 else MINUS}{xs} {unit}", "     ", f"Ry = {'+' if y > 0 else MINUS}{ys} {unit}", note=f"{xw} is {'+' if x > 0 else MINUS}x, {yw} is {'+' if y > 0 else MINUS}y"),
        say("The legs are perpendicular, so the resultant is the hypotenuse."),
        {"mistake": [f"R = {xs} + {ys} = {abs(x) + abs(y):.1f}" if unit == "m/s" else f"R = {xs} + {ys} = {abs(x) + abs(y):g}"], "note": "perpendicular legs need Pythagoras, not addition"},
        row(f"R^{{2}} = Rx^{{2}} + Ry^{{2}} = ({xs})^{{2}} + ({ys})^{{2}} = {R2:.2f}"),
        row("R =", {"sqrt": f"{R2:.2f}"}, f"= {s4['R']:.2f} {unit}"),
        ans("R =", f"{V.two_sf(s4['R'])} {unit}", note="round last"),
        row("tan(angle) =", {"frac": [ys, xs]}, f"= {abs(y) / abs(x):.2f}", note="from the equations box"),
        row(f"angle = tan^{{−1}}({abs(y) / abs(x):.2f}) = {ref:.1f}°"),
        say(f"Rx is {'positive' if x > 0 else 'negative'} ({xw}) and Ry is {'positive' if y > 0 else 'negative'} ({yw}), "
            f"so the vector points {yw} of {xw}."),
        ans("direction:", f"{ref:.0f}° {ns} of {ew}"),
    ]
    if v4.get("frame") == "v":
        steps += [row(f"90° − {ref:.0f}° = {90 - ref:.0f}°", note="the same direction from north"), ans("or:", s4["eq"]["c2"])]
    steps += [say(f"Also correct: {s4['eq']['deg']}, {s4['eq']['neg']}, and {s4['eq']['c2'] if v4.get('frame') != 'v' else s4['eq']['c1']}.", "arrow")]
    return steps


# ---- Q5: two vectors at angles ----------------------------------------------------------

def q5_steps(v5, s5):
    u = v5["unit"]
    vs = s5["vs"]                      # [(mag, theta360), ...]
    parts = s5["parts"]
    names = ("A", "B")
    steps = [say("Turn each direction into a 0–360 angle, then use the equations box.")]
    for (m, spec), name in zip(v5["v"], names):
        steps.append(row(f"{name}: {V.standard_formula(*spec)[0]}", ink="pencil"))
    for (m, th), (cx, cy), name in zip(vs, parts, names):
        steps.append(row(f"{name}x = {m:g} cos {th:.0f}° = {n(cx)} {u}", "     ", f"{name}y = {m:g} sin {th:.0f}° = {n(cy)} {u}"))
    # the slip: the first component that is negative, summed with its sign dropped
    slip = next(((nm, c, v) for (cx, cy), nm in zip(parts, names)
                 for c, v in (("x", cx), ("y", cy)) if v < 0), None)
    if slip:
        nm, comp, val = slip
        i = 0 if comp == "x" else 1
        other = parts[1][i] if nm == "A" else parts[0][i]
        point = {"x": "west", "y": "south"}[comp]
        r = "Rx" if comp == "x" else "Ry"
        steps.append({"mistake": [f"{r} = {abs(val):.2f} + {abs(other):.2f} = {abs(val) + abs(other):.2f}"],
                      "note": f"{nm}{comp} points {point}, so it is negative"})
    steps += [
        say("Add the x parts and the y parts. Keep the extra digits."),
        row(f"Rx = Ax + Bx = {n(parts[0][0])} + ({n(parts[1][0])}) = {n(s5['rx'])} {u}") if parts[1][0] < 0 else
        row(f"Rx = Ax + Bx = {n(parts[0][0])} + {n(parts[1][0])} = {n(s5['rx'])} {u}"),
        row(f"Ry = Ay + By = {n(parts[0][1])} + ({n(parts[1][1])}) = {n(s5['ry'])} {u}") if parts[1][1] < 0 else
        row(f"Ry = Ay + By = {n(parts[0][1])} + {n(parts[1][1])} = {n(s5['ry'])} {u}"),
        row(f"R^{{2}} = ({n(s5['rx'])})^{{2}} + ({n(s5['ry'])})^{{2}} = {s5['rx'] ** 2 + s5['ry'] ** 2:.2f}"),
        row("R =", {"sqrt": f"{s5['rx'] ** 2 + s5['ry'] ** 2:.2f}"}, f"= {s5['R']:.2f} {u}"),
        ans("R =", f"{V.two_sf(s5['R'])} {u}", note="round last"),
    ]
    ref = math.degrees(math.atan(abs(s5["ry"]) / abs(s5["rx"])))
    yw = "north" if s5["ry"] > 0 else "south"
    xw = "east" if s5["rx"] > 0 else "west"
    steps += [
        row("tan(angle) =", {"frac": [f"{abs(s5['ry']):.2f}", f"{abs(s5['rx']):.2f}"]}, f"= {abs(s5['ry']) / abs(s5['rx']):.2f}"),
        row(f"angle = tan^{{−1}}({abs(s5['ry']) / abs(s5['rx']):.2f}) = {ref:.1f}°", note="a size only; the signs give the quadrant"),
        say(f"Rx is {'positive' if s5['rx'] > 0 else 'negative'} and Ry is {'positive' if s5['ry'] > 0 else 'negative'}, so the result points {yw} of {xw}."),
        ans("direction:", f"{ref:.0f}° {yw} of {xw}"),
        say(f"Also correct: {s5['eq']['deg']}, {s5['eq']['neg']}, and {s5['eq']['c2']}.", "arrow"),
    ]
    return steps


# ---- documents ---------------------------------------------------------------------------

def picture(cell, png, width_in, first=False):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(png, width=Inches(width_in))


def worked_key(L, d, out, tmp):
    pal = Palette("physics")
    doc = Q.new_doc(pal, f"SHULL SCIENCE / PHYSICS / UNIT 02 / QUIZ {L} / TEACHER KEY")
    Q.header(doc, pal, f"{V.TITLE_LINES[1]}, VERSION {L}", key=True)
    gap(doc, 4)
    c = one_cell(doc); borders(c, pal.display, sz=18, edges=("left",))
    para(c, CONVENTION, 8.5, color=pal.ink, first=True)
    gap(doc, 5)
    t = doc.add_table(rows=1, cols=3)
    fix_widths(t, [2.5, 2.5, 2.5])
    tight_cells(t, top=30, bottom=30, left=60, right=60)
    for cell, it in zip(t.rows[0].cells, d["restate"]):
        borders(cell, pal.hair, sz=6)
        para(cell, f"{it['n']}.   2 pts", 8, bold=True, color=pal.accent, caps_track=True, first=True)
        para(cell, f"Restate {it['given']}" + ("." if it["n"] != 2 else "."), 9.5, color=pal.ink)
        steps = restate_steps({**it, "letter": L})
        png, _ = solution_image({"prompt": f"{L}{it['n']}", "steps": steps}, os.path.join(tmp, f"{L}q{it['n']}.png"), width_in=2.35, size=19)
        picture(cell, png, 2.35)
    q4, q5 = d["q4"], d["q5"]
    v4, v5 = d["v"]["q4"], d["v"]["q5"]
    for qn, text_, steps, pts in ((4, v4["text"], q4_steps(v4, q4), "1 setup for R · 1 R with units · 1 angle setup · 1 angle value · 1 direction named (axis or 0–360)."),
                                  (5, v5["text"], q5_steps(v5, q5), "1 components of A · 1 components of B · 1 Rx and Ry sums · 1 R with units · 1 direction with a named reference.")):
        if qn == 5:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            gap(doc, 6)
        c = one_cell(doc); borders(c, pal.hair)
        para(c, f"{qn}.   5 pts", 8, bold=True, color=pal.accent, caps_track=True, first=True)
        para(c, text_, 9.5, color=pal.ink)
        png, _ = solution_image({"prompt": f"{L}{qn}", "steps": steps}, os.path.join(tmp, f"{L}q{qn}.png"), width_in=7.0, size=27)
        picture(c, png, 7.0)
        para(c, "Points: " + pts, 7.5, color=pal.label)
    gap(doc, 6)
    c = one_cell(doc); borders(c, pal.hair)
    para(c, "GRADING NOTES", 7.5, bold=True, color=pal.accent, caps_track=True, first=True)
    para(c, "Common errors: dropping the sign of a component · giving R with no direction · naming the wrong quadrant "
            "(read it from the signs of Rx and Ry, never from the calculator) · rounding the components too early.", 8.5, color=pal.ink)
    para(c, "Judgment call for you: a student who rounds the components to two significant figures can land one unit off "
            "in R and a degree off in the angle. Decide whether to accept that; the key above keeps the extra digits.", 8.5, color=pal.ink)
    para(c, "A direction is correct if it names its reference axis or uses the 0–360 or negative form. Any correct form earns the point.", 8.5, color=pal.ink)
    Q.finish(doc, out)


def main():
    outdir = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "out")
    os.makedirs(outdir, exist_ok=True)
    data = V.build()
    V.VERSIONS_Q2 = {L: v["q2"] for L, v in V.VERSIONS.items()}
    bad = 0
    with tempfile.TemporaryDirectory() as tmp:
        for L in "ABC":
            p = os.path.join(outdir, document_name("Quiz", ["2.2"], 2, version=L, key=True))
            worked_key(L, data[L], p, tmp)
            pages = Q.pages(p)
            print(f"{'ok ' if pages <= 2 else 'BAD'} {pages} page  {os.path.basename(p)}")
            bad += pages > 2
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
