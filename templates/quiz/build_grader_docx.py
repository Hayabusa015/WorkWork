#!/usr/bin/env python3
"""Build the one-page quick grader for a quiz, both versions side by side.

    python3 templates/quiz/build_grader_docx.py <spec_a.json> <spec_b.json> <out.docx>

A teacher lays this beside a stack of papers. Every answer, every acceptable form of a direction and
every point split is read from the two quiz specs, so the grader cannot disagree with the keys:
the same fact is not typed twice.

What it reads, per question:
    angle questions (answerLines)   the given direction from the prompt, and the "same direction:"
                                    row for every form that earns credit
    calculation questions           finalAnswer ("53 m at 122°"), the "same direction:" row, the
                                    "Grading (n): ..." label, and each "mistake" step
"""
import json, os, re, sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from _shull_docx import (          # noqa: E402
    COURSE_CODE, Palette, unit_title, borders, para, run, fix_widths, no_split, page_setup,
    running_footer, cell_margins, gap, trim_tail, shade,
)


def clean(s):
    return re.sub(r"_\{(\w+)\}", r"\1", s).replace("θ", "direction")


def forms(steps):
    row = next(s["row"][0] for s in steps if "row" in s and str(s["row"][0]).startswith("same direction:"))
    return [f.strip() for f in row.split(":", 1)[1].split("=")]


def digest(spec):
    out = []
    for q in spec["sectionsContent"][0]["questions"]:
        steps = q["steps"]
        d = {"points": q["points"], "forms": forms(steps)}
        grading = next((s["label"] for s in steps if "label" in s and s["label"].startswith("Grading")), "")
        d["grading"] = re.sub(r"^Grading \(\d+\):\s*", "", grading)
        # A and B name the vector by its own letter ("R setup", "v setup"); the split is the same
        d["grading"] = re.sub(r"(?:(?<=^)|(?<=· ))[A-Za-z] (setup|and unit)", r"magnitude \1", d["grading"])
        if q.get("answerLines"):
            d["kind"] = "angle"
            d["given"] = re.search(r"Restate (.+?) in two other ways", q["prompt"]).group(1)
        else:
            d["kind"] = "calc"
            mag, ang = q["finalAnswer"].split(" at ")
            d["mag"], d["ang"] = mag, ang
            d["slips"] = [(clean(" ".join(x for x in s["mistake"] if isinstance(x, str))), clean(s.get("note", "")))
                          for s in steps if "mistake" in s]
        out.append(d)
    return out


def cell_text(cell, lines, pal):
    """lines: (text, size, bold, color) tuples."""
    first = True
    for text, size, bold, color in lines:
        p = para(cell, text, size, bold=bold, color=color, first=first)
        p.paragraph_format.space_after = Pt(1.5)
        first = False


def answer_cell(cell, d, pal):
    cell_margins(cell, top=60, bottom=60, left=90, right=80)
    if d["kind"] == "angle":
        lines = [(f"Given  {d['given']}", 9.5, True, pal.ink),
                 ("Any two of:", 8, False, pal.label)]
        lines += [(f, 9.5, False, pal.ink) for f in d["forms"]]
    else:
        lines = [(f"{d['mag']}   at   {d['ang']}", 10.5, True, pal.ink),
                 ("Direction also:", 8, False, pal.label)]
        lines += [(f, 9.5, False, pal.ink) for f in d["forms"]]
    cell_text(cell, lines, pal)


def main():
    if len(sys.argv) < 4:
        print("usage: build_grader_docx.py <spec_a.json> <spec_b.json> <out.docx>", file=sys.stderr)
        return 1
    sa, sb = (json.load(open(p, encoding="utf-8")) for p in sys.argv[1:3])
    a, b = digest(sa), digest(sb)
    assert [x["points"] for x in a] == [x["points"] for x in b] and len(a) == len(b), "versions must match in points"
    pal = Palette(sa["course"])
    unit = f"U{int(sa['unit']):02d}"
    code = sa["sectionsContent"][0]["code"]
    total = sum(x["points"] for x in a)

    doc = Document()
    s = page_setup(doc)

    t = doc.add_table(rows=1, cols=2)
    fix_widths(t, [5.55, 1.95])
    h, hr = t.rows[0].cells
    for cc in (h, hr):
        borders(cc, pal.hair, sz=4, edges=("top", "bottom"))
    borders(h, pal.display, sz=30, edges=("left",))
    cell_margins(h, top=80, bottom=80, left=160, right=80)
    cell_margins(hr, top=80, bottom=80, left=80, right=60)
    para(h, f"UNIT {int(sa['unit'])} — {unit_title(sa['course'], sa['unit']).upper()} · TEACHER KEY", 7.5,
         bold=True, color=pal.accent, caps_track=True, first=True)
    para(h, "Vectors Quiz: Quick Grader", 15, bold=True, color=pal.ink)
    pp = para(hr, f"{unit} / S{code}", 11, bold=True, color=pal.accent, caps_track=True, first=True)
    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pp = para(hr, f"{total} POINTS  ·  A AND B", 6.5, color=pal.label, caps_track=True)
    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    gap(doc, 8)

    widths = [0.35, 0.40, 2.45, 2.45, 1.85]
    g = doc.add_table(rows=1 + len(a), cols=5)
    fix_widths(g, widths)
    for cc, lab in zip(g.rows[0].cells, ("Q", "PTS", "VERSION A", "VERSION B", "POINTS GO TO")):
        borders(cc, pal.ink, sz=6, edges=("bottom",))
        cell_margins(cc, top=0, bottom=40, left=90 if lab not in ("Q",) else 40, right=60)
        para(cc, lab, 7, bold=True, color=pal.label, caps_track=True, first=True)
    for i, (qa, qb) in enumerate(zip(a, b), start=1):
        row = g.rows[i]
        no_split(row)
        for cc in row.cells:
            borders(cc, pal.hair, sz=4, edges=("bottom",))
        q, pts, ca, cb, cg = row.cells
        cell_margins(q, top=60, bottom=60, left=40, right=40)
        cell_margins(pts, top=60, bottom=60, left=90, right=40)
        para(q, str(i), 12, bold=True, color=pal.ink, first=True)
        para(pts, str(qa["points"]), 11, bold=True, color=pal.ink, first=True)
        answer_cell(ca, qa, pal)
        answer_cell(cb, qb, pal)
        cell_margins(cg, top=60, bottom=60, left=90, right=40)
        txt = qa["grading"] if qa["grading"] == qb["grading"] else qa["grading"] + "  /  " + qb["grading"]
        parts = [x.strip() for x in re.split(r"\s*[·•]\s*", txt)]
        cell_text(cg, [(x, 9, False, pal.ink) for x in parts], pal)

    # slips that students really make, taken from the "mistake" steps of the keys
    slips = []
    for ver, ds in (("A", a), ("B", b)):
        for i, d in enumerate(ds, start=1):
            for what, why in d.get("slips", []):
                slips.append((f"{ver}{i}", what, why))
    if slips:
        gap(doc, 10)
        c = doc.add_table(rows=1, cols=1)
        fix_widths(c, [7.5])
        cc = c.rows[0].cells[0]
        borders(cc, pal.hair, sz=4, edges=("top",))
        cell_margins(cc, top=60, bottom=0, left=0, right=0)
        para(cc, "Slips to watch for", 9.5, bold=True, color=pal.ink, first=True)
        for tag, what, why in slips:
            p = para(cc, "", 9, first=False)
            run(p, f"{tag}   ", 9, bold=True, color=pal.ink)
            run(p, f"{what}", 9, bold=True, color=pal.ink)
            run(p, f"   {why}", 9, color=pal.label)
    gap(doc, 8)
    p = doc.add_paragraph()
    run(p, "Calculation answers are rounded to two significant figures. The point splits are the ones printed on each key.", 8,
        color=pal.label)
    running_footer(s, f"SHULL SCIENCE          {unit} · S{code} · QUIZ KEY · QUICK GRADER", pal, with_page_numbers=True)
    trim_tail(doc)
    doc.save(sys.argv[3])
    print(f"wrote {sys.argv[3]}  —  {len(a)} questions, {total} points, {len(slips)} slips")
    return 0


if __name__ == "__main__":
    sys.exit(main())
