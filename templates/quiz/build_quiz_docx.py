#!/usr/bin/env python3
"""Build a SHULL quiz (the student copy) from a spec.

    python3 templates/quiz/build_quiz_docx.py <spec.json> <out.docx>

A quiz is not a practice set. A practice set sends students to their notebooks, so the practice-set
builder carries no work space and a card per question. A quiz is written on, so the cards here
are the work space, and short questions share a row so the page is not mostly margin.

Same masthead, palette, fonts and footer as every other SHULL print document. The answer key is
built by build_worksheet_key_docx.py from the same spec, so the two cannot drift.

Spec, one section:
    code, title, directions         plain text; no label, no box
    equations                       short strings; "_x" or "_{ix}" is a subscript, "^2" a superscript,
                                    "θ" and "Δ" are drawn (the brand face has no Greek)
    questions[]                     points, prompt, and either
        inRow: true                 shares one row with the other inRow questions; answerLines
                                    ruled lines under the prompt
        cardHeightIn                the card is this tall, and the blank space is the work area

Theta. The brand face (Archivo) has no Greek, and standards/QA_GATE.md fails a PDF that falls back
to a second font. So theta is set as a small inline picture cut from a Greek-capable face and sized
to the equation text. A picture is not a font, so the page is still Archivo only.
"""
import json, os, re, sys, tempfile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from _shull_docx import (          # noqa: E402
    COURSE_CODE, Palette, known_sections, unit_title, borders, para, run, fix_widths,
    no_split, page_setup, running_footer, cell_margins, gap, trim_tail, tight_cells,
    unpad_cell, rule_lines, text_width_in, hexof,
)

THETA_FACE = "/usr/share/fonts/opentype/inter/Inter-Bold.otf"
THETA_FACE_REGULAR = "/usr/share/fonts/opentype/inter/Inter-Regular.otf"
THETA_CAP = 0.74            # theta stands about this tall as a fraction of the type size
_TMP = tempfile.mkdtemp(prefix="shull_quiz_")


def refuse(msg):
    print(f"build_quiz_docx: {msg}", file=sys.stderr)
    return 1


GLYPHS = "θΔ"              # Greek the brand face cannot set; each is drawn as a small picture


def glyph_png(ch, bold=True):
    path = os.path.join(_TMP, f"glyph_{ord(ch):x}_{'b' if bold else 'r'}.png")
    if not os.path.exists(path):
        from PIL import Image, ImageDraw, ImageFont
        f = ImageFont.truetype(THETA_FACE if bold else THETA_FACE_REGULAR, 900)
        im = Image.new("L", (1600, 1600), 0)
        ImageDraw.Draw(im).text((200, 200), ch, font=f, fill=255)
        im = im.crop(im.getbbox())
        out = Image.new("RGBA", im.size, (0, 0, 0, 0))
        out.putalpha(im)
        out.save(path)
    return path


def marked(paragraph, text, size, *, bold=True, color=None, italic_vars=False):
    """Write `text` into a paragraph, turning _x, ^2 and θ into subscripts, superscripts and a
    drawn theta."""
    buf = ""
    i = 0

    def flush():
        nonlocal buf
        if not buf:
            return
        if not italic_vars:
            run(paragraph, buf, size, bold=bold, color=color)
        else:
            # variables slant; cos, sin, tan, units, digits and operators stay upright
            for piece in re.split(r"(cos|sin|tan|m/s)", buf):
                for sub in re.findall(r"[A-Za-z]+|[^A-Za-z]+", piece):
                    slant = sub.isalpha() and sub not in ("cos", "sin", "tan", "m", "s")
                    run(paragraph, sub, size, bold=bold, color=color, italic=slant)
        buf = ""

    while i < len(text):
        ch = text[i]
        if ch in GLYPHS:
            flush()
            png = glyph_png(ch, bold)
            h = size * THETA_CAP
            pic = paragraph.add_run().add_picture(png, height=Pt(h))
            for att in ("distT", "distB", "distL", "distR"):   # LibreOffice pads an inline picture
                pic._inline.set(att, "0")                       # unless the margins are stated
        elif ch in "_^" and i + 1 < len(text):
            flush()
            if text[i + 1] == "{" and "}" in text[i + 2:]:      # _{ix}: a multi-character sub/superscript
                end = text.index("}", i + 2)
                body = text[i + 2:end]
                i = end
            else:
                body = text[i + 1]
                i += 1
            r = run(paragraph, body, size, bold=bold, color=color)
            if ch == "_":
                r.font.subscript = True
            else:
                r.font.superscript = True
        else:
            buf += ch
        i += 1
    flush()


def _plain(t):
    return re.sub(r"[_^{}]", "", t).replace("Δ", "M").replace("θ", "O")


def _frac_cell(cell, num, den, size, pal, bold=True, italic_vars=False):
    cell_margins(cell, top=0, bottom=0, left=15, right=15)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "8"); bt.set(qn("w:space"), "1")
    bt.set(qn("w:color"), hexof(pal.ink)); bd.append(bt); pPr.append(bd)
    marked(p, num, size, bold=bold, color=pal.ink, italic_vars=italic_vars)
    q = cell.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after = Pt(0)
    marked(q, den, size, bold=bold, color=pal.ink, italic_vars=italic_vars)


def equation_cell(cell, eq, pal, size=11, bold=True, italic_vars=False):
    """One equation. "[n|d]" is a stacked fraction: numerator over a horizontal bar over denominator."""
    parts = [x for x in re.split(r"(\[[^\]]+\|[^\]]+\])", eq) if x != ""]
    if not any(x.startswith("[") for x in parts):
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        marked(p, eq, size, bold=bold, color=pal.ink, italic_vars=italic_vars)
        return
    t = cell.add_table(rows=1, cols=len(parts))
    widths = []
    for x in parts:
        if x.startswith("["):
            n, d = x[1:-1].split("|")
            widths.append(max(text_width_in(_plain(n), size, bold), text_width_in(_plain(d), size, bold)) + 0.10)
        else:
            widths.append(text_width_in(_plain(x), size, bold) + 0.06)
    fix_widths(t, [round(w, 3) for w in widths])
    tight_cells(t, left=0, right=0)
    for x, c in zip(parts, t.rows[0].cells):
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if x.startswith("["):
            n, d = x[1:-1].split("|")
            _frac_cell(c, n, d, size, pal, bold, italic_vars)
        else:
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            marked(p, x, size, bold=bold, color=pal.ink, italic_vars=italic_vars)
    unpad_cell(cell)


def equation_block(doc, eqs, pal, label="Equations", width=7.5, per_row=3, note=None, size=11, bold=True,
                   italic_vars=False, row_in=None):
    gap(doc, 5)
    t = doc.add_table(rows=1, cols=1)
    fix_widths(t, [width])
    c = t.rows[0].cells[0]
    borders(c, pal.hair, sz=4, edges=("top", "bottom"))
    cell_margins(c, top=60, bottom=50, left=0, right=0)
    p = para(c, label, 9.5, bold=True, color=pal.ink, first=True)
    p.paragraph_format.space_after = Pt(3)
    if note:
        run(p, "   " + note, 9, color=pal.label)
    rows = (len(eqs) + per_row - 1) // per_row
    g = c.add_table(rows=rows, cols=per_row)
    colw = round(width / per_row - 0.02, 2)
    fix_widths(g, [colw] * per_row)
    tight_cells(g, top=30, bottom=40, left=0, right=60)
    for i, eq in enumerate(eqs):
        cell = g.rows[i // per_row].cells[i % per_row]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        equation_cell(cell, eq, pal, size, bold, italic_vars)
    if row_in:
        for r in g.rows:
            no_split(r, row_in)
    unpad_cell(c)


def equation_groups(doc, groups, pal, width=7.5, size=9.5, bold=False, italic_vars=True, row_in=0.26,
                    split=(4.35, 3.15)):
    """Equations in labelled groups, two columns of groups.

    Each group is {"label", "note", "eqs", "columns", "col": "L"|"R"}. A group's label says what the
    equations are for ("Vertical (y)"), its note says how to use them. Used where a sheet must say
    which direction an equation belongs to: the kinematic equations are y-only in projectile motion.
    """
    gap(doc, 5)
    t = doc.add_table(rows=1, cols=2)
    fix_widths(t, list(split))
    cells = {"L": t.rows[0].cells[0], "R": t.rows[0].cells[1]}
    for key, c in cells.items():
        borders(c, pal.hair, sz=4, edges=("top", "bottom"))
        cell_margins(c, top=60, bottom=50, left=0 if key == "L" else 140, right=0)
    borders(cells["R"], pal.hair, sz=4, edges=("left",))
    for key, c in cells.items():
        mine = [g for g in groups if g.get("col", "L") == key]
        colw = split[0] if key == "L" else split[1] - 0.14
        for gi, grp in enumerate(mine):
            last = c.paragraphs[-1]
            if gi == 0:
                p = c.paragraphs[0]
            elif not last.text.strip() and not last.runs:
                p = last                      # the empty paragraph python-docx leaves after a table
            else:
                p = c.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(0 if gi == 0 else 5)
            run(p, grp["label"], 9, bold=True, color=pal.ink)
            if grp.get("note"):
                run(p, "   " + grp["note"], 8.5, color=pal.label)
            cols = int(grp.get("columns", 1))
            rows = (len(grp["eqs"]) + cols - 1) // cols
            g = c.add_table(rows=rows, cols=cols)
            fix_widths(g, [round(colw / cols - 0.02, 2)] * cols)
            tight_cells(g, top=25, bottom=35, left=0, right=50)
            for i, eq in enumerate(grp["eqs"]):
                cell = g.rows[i // cols].cells[i % cols]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                equation_cell(cell, eq, pal, size, bold, italic_vars)
            if row_in:
                for r in g.rows:
                    no_split(r, row_in)
        unpad_cell(c)


def masthead(doc, spec, sec, pal):
    unit = f"U{int(spec['unit']):02d}"
    t = doc.add_table(rows=1, cols=2)
    fix_widths(t, [5.55, 1.95])
    h, hr = t.rows[0].cells
    for cc in (h, hr):
        borders(cc, pal.hair, sz=4, edges=("top", "bottom"))
    borders(h, pal.display, sz=30, edges=("left",))
    cell_margins(h, top=80, bottom=80, left=160, right=80)
    cell_margins(hr, top=80, bottom=80, left=80, right=60)
    para(h, f"UNIT {int(spec['unit'])} — {unit_title(spec['course'], spec['unit']).upper()}", 7.5,
         bold=True, color=pal.accent, caps_track=True, first=True)
    para(h, sec["title"], 15, bold=True, color=pal.ink)
    pp = para(hr, f"{unit} / S{sec['code']}", 11, bold=True, color=pal.accent,
              caps_track=True, first=True)
    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pp = para(hr, spec.get("docType", "Quiz").replace("_", " ").upper(), 6.5, color=pal.label,
              caps_track=True)
    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def name_row(doc, total, pal):
    gap(doc, 5)
    widths = [3.05, 1.75, 1.10, 1.60]
    t = doc.add_table(rows=1, cols=4)
    fix_widths(t, widths)
    for cc, lab in zip(t.rows[0].cells, ("NAME", "DATE", "PERIOD", "SCORE")):
        borders(cc, pal.ink, sz=4, edges=("bottom",))
        cell_margins(cc, top=0, bottom=40, left=0, right=120)
        p = para(cc, lab, 6.5, bold=True, color=pal.label, caps_track=True, first=True)
        if lab == "SCORE":
            p.paragraph_format.tab_stops.add_tab_stop(Inches(widths[3] - 0.10), WD_TAB_ALIGNMENT.RIGHT)
            run(p, "\t", 9)
            run(p, f"/  {total}", 10, bold=True, color=pal.ink)


def equations(doc, eqs, pal, width=7.5, per_row=3):
    gap(doc, 5)
    t = doc.add_table(rows=1, cols=1)
    fix_widths(t, [width])
    c = t.rows[0].cells[0]
    borders(c, pal.hair, sz=4, edges=("top", "bottom"))
    cell_margins(c, top=60, bottom=40, left=0, right=0)
    p = para(c, "Equations you may use", 9.5, bold=True, color=pal.ink, first=True)
    p.paragraph_format.space_after = Pt(3)
    rows = (len(eqs) + per_row - 1) // per_row
    g = c.add_table(rows=rows, cols=per_row)
    colw = round(width / per_row - 0.02, 2)
    fix_widths(g, [colw] * per_row)
    tight_cells(g, top=20, bottom=30, left=0, right=60)
    for i, eq in enumerate(eqs):
        cell = g.rows[i // per_row].cells[i % per_row]
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        marked(p, eq, 11, bold=True, color=pal.ink)
    unpad_cell(c)


def card(cell, q, n, pal, inner):
    borders(cell, pal.hair, sz=6)
    cell_margins(cell, top=60, bottom=60, left=110, right=110)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(inner), WD_TAB_ALIGNMENT.RIGHT)
    run(p, f"{n}.", 10.5, bold=True, color=pal.ink)
    run(p, "\t", 8)
    run(p, f"{q['points']} pts", 8, color=pal.label)
    lines = int(q.get("answerLines", 0))
    if not lines:
        para(cell, q["prompt"], 10, color=pal.ink)
        return
    # Fixed rows, so the answer lines sit at the same height in every card of the row, whatever
    # the prompt does. (Adjacent ruled paragraphs merge into one line in Word and LibreOffice.)
    inner_t = cell.add_table(rows=1 + lines, cols=1)
    fix_widths(inner_t, [inner])
    tight_cells(inner_t, left=0, right=0)
    no_split(inner_t.rows[0], 0.50)
    para(inner_t.rows[0].cells[0], q["prompt"], 10, color=pal.ink, first=True)
    for r in inner_t.rows[1:]:
        no_split(r, 0.34)
        borders(r.cells[0], pal.hair, sz=4, edges=("bottom",))


def questions(doc, qs, pal):
    n = 0
    row_qs = []
    while n < len(qs) and qs[n].get("inRow"):
        row_qs.append(qs[n]); n += 1
    if row_qs:
        k = len(row_qs)
        gut = 0.10
        colw = round((7.5 - gut * (k - 1)) / k, 3)
        widths = []
        for i in range(k):
            widths.append(colw)
            if i < k - 1:
                widths.append(gut)
        t = doc.add_table(rows=1, cols=len(widths))
        fix_widths(t, widths)
        no_split(t.rows[0], max(float(q.get("cardHeightIn", 1.5)) for q in row_qs))
        for i, q in enumerate(row_qs):
            card(t.rows[0].cells[i * 2], q, i + 1, pal, round(colw - 0.2, 2))
        gap(doc, 6)
    for q in qs[n:]:
        n += 1
        t = doc.add_table(rows=1, cols=1)
        fix_widths(t, [7.5])
        no_split(t.rows[0], q.get("cardHeightIn", 2.5))
        card(t.rows[0].cells[0], q, n, pal, 7.3)
        gap(doc, 6)


def main():
    if len(sys.argv) < 3:
        return refuse("usage: build_quiz_docx.py <spec.json> <out.docx>")
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    course = spec["course"]
    if course not in COURSE_CODE:
        return refuse(f"unknown course {course!r}.")
    bad = [s for s in spec["sections"] if s not in known_sections(course)]
    if bad:
        return refuse(f"section(s) {', '.join(bad)} are not in courses/{course}/DECISIONS.md.")
    sec = spec["sectionsContent"][0]
    pal = Palette(course)
    qs = sec["questions"]
    total = sum(int(q["points"]) for q in qs)

    doc = Document()
    s = page_setup(doc)
    masthead(doc, spec, sec, pal)
    name_row(doc, total, pal)
    gap(doc, 5)
    c = doc.add_table(rows=1, cols=1); fix_widths(c, [7.5])
    cell = c.rows[0].cells[0]
    cell_margins(cell, top=0, bottom=0, left=0, right=0)
    para(cell, sec["directions"], 9.5, color=pal.ink, first=True)
    if sec.get("equations"):
        equations(doc, sec["equations"], pal)
    gap(doc, 6)
    questions(doc, qs, pal)
    unit = f"U{int(spec['unit']):02d}"
    running_footer(s, f"SHULL SCIENCE          {unit} · S{sec['code']} · {sec['title'].upper()}", pal,
                   with_page_numbers=True)
    trim_tail(doc)
    doc.save(sys.argv[2])
    print(f"wrote {sys.argv[2]}  —  {COURSE_CODE[course]} {unit} S{sec['code']} QUIZ, {len(qs)} question(s), {total} points")
    return 0


if __name__ == "__main__":
    sys.exit(main())
