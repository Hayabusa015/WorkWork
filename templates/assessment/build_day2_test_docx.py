#!/usr/bin/env python3
"""Build a SHULL Day-2 problems test (.docx) from a JSON spec.

    python3 templates/assessment/build_day2_test_docx.py specs/<spec>.json [outdir]

The Chemistry two-day unit test, Day 2: free-response calculation problems, paper and
pencil, work shown and graded for partial credit. Structurally this is the worksheet
template's question card and work box, reused rather than redrawn - the visual system
for "here is a problem, here is where you solve it" is one fact, and a test is not a
different fact from a practice set just because it is graded.

Two things a graded test must never carry that a practice set does:

  - No `selfCheck` brackets. SHULL-CHG-0021 puts an answer in the margin so a student
    can check their own arithmetic on ungraded practice; on a test that same bracket
    hands over the answer being graded. This builder does not accept the field.
  - No tier tags (warm-up / practice / challenge). The ramp is a practice-set teaching
    device; a test problem is just problem N worth N points.

One spec, `versions`, each a fully independent re-solve - different compounds, different
numbers - never a shuffle of the same one. Each version becomes two files: the student
test and a separate `_Key` carrying the worked solution in place of the blank work box.
`validate_versions()` checks the two versions share a blueprint (same problem count,
same point value per slot, same total) before either is built, so a version cannot drift
into an easier or harder form of "equivalent."

Colour comes from brand/tokens.json. No hex is typed in this file.
"""
import json, os, sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from _shull_docx import (          # noqa: E402
    COURSE_CODE, Palette, known_sections, unit_title, section_span,
    borders, para, run, fix_widths, one_cell, no_split,
    equation_bar, work_box, given_need, page_setup, running_footer,
    cell_margins, gap, trim_tail,
)

HERE = os.path.dirname(os.path.abspath(__file__))
TEXT_W_IN = 7.50


def refuse(msg, *more):
    print("build_day2_test_docx: " + msg, file=sys.stderr)
    for m in more:
        print("   " + m, file=sys.stderr)
    return 1


def validate_versions(versions):
    """Blueprint parity: every version has the same number of problems and the same
    point value in each slot. This is a structural check, not a content one - it
    cannot tell whether a version was genuinely re-solved from scratch or copied with
    the numbers swapped, which is why the skill puts "re-solve every version
    independently" on the person writing the spec, not on this script."""
    if not versions:
        return "spec has no versions."
    base = versions[0]
    base_points = [p["points"] for p in base["problems"]]
    for v in versions[1:]:
        pts = [p["points"] for p in v["problems"]]
        if pts != base_points:
            return (f"version {v['label']} point sequence {pts} does not match "
                    f"version {base['label']}'s {base_points}. Parallel versions carry "
                    "identical blueprints and point totals (QA_GATE #9), not just "
                    "the same problem count.")
    for v in versions:
        for i, p in enumerate(v["problems"], 1):
            if "selfCheck" in p:
                return (f"version {v['label']} problem {i} carries \"selfCheck\". "
                        "That is a practice-set field; a graded test never hands over "
                        "the answer being graded. SHULL-CHG-0021 applies to ungraded "
                        "work only.")
            if p.get("tier"):
                return (f"version {v['label']} problem {i} carries a tier tag. Tiers "
                        "are the practice-set ramp; a test problem is just problem "
                        f"{i} worth {p['points']} point(s).")
    return None


def problem_card(cell, p, n, pal, width, *, key=False):
    inner = round(width - 0.28, 2)
    top = cell.paragraphs[0]
    top.paragraph_format.space_after = Pt(2)
    top.paragraph_format.tab_stops.add_tab_stop(Inches(inner), WD_TAB_ALIGNMENT.RIGHT)
    run(top, f"{n}.", 11, bold=True, color=pal.ink)
    run(top, "\t", 8)
    run(top, f"{p['points']} POINT{'S' if p['points'] != 1 else ''}", 7.5,
        bold=True, color=pal.label)

    para(cell, p["prompt"], 10.5, color=pal.ink)
    for part in p.get("parts", []):
        para(cell, "   " + part, 10.5, color=pal.ink)

    if p.get("given") or p.get("need"):
        given_need(cell, p.get("given", ""), p.get("need", ""), pal, inner)

    if key:
        t = cell.add_table(rows=1, cols=1)
        fix_widths(t, [inner])
        no_split(t.rows[0])
        box = t.rows[0].cells[0]
        borders(box, pal.accent, sz=8)
        cell_margins(box, top=55, bottom=55, left=110, right=110)
        para(box, "ANSWER", 7.5, bold=True, color=pal.accent, caps_track=True, first=True)
        para(box, p["answer"]["value"], 10.5, bold=True, color=pal.ink)
        for step in p["answer"].get("workSteps", []):
            para(box, step, 9.5, color=pal.ink)
    else:
        work_box(cell, "", pal, float(p.get("workHeightIn", 1.6)), inner)
    return cell


def render_problems(doc, problems, pal, *, key=False):
    for i, p in enumerate(problems, 1):
        t = doc.add_table(rows=1, cols=1)
        fix_widths(t, [TEXT_W_IN])
        no_split(t.rows[0])
        borders(t.rows[0].cells[0], pal.hair, sz=6)
        cell_margins(t.rows[0].cells[0], top=55, bottom=55, left=130, right=130)
        problem_card(t.rows[0].cells[0], p, i, pal, TEXT_W_IN, key=key)
        gap(doc, 6)


def build_one(spec, version, out, *, key=False):
    course = spec["course"]
    pal = Palette(course)
    unit = f"U{int(spec['unit']):02d}"
    utitle = unit_title(course, spec["unit"])
    span = section_span(spec["sections"])
    day2 = spec["day2"]

    doc = Document()
    s = page_setup(doc)

    t = doc.add_table(rows=1, cols=2)
    fix_widths(t, [5.35, 2.15])
    h, hr = t.rows[0].cells
    for cc in (h, hr):
        borders(cc, pal.hair, sz=4, edges=("top", "bottom"))
    borders(h, pal.display, sz=30, edges=("left",))
    cell_margins(h, top=80, bottom=80, left=160, right=80)
    cell_margins(hr, top=80, bottom=80, left=80, right=60)
    para(h, f"UNIT {int(spec['unit'])} — {utitle.upper()}", 7.5, bold=True,
         color=pal.accent, caps_track=True, first=True)
    para(h, f"Unit Test — Day 2, Problems", 15, bold=True, color=pal.ink)
    pp = para(hr, f"{unit} / {span}", 11, bold=True, color=pal.accent,
              caps_track=True, first=True)
    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pp = para(hr, f"VERSION {version['label']}", 13, bold=True, color=pal.display,
              caps_track=True)
    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if key:
        pp2 = para(hr, "TEACHER KEY", 7.5, bold=True, color=pal.label, caps_track=True)
        pp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    total = sum(p["points"] for p in version["problems"])
    if not key:
        gap(doc, 5)
        widths = [3.05, 1.75, 1.10, 1.60]
        t = doc.add_table(rows=1, cols=4)
        fix_widths(t, widths)
        for cc, lab in zip(t.rows[0].cells, ("NAME", "DATE", "PERIOD", "SCORE")):
            borders(cc, pal.ink, sz=4, edges=("bottom",))
            cell_margins(cc, top=0, bottom=40, left=0, right=120)
            p = para(cc, lab, 6.5, bold=True, color=pal.label, caps_track=True,
                     first=True)
            if lab == "SCORE":
                p.paragraph_format.tab_stops.add_tab_stop(
                    Inches(widths[3] - 0.10), WD_TAB_ALIGNMENT.RIGHT)
                run(p, "\t", 9)
                run(p, f"/  {total}", 10, bold=True, color=pal.ink)
    else:
        gap(doc, 5)
        c = one_cell(doc)
        para(c, f"Answer key — Version {version['label']}  ·  {total} points total",
             9.5, bold=True, color=pal.label, first=True)

    gap(doc, 4)
    if day2.get("equations"):
        equation_bar(doc, day2["equations"], pal, day2.get("equationLabel", "EQUATIONS YOU MAY USE"))

    if day2.get("directions"):
        gap(doc, 3)
        c = one_cell(doc)
        borders(c, pal.ink, sz=8, edges=("top",))
        para(c, "DIRECTIONS", 7.5, bold=True, color=pal.accent, caps_track=True, first=True)
        para(c, day2["directions"], 9.5)

    gap(doc, 4)
    render_problems(doc, version["problems"], pal, key=key)

    running_footer(
        s, f"SHULL SCIENCE          {unit} · {span} · Day 2 · Version {version['label']}"
        + ("  ·  KEY" if key else ""),
        pal, with_page_numbers=True)
    trim_tail(doc)
    doc.save(out)
    return total


def main():
    if len(sys.argv) < 2:
        return refuse("usage: build_day2_test_docx.py <spec.json> [outdir]")
    spec = json.load(open(sys.argv[1]))
    course = spec["course"]
    if course not in COURSE_CODE:
        return refuse(f"unknown course {course!r}.")
    outdir = sys.argv[2] if len(sys.argv) > 2 else HERE

    valid = known_sections(course)
    bad = [s for s in spec["sections"] if s not in valid]
    if bad:
        return refuse(f"section(s) {', '.join(bad)} are not in "
                      f"courses/{course}/DECISIONS.md.",
                      "Nothing is built against a code not in the decisions file.")

    day2 = spec.get("day2")
    if not day2:
        return refuse("spec has no \"day2\" block.")
    err = validate_versions(day2["versions"])
    if err:
        return refuse(err)

    code = COURSE_CODE[course]
    unit = f"U{int(spec['unit']):02d}"
    span = section_span(spec["sections"])

    written = []
    for version in day2["versions"]:
        base = f"SHULL_{code}_Test_{unit}_{span}_Day2_{version['label']}"
        student_out = os.path.join(outdir, base + ".docx")
        key_out = os.path.join(outdir, base + "_Key.docx")
        total = build_one(spec, version, student_out, key=False)
        build_one(spec, version, key_out, key=True)
        written.append((version["label"], total, student_out, key_out))

    for label, total, student_out, key_out in written:
        print(f"wrote {student_out}  —  Version {label}, "
              f"{len(day2['versions'][0]['problems'])} problems, {total} points")
        print(f"wrote {key_out}  —  Version {label} key")
    return 0


if __name__ == "__main__":
    sys.exit(main())
