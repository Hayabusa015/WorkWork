#!/usr/bin/env python3
"""The .docx primitives every SHULL print builder shares.

Guided notes and worksheets are the same design system on paper. When they each
owned a copy of "how a work box is drawn", the copies drifted within a day: the
notes builder learned that `cantSplit` holds a row and the worksheet builder did
not. A fact lives in exactly one place, and the drawing of a work box is a fact.

Nothing here decides *what* goes on a page. Callers pass their own widths, so a
Cornell notes column and a full-width worksheet share the drawing and not the
geometry.

    from _shull_docx import palette, para, work_box, equation_bar

Colour is read from brand/tokens.json. No hex is typed in this file or in any
file that imports it.
"""
import json, os, re
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
COURSE_CODE = {"chemistry": "CHEM", "physics": "PHYS", "geology": "GEO"}

T = json.load(open(os.path.join(REPO, "brand", "tokens.json")))
G = T["ground"]
FONT = T["typography"]["stack"].split(",")[0].strip().strip('"')
FLOOR = T["typography"]["floors"]["printBody"]["pt"]
FOOTER_FLOOR = T["typography"]["floors"]["printFooter"]["pt"]

BULLETS = "•·▸►‣-–— "


def debullet(x):
    """The source packets wrote their own bullets into the text; we supply one."""
    return str(x).lstrip(BULLETS).strip()


def hexof(h):
    return h.lstrip("#").upper()


class Palette:
    """The eight colours a print builder is allowed to use, resolved once.

    `accent` is the text-safe deep course colour; `display` is the brighter one and
    is for rules and borders only - it does not clear contrast as body text.
    """

    def __init__(self, course):
        self.course = course
        self.accent = T["courses"][course]["primaryDeep"]["hex"]
        self.display = T["courses"][course]["primary"]["hex"]
        self.ink = G["asphalt"]["hex"]
        self.hair = G["ruleHairline"]["hex"]
        self.surface = G["parchment"]["hex"]
        self.label = G["label"]["hex"]
        self.footer = G["footer"]["hex"]
        self.white = G["white"]["hex"]

    @property
    def watermark(self):
        """A watermark is defined by ink coverage, not by a colour someone picked.

        `print.watermarkOpacityPct` in tokens.json is the standard; this converts the
        normal band's low end to the grey that puts that much ink on white paper, so
        the number in the design system is the number on the page.
        """
        pct = T["print"]["watermarkOpacityPct"]["normal"][0]
        v = round(255 * (1 - pct / 100.0))
        return "#{0:02X}{0:02X}{0:02X}".format(v)


def known_sections(course):
    text = open(os.path.join(REPO, "courses", course, "DECISIONS.md")).read()
    end = text.find("## Course sequencing rules")
    return set(re.findall(r"(?<![\d.])\d{1,2}\.\d(?![\d])", text[:end] if end > 0 else text))


def unit_title(course, unit):
    """The unit's name, read from that course's DECISIONS.md.

    It is a course fact, so it lives in exactly one place and a spec does not get to
    type it - `validate_layers.py` refuses a spec that does. The three roadmaps are
    formatted differently ("**U7 The Mole...** (7 sections)", "**U1 · Motion...** — 4
    sections", "### U1 · The Universe... — 6 sections"), so all three shapes are read
    rather than one being imposed on files that were correct already.
    """
    text = open(os.path.join(REPO, "courses", course, "DECISIONS.md")).read()
    pat = re.compile(r"^(?:#+\s*)?\*{0,2}U0*%d\b[ ·:—-]*([^*\n(]+?)\*{0,2}\s*(?:\(|—|$)"
                     % int(unit), re.M)
    m = pat.search(text)
    if not m:
        raise SystemExit(f"unit_title: no U{unit} heading in courses/{course}/DECISIONS.md")
    return m.group(1).strip(" ·—-")


def unit_phase(course, unit):
    """The curriculum phase a unit sits in, read from the course's DECISIONS.md, where
    it is a heading: "### Phase 1 — Lab Ready (U0–U3)". A course whose roadmap has no
    phases (Geology) returns None; the caller leaves the phase out rather than invent
    one."""
    text = open(os.path.join(REPO, "courses", course, "DECISIONS.md")).read()
    for m in re.finditer(r"^#+\s*Phase\s+(\d+)\b[^\n(]*\(U(\d+)\s*[–-]\s*U(\d+)\)",
                         text, re.M):
        if int(m.group(2)) <= int(unit) <= int(m.group(3)):
            return int(m.group(1))
    return None


def shade(cell, hexval):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), hexof(hexval))
    cell._tc.get_or_add_tcPr().append(el)


# OOXML fixes the order of the children of w:tcBorders. A file that lists them in
# another order is schema-invalid, and Word repairs it silently while LibreOffice
# may not - so the order is written down once and the merge below respects it.
_BORDER_ORDER = ("top", "start", "left", "bottom", "end", "right", "insideH", "insideV",
                 "tl2br", "tr2bl")


def borders(cell, hexval, sz=6, edges=("top", "left", "bottom", "right"), val="single"):
    """A cell border. `val="dashed"` is the cut line on a cut-and-glue activity - the
    dash is the instruction, so it has to be the border style and not a drawn shape.

    Calls STACK. Two calls on one cell used to append two `w:tcBorders` elements,
    which is invalid - readers were free to take either, and the second call only
    appeared to work because both renderers happened to merge them. It now merges
    into the one element, replacing an edge already set, so a weight ladder can be
    built the obvious way: the hairline box first, then the heavier accent edge.

        borders(cue, pal.hair, sz=8)                       # the box
        borders(cue, pal.display, sz=12, edges=("left",))  # the rail on top of it
    """
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for e in edges:
        old = b.find(qn(f"w:{e}"))
        if old is not None:
            b.remove(old)
        x = OxmlElement(f"w:{e}")
        x.set(qn("w:val"), val); x.set(qn("w:sz"), str(sz))
        x.set(qn("w:color"), hexof(hexval))
        b.append(x)
    order = {n: i for i, n in enumerate(_BORDER_ORDER)}
    for child in sorted(list(b), key=lambda c: order.get(c.tag.split("}")[-1], 99)):
        b.append(child)


def checkbox(paragraph, pal, size=9):
    """An empty tick box, drawn in the brand face rather than borrowed from another.

    U+2610 ☐ is not in Archivo. Setting it pulls DejaVu into the PDF, and the box a
    student actually ticks is then drawn by a font nobody chose - the same failure the
    lab template hit with U+25AA. A bordered run of spaces is a real square in Archivo
    metrics, so the file renders in one face. The space count is measured, not guessed,
    so the box stays square if the size changes.
    """
    sp = " " * max(2, round(size / (text_width_in(" ", size) * 72.0)))
    r = paragraph.add_run(sp)
    r.font.name = FONT
    r.font.size = Pt(size)
    bdr = OxmlElement("w:bdr")
    bdr.set(qn("w:val"), "single"); bdr.set(qn("w:sz"), "6")
    bdr.set(qn("w:space"), "0"); bdr.set(qn("w:color"), hexof(pal.ink))
    r._element.get_or_add_rPr().append(bdr)
    return r


def check_item(cell, text, pal, size=9):
    """A checklist line: a drawn box, then the text."""
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    checkbox(p, pal, size)
    run(p, "   " + text, size, color=pal.ink)
    return p


def run(paragraph, text, size, *, bold=False, color=None, italic=False):
    r = paragraph.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(hexof(color))
    return r


def para(cell, text, size, *, bold=False, color=None, spacing=0.14, caps_track=False, first=False):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(hexof(color))
    if caps_track:
        # True is the standard label tracking; an int is a looser or tighter one, for
        # a label that has to read as quieter than the headings around it. Tracking is
        # part of how loud a caps label is, so demoting one means turning this down as
        # well as the size and the colour.
        val = 26 if caps_track is True else int(caps_track)
        el = OxmlElement("w:spacing"); el.set(qn("w:val"), str(val))
        r._element.get_or_add_rPr().append(el)
    return p


def rule_lines(cell, n, hexval):
    for _ in range(n):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        pPr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr"); x = OxmlElement("w:bottom")
        x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:color"), hexof(hexval))
        b.append(x); pPr.append(b)


def fix_widths(table, widths):
    """Word autofits tables and will happily discard the column widths python-docx
    sets. Every split in this system is deliberate, so it is pinned: autofit off, a
    fixed layout, and the width written on every cell."""
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w)
    for col, w in zip(table.columns, widths):
        col.width = Inches(w)


def gap(doc, pt=6):
    """A vertical gap of exactly `pt` points.

    `doc.add_paragraph()` with a space_after is not a gap of that size - it is a full
    empty body line PLUS the space after it, so a 3pt spacer actually costs about 16pt.
    Six of them between the questions of one section is most of an inch of nothing, and
    on a one-page-per-section sheet that inch is what pushes the closing checklist onto
    a second sheet of paper. Pinning the line makes the gap the size it says it is.
    """
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "0")
    sp.set(qn("w:line"), str(int(pt * 20))); sp.set(qn("w:lineRule"), "exact")
    pPr.append(sp)
    r = p.add_run()
    r.font.name = FONT
    r.font.size = Pt(max(1, pt - 2))
    return p


def trim_tail(doc):
    """Make sure the body ends with a paragraph, and that the paragraph is 1pt.

    OOXML requires the document body to end with a paragraph, and a body built entirely
    out of tables does not have one - python-docx will happily write `</w:tbl></w:body>`.
    Word and LibreOffice both repair that by inserting a paragraph of their own, at full
    body height, and when the content already reaches the bottom of the page that
    inserted paragraph does not fit: the file gains a completely blank final page
    carrying nothing but the running footer.

    So the paragraph is written deliberately and pinned to a 1pt exact line. The file
    becomes valid and the page count is the content's, rather than an artefact of how
    the file had to be built.
    """
    body = doc.element.body
    last = None
    for el in body:
        if el.tag.split("}")[-1] != "sectPr":
            last = el
    def pin(el):
        pPr = el.get_or_add_pPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "0")
        sp.set(qn("w:line"), "20"); sp.set(qn("w:lineRule"), "exact")
        pPr.append(sp)
        rPr = OxmlElement("w:rPr")
        for name in ("w:sz", "w:szCs"):
            e = OxmlElement(name); e.set(qn("w:val"), "2"); rPr.append(e)
        pPr.append(rPr)

    if last is not None and last.tag.split("}")[-1] == "p" \
            and not "".join(last.itertext()).strip():
        pin(last)
        return
    p = doc.add_paragraph()
    pin(p._p)
    sect = body.find(qn("w:sectPr"))
    if sect is not None:
        body.remove(p._p)
        sect.addprevious(p._p)


def cell_margins(cell, top=0, bottom=0, left=40, right=40):
    """Margins on ONE cell, in twentieths of a point.

    Applying them to a whole table is the usual case, but not always the right one: a
    right-aligned tag needs no margin on its left and a real gap on its right, and
    zeroing the row wholesale ran the tag straight into the text it labels.
    """
    mar = OxmlElement("w:tcMar")
    for edge, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), str(int(v))); e.set(qn("w:type"), "dxa")
        mar.append(e)
    cell._tc.get_or_add_tcPr().append(mar)


def tight_cells(t, top=0, bottom=0, left=40, right=40):
    for row in t.rows:
        for c in row.cells:
            cell_margins(c, top, bottom, left, right)


def one_cell(doc, width=7.5, protect=False):
    """`protect=True` holds the whole block together across a page break (cantSplit) -
    for a box like the section summary, where one orphaned line stranded on an
    otherwise-blank page is worse than the whole box moving down as a unit."""
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.LEFT
    fix_widths(t, [width])
    if protect:
        no_split(t.rows[0])
    return t.rows[0].cells[0]


def no_split(row, height_in=None):
    """Hold a table row together across a page break, and optionally floor its height.

    hRule "atLeast" means the box grows with what a student writes into it and never
    shrinks below the floor. cantSplit holds ONE row - it does nothing across two, so
    anything that must stay together goes in a single row.
    """
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    if height_in is not None:
        h = OxmlElement("w:trHeight")
        h.set(qn("w:val"), str(int(float(height_in) * 1440)))
        h.set(qn("w:hRule"), "atLeast")
        trPr.append(h)


# The fraction bar has to be at least as wide as the widest of numerator and
# denominator or the term wraps inside the fraction, which is worse than a loose bar.
# A character-count guess has to be padded for that safety; the shipped Archivo files
# say exactly how wide the term is, so the pad can be the small optical one a fraction
# bar wants and nothing more.
_FONTS = {}


def text_width_in(text, size_pt, bold=False):
    key = "Bold" if bold else "Regular"
    if key not in _FONTS:
        from fontTools.ttLib import TTFont
        tt = TTFont(os.path.join(REPO, "brand", "fonts", f"Archivo-{key}.ttf"))
        _FONTS[key] = (tt.getBestCmap(), tt["hmtx"], tt["head"].unitsPerEm)
    cmap, hmtx, upem = _FONTS[key]
    total = 0
    for ch in text:
        g = cmap.get(ord(ch))
        total += hmtx[g][0] if g else upem // 2
    return total * size_pt / upem / 72.0


def unpad_cell(cell):
    """Strip the blank paragraphs python-docx leaves around a nested table.

    A cell that gets a table starts with an empty paragraph above it and python-docx
    adds another below (Word requires a cell to end with one). Both are full-size and
    both count toward the row height, so a vertically centred neighbour - the "v  ="
    beside a fraction - centres against that padding and lands above the fraction bar
    instead of on it. The leading blank can go outright; the trailing one has to stay,
    so it is collapsed to 1pt.
    """
    body = cell._tc
    kids = list(body)
    tags = [k.tag.split("}")[-1] for k in kids]
    if "tbl" not in tags:
        return
    first_tbl = tags.index("tbl")
    for k, tag in list(zip(kids, tags))[:first_tbl]:
        if tag == "p" and not "".join(k.itertext()).strip():
            body.remove(k)
    for k, tag in reversed(list(zip(list(body), [x.tag.split("}")[-1] for x in body]))):
        if tag != "p":
            break
        if "".join(k.itertext()).strip():
            break
        pPr = k.get_or_add_pPr()
        rPr = OxmlElement("w:rPr")
        for name in ("w:sz", "w:szCs"):
            e = OxmlElement(name); e.set(qn("w:val"), "2"); rPr.append(e)
        pPr.append(rPr)
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:after"), "0"); sp.set(qn("w:before"), "0")
        sp.set(qn("w:line"), "20"); sp.set(qn("w:lineRule"), "exact")
        pPr.append(sp)
        break


# A stacked fraction, never a/b inline. SHULL_DESIGN_SYSTEM section 8, SHULL-CHG-0017.
# Built as a ruled two-row table rather than Word's own equation objects (OMML):
# OMML is valid and Word renders it, but LibreOffice will not import it from .docx, so
# nothing in this pipeline could see what shipped. A ruled table renders identically in
# Word, LibreOffice, Google Docs and print, and can be checked.
def stacked_frac(cell, num, den, ink_hex, rule_hex, size=10.5):
    t = cell.add_table(rows=2, cols=1)
    tight_cells(t)
    term = max(text_width_in(num, size), text_width_in(den, size))
    fix_widths(t, [max(0.40, round(term + 0.16, 3))])
    top, bot = t.rows[0].cells[0], t.rows[1].cells[0]
    borders(top, rule_hex, sz=6, edges=("bottom",))
    for c, txt in ((top, num), (bot, den)):
        p = para(c, txt, size, color=ink_hex, first=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
    unpad_cell(cell)
    return t


def equation_bar(doc, equations, pal, label, width=7.5):
    """The equations students may reference, at the top of the page.

    His rule: on any Physics or Chemistry sheet with math, the equations go at the top
    so students know what they are allowed to reach for. Two across, no box around the
    block - the only rules on it are the fraction bars, which are the point of it.
    """
    c = one_cell(doc, width)
    para(c, label, 7.5, bold=True, color=pal.accent, caps_track=True, first=True)
    cols = 2
    half = round(width / 2 - 0.13, 2)
    rows = (len(equations) + cols - 1) // cols
    t = c.add_table(rows=rows, cols=cols)
    fix_widths(t, [half, half])
    tight_cells(t, top=30, bottom=40, left=0, right=60)
    for i, eq in enumerate(equations):
        cell = t.rows[i // cols].cells[i % cols]
        if eq.get("note"):
            p = para(cell, eq["note"], 7.5, bold=True, color=pal.accent,
                     caps_track=True, first=True)
            p.paragraph_format.space_after = Pt(1)
        if eq.get("num") and eq.get("den"):
            # lhs and the fraction side by side. A nested table always starts its own
            # line, so the equals sign needs its own cell to sit beside the bar.
            inner = cell.add_table(rows=1, cols=2)
            fix_widths(inner, [0.80, half - 1.02])
            tight_cells(inner, left=0, right=0)
            lc, fc = inner.rows[0].cells
            # Middle-aligned so the equals sign lands on the fraction bar, not above it.
            lc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            para(lc, (eq.get("lhs", "") + "  =") if eq.get("lhs") else "", 10.5,
                 bold=True, first=True).alignment = WD_ALIGN_PARAGRAPH.RIGHT
            stacked_frac(fc, eq["num"], eq["den"], pal.ink, pal.hair)
            unpad_cell(cell)
        else:
            para(cell, eq.get("plain", ""), 10.5, bold=True, first=not eq.get("note"))
    return c


def work_box(cell, label, pal, height_in=1.4, width=5.06, watermark=None, *,
             label_size=7.5, label_bold=True, label_caps=True, label_color=None):
    """An open bordered box a student solves a problem in.

    Ruled lines are for prose; open boxes are for math. The box grows with the work
    and never breaks across a page.

    `watermark` puts a faint centred instruction in the empty space - his ask for the
    Chemistry sheets, "a watermark that says Show work here in the area". Its grey
    comes from print.watermarkOpacityPct in tokens.json, so the ink standard and the
    page agree by construction.

    The label keywords let a caller set the label as a quiet sentence-case caption
    ("WORK / show your reasoning and units", small and grey) instead of a tracked caps
    heading. The defaults are the worksheet's, unchanged.
    """
    t = cell.add_table(rows=1, cols=1)
    fix_widths(t, [width])
    no_split(t.rows[0], height_in)
    inner = t.rows[0].cells[0]
    borders(inner, pal.hair, sz=6)
    # The label is read; the border is not. Hairline on white measures 1.6:1.
    if label:
        para(inner, label, label_size, bold=label_bold, color=label_color or pal.accent,
             caps_track=label_caps, first=True)
    if watermark:
        p = para(inner, watermark, 12, color=pal.watermark, caps_track=True,
                 first=not label)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(int(max(6, height_in * 72 * 0.32)))
    return inner


def given_need(cell, given, need, pal, inner_w):
    t = cell.add_table(rows=2, cols=2)
    fix_widths(t, [0.72, inner_w - 0.72])
    for ri, (lab, val) in enumerate((("GIVEN", given), ("NEED", need))):
        a, b = t.rows[ri].cells
        borders(a, pal.hair, sz=4); borders(b, pal.hair, sz=4)
        para(a, lab, 7.5, bold=True, color=pal.accent, caps_track=True, first=True)
        para(b, val, 9.5, first=True)


def fillin_table(cell, spec, pal, inner_w, *, header_fill=None, header_caps=True,
                 label_bold=True, header_size=7.5, body_size=9.5):
    """A small comparison table a student fills in by hand.

    His example: charge/mass/location for proton, neutron, electron was three
    separate "what's the charge, mass, and location of X" questions, asked once per
    particle, each with its own blank lines - the same three-part question in disguise
    three times over. When several open questions are really one row of a comparison
    table, this draws the table instead: one header row naming what belongs in each
    column, then one bordered, hand-writeable row per item.

    `spec` is `{"headers": [...], "rows": [[...], ...], "widths": [...]}`. `widths` is
    optional column fractions of `inner_w` (must sum to ~1); omitted, the first column
    (the row's own label - "Proton", "Cl-35") gets a narrower share than the rest. An
    empty string in a cell is a blank for the student to write in; text there is the
    key's answer, and the key's answers are bold - the same distinction a must-write
    line marks with its rule, made here with weight instead of a line, because a table
    cell has no left edge to draw one on.

    No body-cell fill anywhere - hairline borders only, the same print-ink rule as
    `given_need` and `diagram_block`. Every body row is `no_split` and floored tall
    enough to hand-write a word or a number into, because a table with zero-height
    blank cells is not a graphic organizer, it is a smaller way to fail a student.

    `header_fill` is the one exception, and only the header row takes it: section 8
    allows "a very light grey tint" where shading is genuinely needed for
    scanability. Pass `pal.surface` (ground.parchment) - never a course colour. The
    default is no fill, so the worksheet is unchanged.
    """
    headers = spec["headers"]
    rows = spec["rows"]
    ncols = len(headers)
    widths = spec.get("widths")
    if widths:
        widths = [round(inner_w * w, 3) for w in widths]
    else:
        label_w = min(round(inner_w * 0.22, 3), 1.15)
        rest_w = round((inner_w - label_w) / (ncols - 1), 3)
        widths = [label_w] + [rest_w] * (ncols - 1)

    t = cell.add_table(rows=1 + len(rows), cols=ncols)
    fix_widths(t, widths)
    tight_cells(t, top=24, bottom=24, left=50, right=50)

    for ci, h in enumerate(headers):
        c = t.rows[0].cells[ci]
        borders(c, pal.hair, sz=4)
        if header_fill:
            shade(c, header_fill)
        para(c, h, header_size, bold=True, color=pal.accent, caps_track=header_caps,
             first=True)

    row_h = float(spec.get("rowHeightIn", 0.34))
    for ri, vals in enumerate(rows):
        tr = t.rows[ri + 1]
        no_split(tr, row_h)
        for ci, val in enumerate(vals):
            c = tr.cells[ci]
            borders(c, pal.hair, sz=4)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if ci == 0:
                para(c, val, 9 if label_bold else body_size, bold=label_bold,
                     color=pal.ink, first=True)
            else:
                filled = bool(str(val).strip())
                para(c, val, body_size, bold=filled, color=pal.ink, first=True)
    return t


def diagram_block(cell, dgm, pal, inner_w, base_dir=HERE):
    """A figure students label.

    His example: the layers of the Earth, or a plate boundary with the trench and the
    mid-ocean ridge to identify. The figure gets real room and the label lines are
    numbered beside it, so a student writes on the sheet rather than in a margin.

    The heading and caption live INSIDE the non-splitting row. cantSplit holds ONE row
    and does nothing across two, so kept outside it the heading stranded on the
    previous page while the figure moved to the next - a box with no instruction above
    it, which is worse than no heading at all.
    """
    fig_w = inner_w * 0.61
    t = cell.add_table(rows=1, cols=2)
    fix_widths(t, [fig_w, inner_w * 0.39])
    fig, labels = t.rows[0].cells
    para(fig, dgm.get("label", "LABEL THIS DIAGRAM"), 7.5, bold=True, color=pal.accent,
         caps_track=True, first=True)
    if dgm.get("caption"):
        para(fig, dgm["caption"], 9)
    borders(fig, pal.hair, sz=4)
    no_split(t.rows[0], float(dgm.get("heightIn", 3.0)))
    img = dgm.get("image")
    path = os.path.join(base_dir, img) if img else None
    if path and os.path.exists(path):
        r = fig.add_paragraph().add_run()
        r.add_picture(path, width=Inches(fig_w - 0.20))
        fig.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        # NOT first=True. The heading already claimed paragraph 0, and passing it twice
        # appends a run to the same paragraph instead of making a new one - the heading
        # and the figure note ran together as one line.
        p = para(fig, dgm.get("figureNote", "[ FIGURE ]"), 9, color=pal.hair, caps_track=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
    para(labels, "LABEL", 7.5, bold=True, color=pal.accent, caps_track=True, first=True)
    for i in range(int(dgm.get("labelSlots", 6))):
        # Just the number. No box, no square, no circle. SHULL-CHG-0015, no exceptions.
        para(labels, f"{i + 1}", 9.5, bold=True, color=pal.ink)
        rule_lines(labels, 1, pal.hair)
    return cell


def page_setup(doc, margin_lr=0.5, margin_tb=0.44, base_pt=10):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(base_pt)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Inches(margin_lr)
    s.top_margin = s.bottom_margin = Inches(margin_tb)
    return s


def running_footer(section, text, pal, with_page_numbers=False):
    """The footer every printed page carries.

    Outlined and small: printFooter in tokens.json is the only floor a student page is
    allowed below printBody, and this is the exception it exists for.
    """
    f = section.footer.paragraphs[0]
    f.text = text
    for r in f.runs:
        r.font.name = FONT
        r.font.size = Pt(FOOTER_FLOOR)
        r.font.color.rgb = RGBColor.from_string(hexof(pal.footer))
    if with_page_numbers:
        r = f.add_run("          PAGE ")
        r.font.name = FONT; r.font.size = Pt(FOOTER_FLOOR)
        r.font.color.rgb = RGBColor.from_string(hexof(pal.footer))
        _field(f, "PAGE", pal)
        r = f.add_run(" OF ")
        r.font.name = FONT; r.font.size = Pt(FOOTER_FLOOR)
        r.font.color.rgb = RGBColor.from_string(hexof(pal.footer))
        _field(f, "NUMPAGES", pal)
    return f


def field(paragraph, instr, pal, *, size=None, bold=False, color=None):
    """A Word field. Page numbers have to be computed by the reader, not by us -
    a worksheet that says "page 2 of 3" in fixed text lies the moment it is edited.

    `instr` is the whole field code, switches included (`PAGE \\# "00"`)."""
    r = paragraph.add_run()
    r.font.name = FONT; r.font.size = Pt(size or FOOTER_FLOOR)
    r.bold = bold
    r.font.color.rgb = RGBColor.from_string(hexof(color or pal.footer))
    beg = OxmlElement("w:fldChar"); beg.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r._r.append(beg); r._r.append(it); r._r.append(end)
    return r


def _field(paragraph, instr, pal):
    return field(paragraph, instr, pal)


# w:sectPr children in schema order, so a page-number format lands where both
# readers expect it rather than wherever an append happens to put it.
_SECT_ORDER = ("headerReference", "footerReference", "footnotePr", "endnotePr", "type",
               "pgSz", "pgMar", "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols",
               "formProt", "vAlign", "noEndnote", "titlePg", "textDirection", "bidi",
               "rtlGutter", "docGrid", "printerSettings", "sectPrChange")


def page_number_format(section, fmt="decimalZero"):
    """Set how a section's PAGE field renders - "decimalZero" is 01, 02 ... 10.

    A `\\# "00"` picture switch on the field does the same job in Word, but
    LibreOffice ignores it and prints "1". The section-level number format is honoured
    by both, so a two-digit page number survives the PDF this pipeline checks as well
    as the .docx the teacher edits. Callers use both: this for LibreOffice, the switch
    for any reader that honours field switches over section formats.
    """
    sp = section._sectPr
    old = sp.find(qn("w:pgNumType"))
    if old is not None:
        sp.remove(old)
    el = OxmlElement("w:pgNumType"); el.set(qn("w:fmt"), fmt)
    sp.append(el)
    _sort_children(sp, _SECT_ORDER)
    return el


# ---------------------------------------------------------------------------------
# Schema order. OOXML fixes the order of the children of pPr, rPr, tcPr and tblPr,
# and every helper in this file appends - so a cell that is bordered, then margined,
# then shaded, carries its children in call order, not schema order. LibreOffice
# reads either. Word is stricter, and "Word found unreadable content" on a packet the
# teacher opens at 7:40 a.m. is the failure the .docx was chosen to avoid. One pass
# at save time puts every property list back in schema order.
_PPR_ORDER = ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
              "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
              "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
              "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
              "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
              "suppressOverlap", "jc", "textDirection", "textAlignment",
              "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
              "pPrChange")
_RPR_ORDER = ("rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
              "dstrike", "outline", "shadow", "emboss", "imprint", "noProof",
              "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
              "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd",
              "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
              "specVanish", "oMath")
_TCPR_ORDER = ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
               "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark")
_TBLPR_ORDER = ("tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
                "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
                "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook")
_PBDR_ORDER = ("top", "left", "bottom", "right", "between", "bar")


def _sort_children(el, order):
    rank = {n: i for i, n in enumerate(order)}
    kids = list(el)
    # Keep only the LAST of any duplicated single-occurrence child: two w:spacing
    # elements in one pPr is invalid, and the later one is the one the caller meant.
    seen, keep = set(), []
    for k in reversed(kids):
        tag = k.tag.split("}")[-1]
        if tag in rank and tag in seen:
            el.remove(k)
            continue
        seen.add(tag)
        keep.append(k)
    for k in sorted(reversed(keep), key=lambda c: rank.get(c.tag.split("}")[-1], 999)):
        el.append(k)


def schema_order(doc):
    """Put every pPr / rPr / tcPr / tblPr / pBdr in the body, headers and footers into
    schema order. Call once, immediately before `doc.save()`."""
    parts = [doc.element.body]
    for s in doc.sections:
        parts += [s.header._element, s.footer._element]
    for root in parts:
        for tag, order in (("w:pPr", _PPR_ORDER), ("w:rPr", _RPR_ORDER),
                           ("w:tcPr", _TCPR_ORDER), ("w:tblPr", _TBLPR_ORDER),
                           ("w:pBdr", _PBDR_ORDER)):
            for el in root.iter(qn(tag)):
                _sort_children(el, order)


# ---------------------------------------------------------------------------------
# Height estimation. A .docx has no layout of its own - the reader lays it out - so a
# builder that promises "this designed page is one sheet of paper" has to predict what
# the reader will do. This walks the XML actually written and predicts the height of a
# paragraph or table in points, using the shipped Archivo metrics for line wrapping.
#
# It is exact for what it is exact for: a paragraph pinned to an exact line height
# (every paragraph a notes page writes is) and a row with an explicit height. Line
# wrapping is a greedy fill against real advance widths, which is what both Word and
# LibreOffice do. It is not a layout engine - callers keep a safety margin, and the
# rendered PDF is still what QA looks at.
_TWIP = 1 / 20.0
_DEFAULT_CELL_MAR = (0, 108, 0, 108)       # top, left, bottom, right - "Normal Table"
_AUTO_LINE = 1.088                         # hhea ascent + descent, Archivo, per pt


def _attr(el, name, default=None):
    v = el.get(qn(name)) if el is not None else None
    return default if v is None else v


def _para_runs(p, base_pt):
    """(text, size_pt, bold) for every run, plus a list of inline picture heights."""
    out, pics = [], []
    for r in p.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        size = base_pt
        bold = False
        if rpr is not None:
            sz = rpr.find(qn("w:sz"))
            if sz is not None:
                size = int(_attr(sz, "w:val")) / 2.0
            b = rpr.find(qn("w:b"))
            if b is not None and _attr(b, "w:val", "1") not in ("0", "false"):
                bold = True
        for ch in r:
            tag = ch.tag.split("}")[-1]
            if tag == "t":
                out.append((ch.text or "", size, bold))
            elif tag == "tab":
                out.append(("    ", size, bold))
            elif tag == "br" and _attr(ch, "w:type") in (None, "textWrapping"):
                out.append(("\n", size, bold))
            elif tag == "drawing":
                for ext in ch.iter("{http://schemas.openxmlformats.org/drawingml/2006/"
                                   "wordprocessingDrawing}extent"):
                    pics.append(int(ext.get("cy")) / 12700.0)
                    break
    return out, pics


def _wrap_lines(runs, width_in):
    """Greedy line fill over (text, size, bold) runs. Returns the number of lines."""
    if width_in <= 0:
        return 1
    lines, x = 1, 0.0
    word_w, space_w, pending = 0.0, 0.0, False
    for text, size, bold in runs:
        for ch in text:
            if ch == "\n":
                lines += 1; x = 0.0; word_w = 0.0; pending = False
                continue
            w = text_width_in(ch, size, bold)
            if ch == " ":
                if pending:
                    x += word_w; word_w = 0.0; pending = False
                x += w
                continue
            if not pending:
                # a new word starts; wrap if it will not fit where the line has got to
                pending = True
            word_w += w
            if x + word_w > width_in and x > 0:
                lines += 1
                x = 0.0
    return lines


def estimate_height(el, width_in, base_pt=10.0):
    """Predicted rendered height, in points, of a w:p or w:tbl element."""
    tag = el.tag.split("}")[-1]
    if tag == "p":
        ppr = el.find(qn("w:pPr"))
        sp = ppr.find(qn("w:spacing")) if ppr is not None else None
        before = int(_attr(sp, "w:before", 0)) * _TWIP
        after = int(_attr(sp, "w:after", 0)) * _TWIP
        runs, pics = _para_runs(el, base_pt)
        size = max([s for _, s, _ in runs] or [base_pt])
        rule = _attr(sp, "w:lineRule", "auto")
        line = _attr(sp, "w:line")
        if line is not None and rule in ("exact", "atLeast"):
            line_pt = int(line) * _TWIP
            if rule == "atLeast":
                line_pt = max(line_pt, size * _AUTO_LINE)
        else:
            mult = int(line) / 240.0 if line is not None else 1.0
            line_pt = size * _AUTO_LINE * mult
        ind = ppr.find(qn("w:ind")) if ppr is not None else None
        avail = width_in - (int(_attr(ind, "w:left", 0)) + int(_attr(ind, "w:right", 0))) \
            * _TWIP / 72.0
        n = _wrap_lines(runs, avail)
        body = n * line_pt
        if pics:
            body = max(body - line_pt, 0) + sum(pics) + 2
        bdr = 0.0
        pb = ppr.find(qn("w:pBdr")) if ppr is not None else None
        if pb is not None:
            for edge in ("top", "bottom"):
                e = pb.find(qn(f"w:{edge}"))
                if e is not None:
                    bdr += int(_attr(e, "w:sz", 4)) / 8.0 + int(_attr(e, "w:space", 0))
        return before + body + after + bdr
    if tag == "tbl":
        total = 0.0
        for tr in el.findall(qn("w:tr")):
            natural = 0.0
            for tc in tr.findall(qn("w:tc")):
                tcpr = tc.find(qn("w:tcPr"))
                vm = tcpr.find(qn("w:vMerge")) if tcpr is not None else None
                if vm is not None and _attr(vm, "w:val") != "restart":
                    continue
                w = tcpr.find(qn("w:tcW")) if tcpr is not None else None
                cw = int(_attr(w, "w:w", 1440)) * _TWIP / 72.0
                mt, ml, mb, mr = _DEFAULT_CELL_MAR
                mar = tcpr.find(qn("w:tcMar")) if tcpr is not None else None
                if mar is not None:
                    get = lambda e, d: int(_attr(mar.find(qn(f"w:{e}")), "w:w", d))
                    mt, ml, mb, mr = get("top", mt), get("left", ml), get("bottom", mb), \
                        get("right", mr)
                inner = cw - (ml + mr) * _TWIP / 72.0
                h = sum(estimate_height(k, inner, base_pt) for k in tc
                        if k.tag.split("}")[-1] in ("p", "tbl"))
                natural = max(natural, h + (mt + mb) * _TWIP)
            trpr = tr.find(qn("w:trPr"))
            th = trpr.find(qn("w:trHeight")) if trpr is not None else None
            if th is not None:
                val = int(_attr(th, "w:val", 0)) * _TWIP
                rule = _attr(th, "w:hRule", "atLeast")
                row = val if rule == "exact" else max(val, natural)
            else:
                row = natural
            total += row
        return total
    return 0.0
