#!/usr/bin/env python3
"""Build a SHULL practice set (the student copy) in the plain layout.

    python3 templates/worksheet/build_practice_docx.py <spec.json> <out.docx>

For sets students work in their notebooks. What is on the page, top to bottom: the masthead, the
Name/Date/Period/Score line, one short paragraph of directions, the equations they may use, and
the questions. Nothing else. A question is its number, then its text on the same line, then
lettered parts if it has them, then the self-check bracket at the right. Questions are divided by
a thin rule, not boxed.

Same masthead, palette, fonts and footer as every other SHULL print document, and the same
drawn-Greek handling as the quiz builder (the brand face has no theta or delta). The answer key is
built from the same spec by build_worksheet_key_docx.py.

Spec, one section: code, title, directions (plain text), equations (strings; "_i" subscript,
"^2" superscript, "[n|d]" a stacked fraction, "θ" and "Δ" drawn), equationNote (optional, shown
beside the label), equationColumns (default 3), questions[] with points, prompt, parts[], selfCheck.
"""
import json, os, sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATES = os.path.dirname(HERE)
sys.path.insert(0, TEMPLATES)
sys.path.insert(0, os.path.join(TEMPLATES, "quiz"))
from _shull_docx import (          # noqa: E402
    COURSE_CODE, Palette, known_sections, borders, para, run, fix_widths, no_split, page_setup,
    running_footer, cell_margins, gap, trim_tail,
)
from build_quiz_docx import masthead, name_row, equation_block   # noqa: E402

INDENT = 0.36                # the number's column; the text and the parts hang from it
BODY = 11                    # question text
LEAD = 1.12                  # line spacing for question text


def question(doc, q, n, pal):
    t = doc.add_table(rows=1, cols=1)
    fix_widths(t, [7.5])
    no_split(t.rows[0])
    c = t.rows[0].cells[0]
    borders(c, pal.hair, sz=4, edges=("bottom",))
    cell_margins(c, top=100, bottom=90, left=0, right=0)

    p = c.paragraphs[0]
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent, pf.space_after = Inches(INDENT), Inches(-INDENT), Pt(3)
    pf.line_spacing = LEAD
    pf.tab_stops.add_tab_stop(Inches(INDENT))
    run(p, f"{n}.", 12, bold=True, color=pal.accent)
    run(p, "\t", BODY)
    run(p, q["prompt"], BODY, color=pal.ink)
    for part in q.get("parts", []):
        pp = para(c, part, BODY, color=pal.ink)
        pp.paragraph_format.left_indent = Inches(INDENT)
        pp.paragraph_format.space_after = Pt(1.5)
        pp.paragraph_format.line_spacing = LEAD
    if q.get("selfCheck"):
        sp = para(c, q["selfCheck"], 8, color=pal.label)
        sp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        sp.paragraph_format.space_before = Pt(2)


def main():
    if len(sys.argv) < 3:
        print("usage: build_practice_docx.py <spec.json> <out.docx>", file=sys.stderr)
        return 1
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    course = spec["course"]
    bad = [s for s in spec["sections"] if s not in known_sections(course)]
    if bad:
        print(f"build_practice_docx: section(s) {', '.join(bad)} are not in courses/{course}/DECISIONS.md.",
              file=sys.stderr)
        return 1
    sec = spec["sectionsContent"][0]
    pal = Palette(course)
    qs = sec["questions"]
    total = sum(int(q.get("points", 1)) for q in qs)

    doc = Document()
    s = page_setup(doc)
    masthead(doc, spec, sec, pal)
    name_row(doc, total, pal)
    gap(doc, 5)
    c = doc.add_table(rows=1, cols=1)
    fix_widths(c, [7.5])
    cell = c.rows[0].cells[0]
    cell_margins(cell, top=0, bottom=0, left=0, right=0)
    para(cell, sec["directions"], 9.5, color=pal.ink, first=True)
    equation_block(doc, sec["equations"], pal, note=sec.get("equationNote"), size=9.5, bold=False,
                   italic_vars=True, row_in=0.26, per_row=int(sec.get("equationColumns", 3)))
    gap(doc, 4)
    for i, q in enumerate(qs, start=1):
        question(doc, q, i, pal)
    unit = f"U{int(spec['unit']):02d}"
    running_footer(s, f"SHULL SCIENCE          {unit} · S{sec['code']} · {sec['title'].upper()}", pal,
                   with_page_numbers=True)
    trim_tail(doc)
    doc.save(sys.argv[2])
    print(f"wrote {sys.argv[2]}  —  {COURSE_CODE[course]} {unit} S{sec['code']}, {len(qs)} question(s), {total} points")
    return 0


if __name__ == "__main__":
    sys.exit(main())
