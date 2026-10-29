#!/usr/bin/env python3
"""Chemistry U02 test reference sheet: a periodic table read by sublevel (Matt, 2026-10-29).

    python3 templates/worksheet/build_u02_reference.py OUTDIR

One landscape page. Every row carries its sublevel labels in front of each block, in filling
order (row 4 reads 4s, 3d, 4p). The f-block sits below as two 14-element rows, 4f and 5f, with
its insertion point marked between Ba and Lu and between Ra and Lr, so the sublevel counts come
out exactly: 6s holds Cs and Ba, 4f holds La to Yb, 5d starts at Lu. Valence electrons are listed
over the main-group columns. The blank space above the d-block holds the diagonal filling
diagram, the filling order row by row, and a sublevel capacity table (Matt, 2026-10-29: no
exception marking on this sheet).

Every placement is checked against the Aufbau generator before anything is written.
"""
import os, sys
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(REPO, "templates"))
sys.path.insert(0, os.path.join(REPO, "templates", "notes"))
from build_electron_set import (                                    # noqa: E402
    SYMBOLS, config, shorthand, NOBLE_SYM, new_doc, header, Palette, COURSE, UNIT,
)
from _shull_docx import (                                           # noqa: E402
    hexof, borders, fix_widths, no_split, gap, one_cell, running_footer, trim_tail,
)
from build_notes_packet import cell_margins, _run, rich, Ctx       # noqa: E402
from _shull_name import document_name                               # noqa: E402

# Column plan (22 columns): sublevel label before each block, plus the f-block marker.
#   sL g1 g2 | fM | dL g3..g12 | pL g13..g18
SL, FM, DL, PL = 0, 3, 4, 15
COL = {1: 1, 2: 2, **{g: g + 2 for g in range(3, 13)}, **{g: g + 3 for g in range(13, 19)}}
W_LABEL, W_MARK, W_EL = 0.30, 0.44, 0.475
WIDTHS = [W_LABEL, W_EL, W_EL, W_MARK, W_LABEL] + [W_EL] * 10 + [W_LABEL] + [W_EL] * 6
ROW_H = 0.56

EXCEPTIONS = {24, 29}            # Cr and Cu, the two taught (S2.3 notes)
VALENCE = {1: "1", 2: "2", 13: "3", 14: "4", 15: "5", 16: "6", 17: "7", 18: "8"}


def layout():
    """(period, group) -> Z for the main table, and the two f rows. Lu and Lr sit in group 3."""
    main = {(1, 1): 1, (1, 18): 2}
    z = 3
    for per, start in ((2, 3), (3, 11)):
        z = start
        for g in (1, 2):
            main[(per, g)] = z; z += 1
        for g in range(13, 19):
            main[(per, g)] = z; z += 1
    for per, start in ((4, 19), (5, 37)):
        z = start
        for g in range(1, 19):
            main[(per, g)] = z; z += 1
    frows = {}
    for per, start in ((6, 55), (7, 87)):
        z = start
        main[(per, 1)] = z; main[(per, 2)] = z + 1
        frows[per] = list(range(z + 2, z + 16))          # 14 f elements
        z += 16
        for g in range(3, 19):
            main[(per, g)] = z; z += 1
    return main, frows


def last_sublevel(z):
    n, l, _ = config(z)[-1]
    return f"{n}{l}"


def check(main, frows):
    """Every cell's last-filled sublevel must match the label printed in front of its block."""
    for (per, g), z in main.items():
        want = (f"{per}s" if g <= 2 else f"{per - 1}d" if g <= 12 else f"{per}p")
        if z == 2:
            want = "1s"
        got = last_sublevel(z)
        if z in EXCEPTIONS:
            assert got == "3d", (z, got)          # 4s1 3d5 / 4s1 3d10: still a d element
            continue
        assert got == want, (SYMBOLS[z - 1], (per, g), got, want)
    for per, zs in frows.items():
        assert len(zs) == 14
        for z in zs:
            assert last_sublevel(z) == f"{per - 2}f", (SYMBOLS[z - 1], last_sublevel(z))
    assert sorted(list(main.values()) + sum(frows.values(), [])) == list(range(1, 119))
    print("reference table checked: 118 elements, every sublevel label matches the generator")


def exact(row, h):
    no_split(row, h)
    row._tr.trPr.find(qn("w:trHeight")).set(qn("w:hRule"), "exact")


def cell_text(cell, pal, text, size, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER,
              valign=WD_ALIGN_VERTICAL.CENTER):
    cell.vertical_alignment = valign
    cell_margins(cell, top=0, bottom=0, left=10, right=10)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = align
    if text:
        _run(p, text, size, bold=bold, color=color or pal.ink)
    return p


def element(cell, pal, z):
    borders(cell, pal.ink, sz=6)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    cell_margins(cell, top=10, bottom=0, left=25, right=15)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    _run(p, str(z), 6.5, color=pal.label)
    q = cell.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after = Pt(0)
    _run(q, SYMBOLS[z - 1], 11.5, bold=True, color=pal.ink)


def sub_label(cell, pal, text):
    cell_text(cell, pal, text, 7.5, bold=True, color=pal.accent, align=WD_ALIGN_PARAGRAPH.RIGHT)


def diagonal_png(path, pal):
    """The diagonal (Aufbau) filling diagram: s p d f columns, n = 1 to 7 rows. Each arrow runs
    from upper right to lower left through the sublevels with the same n + l, so reading the
    arrows in turn, top to bottom, gives the filling order. 1s and 2s get their own short arrow."""
    from PIL import Image, ImageDraw, ImageFont
    rows = {1: "s", 2: "sp", 3: "spd", 4: "spdf", 5: "spdf", 6: "spd", 7: "sp"}
    dx, dy, x0, y0 = 150, 100, 70, 70
    W, H = x0 + 4 * dx, y0 + 7 * dy
    im = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(im)
    f = ImageFont.truetype(os.path.join(REPO, "brand", "fonts", "Archivo-Bold.ttf"), 62)
    ink = tuple(int(hexof(pal.ink)[i:i + 2], 16) for i in (0, 2, 4))
    acc = tuple(int(hexof(pal.accent)[i:i + 2], 16) for i in (0, 2, 4))
    centre = {}
    for n, ls in rows.items():
        for i, l in enumerate(ls):
            centre[(n, l)] = (x0 + i * dx + 40, y0 + (n - 1) * dy)
    ux, uy = -dx, dy                                   # one step down-left along a diagonal
    L = (ux ** 2 + uy ** 2) ** 0.5
    ux, uy = ux / L, uy / L
    for key in range(1, 11):
        pts = sorted([(n, l) for (n, l) in centre if n + "spdf".index(l) == key])
        if not pts:
            continue
        (xa, ya), (xb, yb) = centre[pts[0]], centre[pts[-1]]
        sx, sy = xa, ya                                # start under the first label's box
        ex, ey = xb + ux * 55, yb + uy * 55            # end just below-left of the last
        d.line((sx, sy, ex, ey), fill=acc, width=10)
        px, py = -uy, ux
        d.polygon([(ex + ux * 30, ey + uy * 30), (ex + px * 19, ey + py * 19),
                   (ex - px * 19, ey - py * 19)], fill=acc)
    for (n, l), (cx, cy) in centre.items():
        text = f"{n}{l}"
        bb = d.textbbox((0, 0), text, font=f)
        tw, th = bb[2] - bb[0], bb[3] - bb[1]
        d.rectangle((cx - tw / 2 - 8, cy - th / 2 - 8, cx + tw / 2 + 8, cy + th / 2 + 8),
                    fill="white")
        d.text((cx - tw / 2 - bb[0], cy - th / 2 - bb[1]), text, font=f, fill=ink)
    im.save(path)


def build(out, imgdir):
    pal = Palette(COURSE)
    ctx = Ctx(pal, False)
    main, frows = layout()
    check(main, frows)
    doc, s = new_doc(landscape=True)
    total_w = sum(WIDTHS)
    header(doc, pal, "UNIT 2 TEST  ·  REFERENCE  ·  DO NOT WRITE ON THIS SHEET",
           "Periodic Table by Sublevel", width=total_w, fields=False)
    gap(doc, 4)

    t = doc.add_table(rows=3 + 7, cols=len(WIDTHS))
    fix_widths(t, WIDTHS)
    # ---- header rows: blocks, group numbers, valence electrons ----
    exact(t.rows[0], 0.20); exact(t.rows[1], 0.17); exact(t.rows[2], 0.20)
    blocks = [("s-BLOCK", COL[1], COL[2]), ("d-BLOCK", COL[3], COL[12]),
              ("p-BLOCK", COL[13], COL[18])]
    for name, a, b in blocks:
        c = t.cell(0, a).merge(t.cell(0, b))
        borders(c, pal.accent, sz=14, edges=("bottom",))
        cell_text(c, pal, name, 7, bold=True, color=pal.accent)
    for g in range(1, 19):
        cell_text(t.cell(1, COL[g]), pal, str(g), 7, bold=True, color=pal.label)
    for g, v in VALENCE.items():
        p_ = cell_text(t.cell(2, COL[g]), pal, v, 8.5, bold=True, color=pal.ink)
        if g == 18:
            _run(p_, "  He: 2", 6.5, color=pal.label)
    dv = t.cell(2, COL[3]).merge(t.cell(2, COL[12]))
    cell_text(dv, pal, "valence = the electrons in the outer s sublevel", 6.5,
              color=pal.label)
    gp_ = cell_text(t.cell(1, FM).merge(t.cell(1, DL)), pal, "group", 6.5, bold=True,
                    color=pal.label, align=WD_ALIGN_PARAGRAPH.RIGHT)
    # the valence row is named in the gap between the s and d blocks
    vp = cell_text(t.cell(2, FM).merge(t.cell(2, DL)), pal, "", 6.5,
                   align=WD_ALIGN_PARAGRAPH.RIGHT)
    rich(vp, "valence e^−^", 6.5, ctx, bold=True, color=pal.label)

    # ---- periods ----
    for per in range(1, 8):
        r = t.rows[2 + per]
        exact(r, ROW_H + (0.07 if per <= 3 else 0))   # periods 1-3 hold the diagram box
        for (pp, g), z in main.items():
            if pp == per:
                element(t.cell(2 + per, COL[g]), pal, z)
        sub_label(t.cell(2 + per, SL), pal, f"{per}s")
        if per >= 2:
            sub_label(t.cell(2 + per, PL), pal, f"{per}p")
        if per >= 4:
            sub_label(t.cell(2 + per, DL), pal, f"{per - 1}d")
        if per in frows:
            zs = frows[per]
            c = t.cell(2 + per, FM)
            borders(c, pal.accent, sz=10, edges=("left", "right"))
            cell_text(c, pal, f"{zs[0]}–{zs[-1]}", 6.5, bold=True, color=pal.accent)
            q = c.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_after = Pt(0)
            _run(q, f"{per - 2}f ↓", 7, bold=True, color=pal.accent)
    # He's 1s label sits right before He, in the g17 cell
    sub_label(t.cell(3, COL[17]), pal, "1s")

    # ---- the blank space above the d-block: diagonal diagram + capacity table ----
    gap_cell = t.cell(3, FM).merge(t.cell(5, COL[12]))
    gap_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    cell_margins(gap_cell, top=20, bottom=0, left=60, right=40)
    inner = gap_cell.add_table(rows=2, cols=3)
    fix_widths(inner, [1.55, 2.0, 1.8])
    exact(inner.rows[0], 0.19)
    for ci, h in enumerate(["DIAGONAL RULE", "FILLING ORDER, ROW BY ROW", "SUBLEVEL CAPACITY"]):
        c = inner.cell(0, ci)
        borders(c, pal.accent, sz=10, edges=("bottom",))
        cell_text(c, pal, h, 6.5, bold=True, color=pal.accent, align=WD_ALIGN_PARAGRAPH.LEFT)
    img = os.path.join(imgdir, "u02_reference_diagonal.png")
    diagonal_png(img, pal)
    c0 = inner.cell(1, 0); c0.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    cell_margins(c0, top=20, bottom=0, left=0, right=0)
    p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = Pt(0)
    p0.add_run().add_picture(img, height=Inches(1.3))
    # filling order written as the periodic-table rows read, one row per line
    c1 = inner.cell(1, 1); c1.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    cell_margins(c1, top=40, bottom=0, left=60, right=0)
    order = {1: "1s", 2: "2s  2p", 3: "3s  3p", 4: "4s  3d  4p", 5: "5s  4d  5p",
             6: "6s  4f  5d  6p", 7: "7s  5f  6d  7p"}
    for k, (row, subs) in enumerate(order.items()):
        q = c1.paragraphs[0] if k == 0 else c1.add_paragraph()
        q.paragraph_format.space_after = Pt(1)
        q.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))
        _run(q, f"Row {row}", 7, color=pal.label)
        _run(q, "\t" + subs, 8, bold=True, color=pal.ink)
    # capacity: one row per sublevel
    c2 = inner.cell(1, 2); c2.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    cell_margins(c2, top=40, bottom=0, left=60, right=0)
    cap = c2.add_table(rows=5, cols=3)
    fix_widths(cap, [0.55, 0.6, 0.55])
    for ri, rowvals in enumerate([["", "ORBITALS", "MAX e^−^"], ["s", "1", "2"],
                                  ["p", "3", "6"], ["d", "5", "10"], ["f", "7", "14"]]):
        exact(cap.rows[ri], 0.18)
        for ci, v in enumerate(rowvals):
            c = cap.cell(ri, ci)
            if ri == 0:
                borders(c, pal.hair, sz=6, edges=("bottom",))
            cp = cell_text(c, pal, "", 8)
            rich(cp, v, 6.5 if ri == 0 else 8.5, ctx, bold=(ri == 0 or ci == 0),
                 color=pal.label if ri == 0 else (pal.accent if ci == 0 else pal.ink))
    # the cell's own empty first paragraph would push the table down a line; make it 1 pt
    from docx.shared import Pt as _Pt
    pf = c2.paragraphs[0].paragraph_format
    pf.space_after = _Pt(0); pf.line_spacing = _Pt(1)
    # the empty first-row cells between H and He stay blank; no borders anywhere outside cells

    # ---- f-block rows: one continuous 14-element strip, lined up under the d-block ----
    gap(doc, 6)
    lead = WIDTHS[SL] + WIDTHS[1] + WIDTHS[2] + WIDTHS[FM]
    ft = doc.add_table(rows=3, cols=2 + 14)
    fix_widths(ft, [lead, W_LABEL] + [W_EL] * 14)
    exact(ft.rows[0], 0.20)
    hc = ft.cell(0, 2).merge(ft.cell(0, 15))
    borders(hc, pal.accent, sz=14, edges=("bottom",))
    cell_text(hc, pal, "f-BLOCK  ·  fits between groups 2 and 3 in rows 6 and 7", 7, bold=True,
              color=pal.accent)
    for i, per in enumerate((6, 7)):
        exact(ft.rows[i + 1], ROW_H)
        name = ("LANTHANIDES", "ACTINIDES")[i]
        cell_text(ft.cell(i + 1, 0), pal, name, 7, bold=True, color=pal.label,
                  align=WD_ALIGN_PARAGRAPH.RIGHT)
        sub_label(ft.cell(i + 1, 1), pal, f"{per - 2}f")
        for k, z in enumerate(frows[per]):
            element(ft.cell(i + 1, 2 + k), pal, z)

    # ---- notes ----
    gap(doc, 6)
    nt = doc.add_table(rows=1, cols=2)
    fix_widths(nt, [total_w / 2, total_w / 2])
    notes = [
        ("READING A ROW",
         "Read left to right and say the sublevel labels as you pass them. Row 4: 4s, then 3d, "
         "then 4p. The d sublevel is always one level lower than the row; f is two lower."),
        ("NOBLE-GAS SHORTHAND",
         "Start from the noble gas at the end of the row above, in brackets: [He] [Ne] [Ar] "
         "[Kr] [Xe] [Rn]. Then write the rest of the row in filling order."),
    ]
    for i, (h, body) in enumerate(notes):
        c = nt.cell(0, i)
        c.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        borders(c, pal.hair, sz=8, edges=("top",))
        cell_margins(c, top=40, bottom=0, left=40, right=120)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        _run(p, h, 6.8, bold=True, color=pal.ink)
        q = c.add_paragraph(); q.paragraph_format.space_after = Pt(0)
        rich(q, body, 8, ctx, color=pal.ink)

    running_footer(s, f"SHULL SCIENCE          {UNIT}  ·  UNIT TEST REFERENCE", pal,
                   with_page_numbers=False)
    trim_tail(doc)
    doc.save(out)


def main():
    outdir = sys.argv[1] if len(sys.argv) > 1 else "."
    imgdir = os.path.join(outdir, ".dots"); os.makedirs(imgdir, exist_ok=True)
    name = document_name("Reference", ["2.2", "2.4"], 2, descriptor="Periodic Table by Sublevel")
    build(os.path.join(outdir, name), imgdir)
    print("wrote", name)


if __name__ == "__main__":
    main()
