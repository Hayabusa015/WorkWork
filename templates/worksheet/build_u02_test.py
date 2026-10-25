#!/usr/bin/env python3
"""Chemistry U02 unit test, first draft (Matt, 2026-10-24): one version.

    python3 templates/worksheet/build_u02_test.py OUTDIR

Day 1 is conceptual and runs on Edulastic, so it is built as a teacher build sheet: every item in
the order the Edulastic editor asks for it (question type, stem, choices, correct answer, points,
why each wrong choice is wrong), plus the PNG images for the label-an-image items. Day 2 is the
paper, handwritten part: configurations, orbital diagrams, ions and one explanation. Unit 2 has
no wave math (Matt, 2026-10-20), so the paper part is write-it-out skills rather than calculations.

Every configuration, valence count and ion on the paper test is computed by the Aufbau generator
and checked against a hand-written table before anything is written. Nothing copies a study
guide question (DECISIONS.md, Assessment shape).
"""
import os, sys
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_u02_assessments as Q                                   # noqa: E402
from build_u02_assessments import (                                 # noqa: E402
    REPO, config, shorthand, valence, flat, unpaired, ion_charge, NOBLE_SYM, SYMBOLS,
    new_doc, header, Palette, COURSE, UNIT, Ctx, rich, _run, one_cell, borders, para, gap,
    fix_widths, no_split, cell_margins, running_footer, trim_tail, pin_tail, hand, hand_config,
    key_strip, cellfmt, part_bar, numcell, tile, tile_png, item, small_table, section_bar,
    page_break_before, NUM_W, INK, document_name,
)

FONT_B = os.path.join(REPO, "brand", "fonts", "Archivo-Bold.ttf")
FONT_R = os.path.join(REPO, "brand", "fonts", "Archivo-Regular.ttf")

# =======================================================================================
# Day 2, paper
# =======================================================================================
W_LONG = [23, 13]            # V reaches 3d (Hund in d), Al stops at 3p^1
W_SHORT = [80, 24]           # Hg has 4f^14; Cr is the taught exception
W_IONS = [37, 4, 34, 13]     # Rb, Be, Se, Al
Q.NAMES.update({23: "Vanadium", 80: "Mercury", 24: "Chromium", 37: "Rubidium",
                4: "Beryllium"})
Q.MASS.update({23: "50.94", 80: "200.59", 24: "52.00", 37: "85.47", 4: "9.01"})

EXPECT = {
    "long": {23: "1s2 2s2 2p6 3s2 3p6 4s2 3d3", 13: "1s2 2s2 2p6 3s2 3p1"},
    "unpaired": {23: 3, 13: 1},
    "short": {80: "[Xe] 6s2 4f14 5d10", 24: "[Ar] 4s1 3d5"},
    # atom: (valence e-, ion, electrons in ion, noble gas it matches)
    "ions": {37: (1, "Rb^+^", 36, "Kr"), 4: (2, "Be^2+^", 2, "He"),
             34: (6, "Se^2−^", 36, "Kr"), 13: (3, "Al^3+^", 10, "Ne")},
}
NOBLE_Z = {"He": 2, "Ne": 10, "Ar": 18, "Kr": 36, "Xe": 54}
PTS = {"long": 4, "short": 2, "why": 2, "ion": 2, "flame": 3}


def check_written():
    for z in W_LONG:
        assert flat(config(z)) == EXPECT["long"][z], (z, flat(config(z)))
        assert unpaired(z) == EXPECT["unpaired"][z], (z, unpaired(z))
    for z in W_SHORT:
        c, r = shorthand(z)
        assert f"[{NOBLE_SYM[c]}] " + flat(r) == EXPECT["short"][z], (z, c, flat(r))
    assert any(l == "f" for _, l, _ in shorthand(80)[1])
    for z, (v, ion, n_e, gas) in EXPECT["ions"].items():
        assert valence(z)[1] == v, (z, valence(z))
        charge = ion_charge(z)
        sign = -1 if charge.endswith("−") else 1
        assert z - sign * int(charge[:-1]) == n_e == NOBLE_Z[gas], (z, charge, n_e, gas)
    total = (len(W_LONG) * PTS["long"] + len(W_SHORT) * PTS["short"] + PTS["why"]
             + len(W_IONS) * PTS["ion"] + PTS["flame"])
    assert total == 25, total
    print("written test answers re-solved: all match")
    return total


def long_row(doc, pal, n, z, key):
    """The quiz's Part 2 row with exact heights, plus an unpaired-electron count."""
    t = doc.add_table(rows=3, cols=3)
    fix_widths(t, [NUM_W, 1.15, 5.9])
    numcell(t.cell(0, 0).merge(t.cell(2, 0)), pal, n)
    tc = t.cell(0, 1).merge(t.cell(2, 1))
    tc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell_margins(tc, top=0, bottom=0, left=0, right=0)
    tile(tc, pal, z, 0.96)
    for row, h in zip(t.rows, (0.38, 0.06, 0.56)):
        no_split(row, h)
        row._tr.trPr.find(qn("w:trHeight")).set(qn("w:hRule"), "exact")
    gp = t.cell(1, 2).paragraphs[0]; gp.paragraph_format.space_after = Pt(0); _run(gp, "", 2)
    cell_margins(t.cell(1, 2), top=0, bottom=0, left=0, right=0)
    lc = t.cell(0, 2)
    pl = cellfmt(lc, pal, top=0, bottom=20, left=80)
    lc.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    pl.paragraph_format.tab_stops.add_tab_stop(Inches(1.25))
    pl.paragraph_format.tab_stops.add_tab_stop(Inches(5.8), WD_TAB_ALIGNMENT.RIGHT)
    _run(pl, "Configuration", 9, bold=True, color=pal.ink)
    _run(pl, "\t", 9)
    if key:
        hand_config(pl, z, size=12)
    _run(pl, "\tUnpaired", 9, bold=True, color=pal.ink)
    Q.blank(pl, str(unpaired(z)), key, 4, size=9)
    bc = t.cell(2, 2)
    borders(bc, pal.hair, sz=6)
    bc.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    cell_margins(bc, top=30, bottom=20, left=80, right=40)
    pb = bc.paragraphs[0]; pb.paragraph_format.space_after = Pt(1)
    _run(pb, "Orbital diagram", 8.5, color=pal.label)
    if key:
        key_strip(bc, pal, z, lab_w=0.22, box_w=0.19)
    gap(doc, 5)


def note(doc, pal, text):
    """Grading guidance on the key, in the arrow ink."""
    c = one_cell(doc); cell_margins(c, top=10, bottom=10, left=60, right=40)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    hand(p, text, 9.5, ink="arrow")


def build_written(out, key, total):
    pal = Palette(COURSE)
    ctx = Ctx(pal, key)
    doc, s = new_doc()
    header(doc, pal, "UNIT 2 TEST  ·  DAY 2  ·  WRITTEN" + ("  ·  TEACHER KEY" if key else ""),
           "Unit 2 Test: Written", score=total)
    gap(doc, 3)
    c = one_cell(doc); borders(c, pal.display, sz=18, edges=("left",))
    p = cellfmt(c, pal, line=False, top=20, bottom=20)
    _run(p, "You may use a periodic table. Write every configuration in order of filling. In an "
            "orbital diagram, each electron gets one arrow.", 9.5, color=pal.ink)
    n = 1

    part_bar(doc, pal, "Part 1   Longhand configurations and orbital diagrams",
             len(W_LONG) * PTS["long"])
    for z in W_LONG:
        long_row(doc, pal, n, z, key)
        n += 1
    if key:
        note(doc, pal, "4 pts each. Configuration 2, diagram 1, unpaired count 1. "
                       "Take 1 off the diagram for each rule broken.")

    part_bar(doc, pal, "Part 2   Noble-gas shorthand",
             len(W_SHORT) * PTS["short"] + PTS["why"])
    t = doc.add_table(rows=1, cols=2)
    fix_widths(t, [3.75, 3.75])
    no_split(t.rows[0])
    for k, z in enumerate(W_SHORT):
        cc = t.cell(0, k)
        cc.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        cell_margins(cc, top=0, bottom=0, left=0, right=160)
        p0 = cc.paragraphs[0]; p0.paragraph_format.space_after = Pt(1)
        _run(p0, str(n), 13, bold=True, color=pal.accent)
        pt_ = cc.add_paragraph(); pt_.paragraph_format.space_after = Pt(2)
        pt_.add_run().add_picture(tile_png(z), width=Inches(0.8))
        lt = cc.add_table(rows=1, cols=1)
        fix_widths(lt, [3.5]); no_split(lt.rows[0], 0.36)
        lcell = lt.rows[0].cells[0]
        lcell.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
        borders(lcell, pal.hair, sz=6, edges=("bottom",))
        cell_margins(lcell, top=0, bottom=20, left=40, right=0)
        q = lcell.paragraphs[0]; q.paragraph_format.space_after = Pt(0)
        if key:
            hand_config(q, z, short=True, size=12)
        pin_tail(cc)
        n += 1
    gap(doc, 4)
    item(doc, pal, n, "Chromium's configuration is not the one the filling order predicts. "
         "Write the predicted shorthand, then explain why the actual one has lower energy.", key,
         lambda q: (hand(q, "Predicted [Ar] 4s^2^ 3d^4^.", 11),
                    hand(q, "   4s and 3d are close in energy. A half-full 3d (five "
                            "unpaired) is more stable.", 10, ink="arrow")),
         lines=2, ctx=ctx)
    n += 1
    if key:
        note(doc, pal, "Shorthand 2 pts each (core 1, rest in order 1). Explanation 2 pts, "
                       "predicted 1 and why 1.")

    page_break_before(doc)
    part_bar(doc, pal, "Part 3   Ions", len(W_IONS) * PTS["ion"])
    c = one_cell(doc); cell_margins(c, top=0, bottom=20, left=40, right=40)
    Q.sg_prompt(c, n, "Complete the table. For the last column, name the noble gas with the "
                "same number of electrons as the ion.", ctx, pal, first=True, after=0)
    pin_tail(c)
    rows, ans = [], []
    for z in W_IONS:
        v, ion, n_e, gas = EXPECT["ions"][z]
        rows.append([Q.name_of(z), None, None, None, None])
        ans.append([None, str(v), ion, str(n_e), gas])
    small_table(doc, pal, ["ATOM", "VALENCE e^−^", "ION IT FORMS", "e^−^ IN THE ION",
                           "SAME AS"], rows, ans, key, [1.9, 1.3, 1.5, 1.5, 1.3], ctx,
                row_h=0.40)
    n += 1
    if key:
        note(doc, pal, "2 pts per row, half a point per box.")

    part_bar(doc, pal, "Part 4   Flame colors", PTS["flame"])
    item(doc, pal, n, "A sodium flame is yellow (589 nm). A lithium flame is crimson red "
         "(671 nm). Explain how the heated atoms give off light, and say which element's "
         "color comes from the larger energy drop and how you know.", key,
         lambda q: hand(q, "Heat excites electrons to higher levels. They drop back and give "
                           "off photons.", 10.5),
         lines=4, ctx=ctx)
    if key:
        c = doc.tables[-1].rows[0].cells[0]
        hand(c.tables[1].rows[0].cells[0].paragraphs[0],
             "Sodium. 589 nm is shorter, so more energy, so a bigger drop.", 10.5)
        note(doc, pal, "1 pt for excited electrons dropping and giving off light. 1 pt for sodium. "
                       "1 pt for shorter wavelength means more energy.")

    running_footer(s, f"SHULL SCIENCE          {UNIT}  ·  UNIT TEST DAY 2" +
                   ("          TEACHER KEY" if key else ""), pal, with_page_numbers=True)
    trim_tail(doc)
    doc.save(out)


# =======================================================================================
# Day 1, Edulastic. Images first.
# =======================================================================================
def _font(path, size):
    from PIL import ImageFont
    return ImageFont.truetype(path, size)


def orbital_blank_png(path, subs):
    """A blank orbital diagram: sublevel labels and empty boxes, for drop-downs on top."""
    from PIL import Image, ImageDraw
    box, lab, gapw = 120, 110, 50
    width = 60 + sum(lab + box * k + gapw for _, k in subs)
    im = Image.new("RGB", (width, 240), "white")
    d = ImageDraw.Draw(im)
    f = _font(FONT_B, 48)
    x = 40
    for name, k in subs:
        bb = d.textbbox((0, 0), name, font=f)
        d.text((x + lab - 16 - (bb[2] - bb[0]), 92), name, font=f, fill=(40, 40, 40))
        x += lab
        for i in range(k):
            d.rectangle((x + i * box, 60, x + (i + 1) * box, 180), outline=(20, 20, 20),
                        width=5)
        x += box * k + gapw
    im.save(path)


def periodic_blocks_png(path):
    """An empty periodic table: every element is a cell, each block outlined heavily, and the
    f block set below. No symbols, so the drop-down boxes are the only labels."""
    from PIL import Image, ImageDraw
    c = 70
    x0, y0 = 40, 40
    im = Image.new("RGB", (x0 * 2 + 18 * c, y0 * 2 + 10 * c), "white")
    d = ImageDraw.Draw(im)
    cells = set()
    cells |= {(1, 1), (1, 18)}
    for per in range(2, 8):
        for g in (1, 2):
            cells.add((per, g))
        for g in range(13, 19):
            cells.add((per, g))
    for per in range(4, 8):
        for g in range(3, 13):
            cells.add((per, g))
    for (per, g) in cells:
        x, y = x0 + (g - 1) * c, y0 + (per - 1) * c
        d.rectangle((x, y, x + c, y + c), outline=(150, 150, 150), width=2)
    fy = y0 + 7.6 * c
    for r in range(2):
        for i in range(14):
            x, y = x0 + (i + 3) * c, fy + r * c
            d.rectangle((x, y, x + c, y + c), outline=(150, 150, 150), width=2)
    ink, w = (20, 20, 20), 6

    def outline(g1, g2, p1, p2):
        d.rectangle((x0 + (g1 - 1) * c, y0 + (p1 - 1) * c, x0 + g2 * c, y0 + p2 * c),
                    outline=ink, width=w)
    outline(1, 2, 2, 7); outline(1, 1, 1, 1)        # s block (H; He sits over the p block)
    outline(3, 12, 4, 7)                            # d block
    outline(13, 18, 2, 7); outline(18, 18, 1, 1)    # p block, with He
    d.rectangle((x0 + 3 * c, fy, x0 + 17 * c, fy + 2 * c), outline=ink, width=w)   # f block
    im.save(path)


LADDER_DROPS = [("A", 2, 1), ("B", 5, 2), ("C", 3, 2), ("D", 4, 3)]


def ladder_drops_png(path):
    """Hydrogen-style energy levels (not to scale) with four drops labeled A to D."""
    from PIL import Image, ImageDraw
    W, H = 2000, 1000
    im = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(im)
    f = _font(FONT_B, 40)
    top, bot = 80, 940
    frac = {1: 0.0, 2: 0.42, 3: 0.62, 4: 0.75, 5: 0.84, 6: 0.90}

    def y(n):
        return bot - frac[n] * (bot - top)
    ink = (50, 50, 50)
    for n in range(1, 7):
        d.line((240, y(n), 1900, y(n)), fill=ink, width=4)
        d.text((110, y(n) - 24), f"n = {n}", font=f, fill=ink)
    d.text((1620, bot + 10 - 60), "not to scale", font=_font(FONT_R, 34), fill=(120, 120, 120))
    for i, (lab, a, b) in enumerate(LADDER_DROPS):
        x = 520 + i * 340
        d.line((x, y(a), x, y(b) - 30), fill=ink, width=7)
        d.polygon([(x, y(b)), (x - 18, y(b) - 36), (x + 18, y(b) - 36)], fill=ink)
        d.text((x + 22, (y(a) + y(b)) / 2 - 24), lab, font=_font(FONT_B, 52), fill=ink)
    im.save(path)


# The items. type = the Edulastic question type to pick; body = what to type in, in order;
# answer = the correct response(s); why = why each wrong choice is wrong.
MC_T, CLOZE, CLASS, MATCH, MATRIX, LABEL = ("Multiple Choice", "Cloze with Drop Down",
                                            "Classification", "Match List", "Choice Matrix",
                                            "Label Image with Drop Down")

ITEMS = [
    # ---------------- S2.1 ----------------
    dict(sec="S2.1", type=MC_T, pts=1,
         stem="In the Bohr model, what does it mean that an electron's energy is quantized?",
         choices=["It can only have certain set amounts of energy.",
                  "It can have any amount of energy, depending on its speed.",
                  "Its energy is highest when it is closest to the nucleus.",
                  "It loses energy steadily as it circles the nucleus."],
         answer="A",
         why="B treats energy as continuous. C reverses the levels: n = 1 is the lowest. "
             "D is the classical spiral the Bohr model was built to fix."),
    dict(sec="S2.1", type=CLOZE, pts=1,
         stem="An electron in the lowest energy level it can occupy is in the [1] state. "
              "After it absorbs energy and moves to a higher level, it is in the [2] state. "
              "To return to a lower level, it must [3] energy.",
         blanks=[("ground", ["ground", "excited"]), ("excited", ["ground", "excited"]),
                 ("release", ["release", "absorb"])],
         why="Mixing up the two states, or thinking an electron absorbs energy on the way down."),
    dict(sec="S2.1", type=MC_T, pts=1,
         stem="Which event moves an electron to a higher energy level?",
         choices=["The atom absorbs a photon.", "The atom gives off a photon.",
                  "The atom loses a proton.", "The atom gains a neutron."],
         answer="A",
         why="B is the drop, not the climb. C and D change the nucleus, not the electron's level."),
    # ---------------- S2.2 ----------------
    dict(sec="S2.2", type=MATRIX, pts=2,
         stem="Mark the most electrons each sublevel can hold.",
         rows=["s", "p", "d", "f"], cols=["2", "6", "10", "14"],
         answer="s = 2, p = 6, d = 10, f = 14",
         why="The usual slip is giving the number of orbitals (1, 3, 5, 7) instead of electrons."),
    dict(sec="S2.2", type=CLOZE, pts=1,
         stem="The n = 4 energy level has [1] sublevels. The last of them is [2]. "
              "Altogether, n = 4 can hold up to [3] electrons.",
         blanks=[("4", ["3", "4", "5"]), ("4f", ["4d", "4f", "4g"]), ("32", ["18", "32", "50"])],
         why="3 sublevels and 18 electrons describe n = 3. There is no g sublevel at this level."),
    dict(sec="S2.2", type=MC_T, pts=1,
         stem="How many orbitals does a p sublevel have, and how many electrons can it hold?",
         choices=["3 orbitals, 6 electrons", "6 orbitals, 6 electrons", "3 orbitals, 3 electrons",
                  "1 orbital, 2 electrons"],
         answer="A",
         why="B counts electrons as orbitals. C allows one electron per orbital. D describes s."),
    # ---------------- S2.3 ----------------
    dict(sec="S2.3", type=LABEL, pts=2, img="u02_test_sulfur_orbitals.png",
         stem="Complete the orbital diagram for sulfur, S. Pick what goes in each box.",
         label="Put one drop-down over each of the 9 boxes, left to right. Options in every "
               "drop-down: ↑↓  ·  ↑  ·  empty.",
         answer="1s ↑↓ · 2s ↑↓ · 2p ↑↓ ↑↓ ↑↓ · 3s ↑↓ · 3p ↑↓ ↑ ↑",
         why="Pairing the third 3p box (↑↓ ↑↓ empty) breaks Hund's rule."),
    dict(sec="S2.3", type=CLASS, pts=2,
         stem="Sort each orbital diagram by the rule it breaks.",
         cols=["Follows all rules", "Breaks Pauli", "Breaks Hund", "Breaks Aufbau"],
         sort=[("2p  [↑] [↑] [↑]", "Follows all rules"),
               ("3p  [↑↓] [↑↓] [↑]", "Follows all rules"),
               ("2s  [↑↑]", "Breaks Pauli"),
               ("1s  [↓↓]", "Breaks Pauli"),
               ("2p  [↑↓] [↑] [   ]", "Breaks Hund"),
               ("3d  [↑↓] [↑↓] [↑] [   ] [   ]", "Breaks Hund"),
               ("1s [↑]  2s [↑]", "Breaks Aufbau"),
               ("1s [↑↓]  2s [   ]  2p [↑]", "Breaks Aufbau")],
         why="Two arrows the same way in one box breaks Pauli. Pairing while a box in the same "
             "sublevel is empty breaks Hund. Skipping a lower sublevel breaks Aufbau."),
    dict(sec="S2.3", type=MC_T, pts=1,
         stem="Which list shows the sublevels in the order they fill?",
         choices=["3p, 4s, 3d, 4p", "3p, 3d, 4s, 4p", "4s, 4p, 3d, 4d", "3s, 3p, 3d, 4s"],
         answer="A",
         why="B and D fill by level number. C puts 4p before 3d."),
    dict(sec="S2.3", type=CLOZE, pts=1,
         stem="Titanium, Ti, has 22 electrons. Its longhand configuration is "
              "1s^2^ 2s^2^ 2p^6^ 3s^2^ 3p^6^ 4s^2^ [1]. Its noble-gas shorthand is [2].",
         blanks=[("3d^2^", ["3d^2^", "4d^2^", "4p^2^", "3d^4^"]),
                 ("[Ar] 4s^2^ 3d^2^", ["[Ar] 4s^2^ 3d^2^", "[Ne] 4s^2^ 3d^2^",
                                       "[Kr] 4s^2^ 3d^2^", "[Ar] 3d^4^"])],
         why="4d and 4p skip 3d. 3d^4^ forgets the 4s pair. [Ne] and [Kr] are the wrong core."),
    dict(sec="S2.3", type=MC_T, pts=1,
         stem="Which element has the configuration [Kr] 5s^2^ 4d^10^ 5p^5^?",
         choices=["Iodine, I", "Bromine, Br", "Xenon, Xe", "Tellurium, Te"],
         answer="A",
         why="Bromine ends 4p^5^. Xenon ends 5p^6^. Tellurium ends 5p^4^."),
    dict(sec="S2.3", type=MC_T, pts=1,
         stem="Which atom's noble-gas shorthand includes 4f^14^?",
         choices=["Barium, Ba", "Hafnium, Hf", "Zirconium, Zr", "Cesium, Cs"],
         answer="B",
         why="Barium stops at 6s^2^, just before 4f fills. Zirconium is in period 5 and uses 4d. "
             "Cesium stops at 6s^1^."),
    dict(sec="S2.3", type=MC_T, pts=1,
         stem="Copper is predicted to be [Ar] 4s^2^ 3d^9^, but it is actually "
              "[Ar] 4s^1^ 3d^10^. Why?",
         choices=["4s and 3d are close in energy, and a full 3d gives lower total energy.",
                  "4s is higher in energy than 4p, so it empties first.",
                  "Copper loses an electron to become an ion.",
                  "Pauli's rule allows only one electron in 4s."],
         answer="A",
         why="B is false. C confuses the atom with an ion. D misstates Pauli: 4s holds two."),
    # ---------------- S2.4 ----------------
    dict(sec="S2.4", type=MATCH, pts=2,
         stem="Match each atom to its number of valence electrons.",
         pairs=[("Sodium, Na", "1"), ("Carbon, C", "4"), ("Nitrogen, N", "5"),
                ("Chlorine, Cl", "7"), ("Argon, Ar", "8")],
         why="Students who count all electrons, or use the group number for groups 13 to 18 "
             "without subtracting 10, miss C, N, Cl and Ar."),
    dict(sec="S2.4", type=MC_T, pts=1,
         stem="Selenium is [Ar] 4s^2^ 3d^10^ 4p^4^. Which electrons are its valence electrons?",
         choices=["4s^2^ 4p^4^, so 6", "4p^4^ only, so 4", "3d^10^ 4p^4^, so 14", "All 34"],
         answer="A",
         why="B leaves out the 4s pair. C counts the last-filled sublevels. D counts every "
             "electron."),
    dict(sec="S2.4", type=CLASS, pts=2,
         stem="Sort each atom by the kind of ion it forms.",
         cols=["Forms a cation", "Forms an anion"],
         sort=[("Potassium, K", "Forms a cation"), ("Magnesium, Mg", "Forms a cation"),
               ("Aluminum, Al", "Forms a cation"), ("Sulfur, S", "Forms an anion"),
               ("Chlorine, Cl", "Forms an anion"), ("Phosphorus, P", "Forms an anion")],
         why="Metals lose electrons (cations); nonmetals gain them (anions)."),
    dict(sec="S2.4", type=CLOZE, pts=2,
         stem="Nitrogen has [1] valence electrons. It [2] [3] electrons to form the ion [4], "
              "which has the same electron arrangement as [5].",
         blanks=[("5", ["3", "5", "7"]), ("gains", ["gains", "loses"]), ("3", ["2", "3", "5"]),
                 ("N^3−^", ["N^3−^", "N^3+^", "N^5+^"]),
                 ("neon", ["neon", "argon", "helium"])],
         why="Losing 5 is possible on paper but costs far more energy than gaining 3. "
             "N^3−^ has 10 electrons, like neon."),
    dict(sec="S2.4", type=MC_T, pts=1,
         stem="Which ion has the same electron arrangement as argon?",
         choices=["K^+^", "Na^+^", "Ca^+^", "Cl"],
         answer="A",
         why="Na^+^ matches neon. Ca^+^ has 19 electrons. Cl is a neutral atom with 17."),
    dict(sec="S2.4", type=MC_T, pts=1,
         stem="When iron, [Ar] 4s^2^ 3d^6^, forms Fe^2+^, which electrons leave?",
         choices=["The two 4s electrons", "Two 3d electrons", "One 4s and one 3d electron",
                  "Two 4p electrons"],
         answer="A",
         why="B assumes the last in is the first out. Iron has no 4p electrons (D)."),
    dict(sec="S2.4", type=LABEL, pts=2, img="u02_test_periodic_blocks.png",
         stem="Label each block of the periodic table.",
         label="Put one drop-down inside each heavy outline: the left two columns, the middle "
               "ten, the right six, and the two rows below. Options in every drop-down: s block "
               "· p block · d block · f block. (H and He sit in their own small outlines; no "
               "drop-down needed there.)",
         answer="left = s block · middle = d block · right = p block · bottom rows = f block",
         why="Swapping d and f is the common slip."),
    # ---------------- S2.5 ----------------
    dict(sec="S2.5", type=CLASS, pts=2,
         stem="Sort each electron move by what happens to a photon.",
         cols=["Absorbs a photon", "Gives off a photon"],
         sort=[("n = 1 → 3", "Absorbs a photon"), ("n = 2 → 4", "Absorbs a photon"),
               ("n = 4 → 6", "Absorbs a photon"), ("n = 5 → 2", "Gives off a photon"),
               ("n = 3 → 1", "Gives off a photon"), ("n = 6 → 2", "Gives off a photon")],
         why="Up the ladder takes energy in; down the ladder gives it out."),
    dict(sec="S2.5", type=MC_T, pts=1, img="u02_test_energy_drops.png",
         stem="The diagram shows four electron drops, A to D. Which drop gives off the photon "
              "with the most energy?",
         choices=["A", "B", "C", "D"], answer="A",
         why="B crosses the most levels but A is the biggest gap. C and D are smaller gaps."),
    dict(sec="S2.5", type=CLOZE, pts=1,
         stem="Hydrogen gives off light at 410 nm and at 656 nm. The 410 nm light has a "
              "[1] wavelength, so its photon carries [2] energy and comes from a [3] drop.",
         blanks=[("shorter", ["shorter", "longer"]), ("more", ["more", "less"]),
                 ("larger", ["larger", "smaller"])],
         why="The trap is thinking a bigger number of nanometers means more energy."),
    dict(sec="S2.5", type=MATCH, pts=2,
         stem="Match each ion to its flame-test color.",
         pairs=[("Na^+^", "Yellow"), ("K^+^", "Lilac"), ("Cu^2+^", "Blue-green"),
                ("Ca^2+^", "Orange-red"), ("Ba^2+^", "Pale green")],
         why="Colors are from the S2.5 notes. Li^+^ and Sr^2+^ are left out because both are red."),
    dict(sec="S2.5", type=MC_T, pts=1,
         stem="Why do atoms in a flame give off light?",
         choices=["Heat excites electrons, and they give off photons as they drop back down.",
                  "Electrons give off photons as they move up to higher levels.",
                  "The atoms burn and the smoke glows.",
                  "Protons jump between energy levels."],
         answer="A",
         why="B has the direction backwards. C and D are not about electrons."),
]


def check_edulastic():
    assert len(ITEMS) == 25, len(ITEMS)
    for it in ITEMS:
        if it["type"] == MC_T:
            assert len(it["choices"]) == 4 and len(set(it["choices"])) == 4
            assert it["answer"] in "ABCD"
        if it["type"] == CLOZE:
            for ans, opts in it["blanks"]:
                assert ans in opts, (ans, opts)
    # Facts the stems rely on
    c, r = shorthand(53); assert (NOBLE_SYM[c], flat(r)) == ("Kr", "5s2 4d10 5p5")
    assert any(l == "f" for _, l, _ in shorthand(72)[1])
    assert not any(l == "f" for _, l, _ in shorthand(56)[1] + shorthand(40)[1])
    assert flat(config(22)).endswith("4s2 3d2")
    assert flat(config(29)).endswith("4s1 3d10")
    assert flat(config(16)) == "1s2 2s2 2p6 3s2 3p4"
    assert [valence(z)[1] for z in (11, 6, 7, 17, 18)] == [1, 4, 5, 7, 8]
    assert valence(34)[1] == 6 and ion_charge(7) == "3−"
    pts = sum(it["pts"] for it in ITEMS)
    print(f"Edulastic items checked: 25 items, {pts} points")
    return pts


def build_edulastic(out, imgdir, pts_total):
    pal = Palette(COURSE)
    ctx = Ctx(pal, True)
    ctx.key = False                       # no blanks are drawn on this sheet
    doc, s = new_doc()
    header(doc, pal, "UNIT 2 TEST  ·  DAY 1  ·  EDULASTIC BUILD SHEET  ·  TEACHER ONLY",
           "Unit 2 Test: Conceptual (Edulastic)", fields=False)
    gap(doc, 3)
    c = one_cell(doc); borders(c, pal.display, sz=18, edges=("left",))
    p = cellfmt(c, pal, line=False, top=20, bottom=20)
    _run(p, f"25 items, {pts_total} points. Each item gives the Edulastic question type to pick, "
            "then what to type, in the order the editor asks for it. The correct answer is in "
            "bold. Images are in the edulastic folder next to this file. Turn on question and "
            "answer shuffling in the test settings in place of separate versions.", 9.5,
         color=pal.ink)

    current = None
    for i, it in enumerate(ITEMS, start=1):
        if it["sec"] != current:
            current = it["sec"]
            section_bar(doc, pal, {"S2.1": "S2.1  ·  THE BOHR MODEL",
                                   "S2.2": "S2.2  ·  ENERGY LEVELS, SUBLEVELS & ORBITALS",
                                   "S2.3": "S2.3  ·  ELECTRON CONFIGURATIONS",
                                   "S2.4": "S2.4  ·  VALENCE ELECTRONS & IONS",
                                   "S2.5": "S2.5  ·  ATOMIC & ELECTRON SPECTRA"}[current])
        c = one_cell(doc); borders(c, pal.hair, sz=4, edges=("bottom",))
        no_split(doc.tables[-1].rows[0])          # an item never breaks across pages
        cell_margins(c, top=50, bottom=50, left=40, right=40)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3), WD_TAB_ALIGNMENT.RIGHT)
        _run(p, f"{i}   ", 11, bold=True, color=pal.accent)
        _run(p, it["type"].upper(), 8.5, bold=True, color=pal.label)
        _run(p, f"\t{it['pts']} pt" + ("s" if it["pts"] > 1 else ""), 8.5, bold=True,
             color=pal.label)

        def line(text, size=10, bold=False, color=None, indent=0.25, after=1):
            q = c.add_paragraph(); q.paragraph_format.space_after = Pt(after)
            q.paragraph_format.left_indent = Inches(indent)
            rich(q, text, size, ctx, bold=bold, color=color)
            return q

        line(it["stem"], 10.5, after=3)
        if it.get("img"):
            q = c.add_paragraph(); q.paragraph_format.left_indent = Inches(0.25)
            q.paragraph_format.space_after = Pt(3)
            q.add_run().add_picture(os.path.join(imgdir, it["img"]),
                                    width=Inches(5.6 if it["type"] == LABEL else 4.6))
            line(f"Image file: edulastic/{it['img']}", 8.5, color=pal.label)
        t = it["type"]
        if t == MC_T:
            for k, ch in enumerate(it["choices"]):
                ok = "ABCD"[k] == it["answer"]
                line(f"{'ABCD'[k]}.  {ch}", 10, bold=ok, color=INK["box"] if ok else None)
        elif t == CLOZE:
            for k, (ans, opts) in enumerate(it["blanks"], start=1):
                q = line(f"[{k}]  ", 10, bold=True)
                for j, o in enumerate(opts):
                    rich(q, o, 10, ctx, bold=(o == ans), color=INK["box"] if o == ans else None)
                    if j < len(opts) - 1:
                        _run(q, "  ·  ", 10, color=pal.label)
        elif t == MATRIX:
            line("Rows: " + "  ·  ".join(it["rows"]), 10)
            line("Columns: " + "  ·  ".join(it["cols"]), 10)
            line("Correct: " + it["answer"], 10, bold=True, color=INK["box"])
        elif t == CLASS:
            line("Column headings: " + "  ·  ".join(it["cols"]), 10)
            for col in it["cols"]:
                items = [x for x, cc in it["sort"] if cc == col]
                q = line(f"{col}:  ", 10, bold=True, color=INK["box"])
                rich(q, "   |   ".join(items), 10, ctx)
            line("Edulastic shuffles the cards; type them in any order.", 8.5, color=pal.label)
        elif t == MATCH:
            for a, b in it["pairs"]:
                q = line(f"{a}   →   ", 10)
                rich(q, b, 10, ctx, bold=True, color=INK["box"])
        elif t == LABEL:
            line(it["label"], 9.5)
            line("Correct: " + it["answer"], 10, bold=True, color=INK["box"])
        line("Why the wrong ones are wrong:  " + it["why"], 8.5, color=pal.label, after=0)
        pin_tail(c)

    running_footer(s, f"SHULL SCIENCE          {UNIT}  ·  UNIT TEST DAY 1  ·  TEACHER ONLY",
                   pal, with_page_numbers=True)
    trim_tail(doc)
    doc.save(out)


def main():
    outdir = sys.argv[1] if len(sys.argv) > 1 else "."
    os.makedirs(outdir, exist_ok=True)
    Q.TILE_DIR = os.path.join(outdir, ".dots"); os.makedirs(Q.TILE_DIR, exist_ok=True)
    imgdir = os.path.join(outdir, "edulastic"); os.makedirs(imgdir, exist_ok=True)
    orbital_blank_png(os.path.join(imgdir, "u02_test_sulfur_orbitals.png"),
                      [("1s", 1), ("2s", 1), ("2p", 3), ("3s", 1), ("3p", 3)])
    periodic_blocks_png(os.path.join(imgdir, "u02_test_periodic_blocks.png"))
    ladder_drops_png(os.path.join(imgdir, "u02_test_energy_drops.png"))

    total = check_written()
    for key in (False, True):
        name = document_name("Test", ["2.1", "2.5"], 2, key=key, descriptor="Day 2 Written")
        build_written(os.path.join(outdir, name), key, total)
        print("wrote", name)
    pts = check_edulastic()
    name = document_name("Test", ["2.1", "2.5"], 2, descriptor="Day 1 Edulastic Build Sheet")
    build_edulastic(os.path.join(outdir, name), imgdir, pts)
    print("wrote", name)


if __name__ == "__main__":
    main()
