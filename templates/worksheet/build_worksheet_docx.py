#!/usr/bin/env python3
"""Build a SHULL worksheet / practice set as .docx from a JSON spec.

    python3 templates/worksheet/build_worksheet_docx.py specs/<spec>.json out.docx

One builder, three course profiles. The three courses do not want the same sheet, and
pretending they do is how a Geology worksheet ends up with a work box.

  Chemistry  Conceptual and math. Every math question gets an open box to solve in,
             carrying a SHOW WORK HERE watermark. Equations sit at the top of the page.
             Diagrams and charts to complete are first-class.
  Physics    NO work areas at all. Students work in their Hayden-McNeil carbonless lab
             notebooks, so work boxes and answer rules are refused outright. The page
             carries the equations and a prior-knowledge recall block instead, and the
             questions ramp: warm-up, practice, challenge, then one multi-topic problem
             that reaches back into earlier units.
  Geology    No math. Visual and expressive: diagrams to label, panels to draw in,
             cards to cut, slots to glue them into. A Geology sheet with an equation
             bar or a work box is refused.

The profile is enforced here, at build time. A rule that only lives in a document is
not a mechanism.

A packet holds one or more sections, each starting a new page - the shape of his own
PHYS U01 practice sets, which run one section per page across the whole unit.

Colour comes from brand/tokens.json. No hex is typed in this file.
"""
import json, os, sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from _shull_name import document_name, short_title   # noqa: E402
from _shull_docx import (          # noqa: E402
    COURSE_CODE, Palette, debullet, known_sections, unit_title,
    borders, para, run, check_item, rule_lines, fix_widths, one_cell, no_split,
    equation_bar, work_box, given_need, diagram_block, page_setup, running_footer,
    text_width_in, cell_margins, gap, shade, unpad_cell, T, trim_tail, rich,
)

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))

TEXT_W_IN = 7.50
NUM_W_IN = 0.30                                  # the question number's own column
BODY_W_IN = TEXT_W_IN - NUM_W_IN

# The ramp, per course. His own PHYS U01 practice sets run 2 warm-up, 2 practice,
# 1 challenge, 1 multi-topic per section - six questions, easy to start and finishing on
# a problem that reaches into another unit. The inherited studio spec said eight
# (2/3/2/1); it is marked INHERITED in docs/LEGACY_SKILL_AND_DESIGN_AUDIT.md, meaning
# nobody confirmed it, and his shipped packet is the better evidence. SHULL-CHG-0019
# records the conflict.
TIERS = ("warm-up", "practice", "challenge", "multi-topic")
TIER_PT = 7                     # the tag is metadata, set below the body size
TIER_TRACK_PT = 1.3             # w:spacing val 26, in points per character
TIER_GAP_IN = 0.10              # the gap between the tag and the prompt it labels
# Exact counts are enforced only where there is evidence for them. Physics has a
# shipped packet to read; Chemistry does not, so it gets the ORDER of the ramp enforced
# and picks its own counts. Inventing a Chemistry count and calling it locked is how
# the 8-question figure got into the system in the first place.
RAMP = {
    "physics":   {"warm-up": 2, "practice": 2, "challenge": 1, "multi-topic": 1},
    "chemistry": None,
    "geology":   None,          # a Geology sheet is an activity, not a question ladder
}


def tier_col_in():
    """Width of the tier column, measured from the longest tier name.

    The tags used to run inline with the prompt, which left every question starting at
    a different place and every wrapped line running back under the tag. Giving the tag
    its own column fixes both - but only if the column is wide enough for the longest
    name, or "MULTI-TOPIC" wraps to two lines and looks worse than what it replaced.
    So it is measured off the shipped Archivo files rather than picked, and it follows
    the tier names if they are ever renamed.
    """
    widest = max(text_width_in(t.upper(), TIER_PT, bold=True)
                 + TIER_TRACK_PT * len(t) / 72.0 for t in TIERS)
    # The tag is set right, so it needs no slack on its left - only the gap to the
    # prompt on its right. The cell's own margins are zeroed to match, because every
    # hundredth of an inch here comes straight off the prompt column, and on a
    # one-page-per-section sheet that is the difference between four pages and five.
    return round(widest + TIER_GAP_IN, 2)
WORK_MIN_IN = 1.6
WATERMARK = "SHOW WORK HERE"
CARD_PAD_IN = 0.28              # one nesting of cell margins
SHORT_PROMPT_CH = 118           # two such questions sit side by side


def tier_colour(tier, pal):
    """Every tier tag is the muted label grey. SHULL-CHG-0055.

    The word carries the tier and the order of the questions carries the difficulty. The
    old ramp (grey, amber, red, ink - SHULL-CHG-0020) put an amber PRACTICE tag beside
    Chemistry's Energy Orange Deep headings, the same grey once photocopied (SHULL-CHG-0054),
    and spent semantic red on CHALLENGE in a room where red means hazard.
    """
    return pal.label


def is_short(q):
    """A question that can share a row with its neighbour: nothing to draw, nothing to
    solve in, no lettered parts, and a prompt that fits two lines of half a page."""
    return not (q.get("diagram") or q.get("draw") or q.get("math")
                or q.get("answerLines") or q.get("parts") or q.get("given")
                or q.get("need")) and len(q.get("prompt", "")) <= SHORT_PROMPT_CH


# --------------------------------------------------------------------------- blocks

def page_break(doc):
    """Start a new page without spending a line to do it.

    A normal paragraph carrying a page break is full body height, so when a section
    ends flush with the bottom of its page that paragraph does not fit - it moves to
    the next page, and only then breaks, leaving a blank page carrying nothing but the
    footer. Pinning the line to 1pt exact makes the break paragraph effectively
    dimensionless, so it always fits where it is written.
    """
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "0")
    sp.set(qn("w:line"), "20"); sp.set(qn("w:lineRule"), "exact")
    pPr.append(sp)
    r = p.add_run()
    r.font.size = Pt(1)
    r.add_break(WD_BREAK.PAGE)
    return p


def label(cell, text, pal, size=7.5, first=False):
    return para(cell, text, size, bold=True, color=pal.accent, caps_track=True, first=first)


def arrow(doc, pal):
    """The reading-order arrow between rows of cards. It is on his sheet and it does
    real work - without it the six cards read as a grid, not a sequence."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run(p, "↓", 11, bold=True, color=pal.display)
    return p


def grid(doc, n, across, height_in, pal, gutter=0.18, arrows=True):
    """A grid of equal cells, `across` per row, each row held together across a break.

    One table per ROW, not one table for the whole grid: cantSplit holds a row and does
    nothing across two, and a card sliced in half by a page break cannot be cut out.
    Columns are separated by a narrow spacer column rather than by cell margins, so the
    dashed cut lines of two neighbouring cards never touch.
    """
    cols = across * 2 - 1
    cell_w = round((TEXT_W_IN - gutter * (across - 1)) / across, 3)
    widths = []
    for i in range(across):
        widths.append(cell_w)
        if i < across - 1:
            widths.append(gutter)
    rows = (n + across - 1) // across
    out = []
    for r in range(rows):
        t = doc.add_table(rows=1, cols=cols)
        fix_widths(t, widths)
        no_split(t.rows[0], height_in)
        for i in range(across):
            if r * across + i < n:
                out.append(t.rows[0].cells[i * 2])
        if r < rows - 1 and arrows:
            arrow(doc, pal)
        else:
            gap(doc, 4)   # cards need air between rows even without an arrow
    return out


def cards_block(doc, block, pal):
    """Cards to cut out. Dashed border, title, description, and a box to draw in.

    Straight off his Nebular Theory sheet, which is the one he named as the look he
    wants: "cut it out, and color or sketch a small illustration for it."
    """
    items = block["cards"]
    across = int(block.get("across", 2))
    h = float(block.get("cardHeightIn", 1.65))
    draw_h = float(block.get("drawHeightIn", 0.95))
    cells = grid(doc, len(items), across, h, pal, arrows=block.get("arrows", True))
    for cell, card in zip(cells, items):
        borders(cell, pal.label, sz=6, val="dashed")
        label(cell, block.get("cutLabel", "✂  CUT ALONG THE DASHED LINE"), pal,
              size=6.5, first=True)
        para(cell, card["title"], 11, bold=True, color=pal.accent)
        para(cell, card["text"], 9)
        label(cell, block.get("drawLabel", "COLOR OR SKETCH THIS STAGE:"), pal, size=6.5)
        t = cell.add_table(rows=1, cols=1)
        fix_widths(t, [round(TEXT_W_IN / across - 0.5, 2)])
        no_split(t.rows[0], draw_h)
        borders(t.rows[0].cells[0], pal.ink, sz=4)
    return cells


def slots_block(doc, block, pal):
    """Numbered slots the cut cards get glued into, sized to match the cards."""
    n = int(block.get("slots", 6))
    across = int(block.get("across", 2))
    h = float(block.get("slotHeightIn", 1.65))
    inner_h = float(block.get("innerHeightIn", 0.80))
    cells = grid(doc, n, across, h, pal, arrows=block.get("arrows", True))
    for i, cell in enumerate(cells, 1):
        borders(cell, pal.ink, sz=8)
        label(cell, f"{block.get('slotLabel', 'GLUE HERE — SLOT')} {i}", pal,
              size=6.5, first=True)
        t = cell.add_table(rows=1, cols=1)
        fix_widths(t, [round(TEXT_W_IN / across - 0.5, 2)])
        no_split(t.rows[0], inner_h)
        borders(t.rows[0].cells[0], pal.label, sz=4, val="dashed")
    return cells


def draw_block(cell, block, pal, width):
    """Blank panels to draw in, captioned underneath. The lighter cousin of `cards` -
    used when there is nothing to cut, only something to show."""
    n = int(block.get("panels", 4))
    across = int(block.get("across", 4))
    col_w = round(width / across, 3)
    panel_h = float(block.get("panelHeightIn", 1.7))
    label(cell, block.get("label", "DRAW IT"), pal)
    if block.get("instruction"):
        para(cell, block["instruction"], 9.5)
    captions = block.get("captions") or []
    rows = (n + across - 1) // across
    t = cell.add_table(rows=rows * 2, cols=across)
    fix_widths(t, [col_w] * across)
    for i in range(n):
        r, c = divmod(i, across)
        panel = t.rows[r * 2].cells[c]
        borders(panel, pal.ink, sz=4)
        no_split(t.rows[r * 2], panel_h)
        # Just the number. No box, no square, no circle. SHULL-CHG-0015.
        para(panel, str(i + 1), 9, bold=True, color=pal.ink, first=True)
        under = t.rows[r * 2 + 1].cells[c]
        borders(under, pal.hair, sz=4, edges=("bottom",))
        cap = captions[i] if i < len(captions) else block.get("captionPrompt", "")
        para(under, cap, 8, color=pal.label, first=True)
    return t


def sort_block(doc, block, pal):
    """Put things in order, in place - no scissors.

    The same learning as the cut-and-glue timeline, on one page and in ten minutes:
    the items are printed scrambled with a blank beside each, and the student writes
    the order. Cut-and-glue is worth the two pages and the scissors sometimes; it is
    not worth them every time, and this is the version for the rest of the time.
    """
    c = one_cell(doc)
    borders(c, pal.hair, sz=4)
    cell_margins(c, top=45, bottom=45, left=130, right=130)
    label(c, block.get("label", "PUT THESE IN ORDER"), pal, first=True)
    if block.get("instruction"):
        para(c, block["instruction"], 9.5)
    items = block["items"]
    t = c.add_table(rows=len(items), cols=2)
    fix_widths(t, [0.55, TEXT_W_IN - 0.55 - 0.30])
    for r, item in enumerate(items):
        box, txt = t.rows[r].cells
        no_split(t.rows[r])
        borders(box, pal.ink, sz=4)
        cell_margins(box, top=40, bottom=40, left=60, right=60)
        para(box, "", 11, first=True)
        borders(txt, pal.hair, sz=4, edges=("bottom",))
        cell_margins(txt, top=40, bottom=40, left=130, right=60)
        para(txt, item, 9.5, color=pal.ink, first=True)
    unpad_cell(c)
    return c


def match_block(doc, block, pal):
    """Matching: terms on the left with a blank, descriptions on the right.

    The descriptions are deliberately NOT in the same order as the terms; the spec
    supplies them already shuffled, because a builder that shuffles them would produce
    a different sheet on every build and no answer key would survive.
    """
    c = one_cell(doc)
    borders(c, pal.hair, sz=4)
    cell_margins(c, top=45, bottom=45, left=130, right=130)
    label(c, block.get("label", "MATCHING"), pal, first=True)
    para(c, block.get("instruction",
                      "Write the letter of the correct description beside each term."), 9.5)
    terms, descs = block["terms"], block["descriptions"]
    rows = max(len(terms), len(descs))
    t = c.add_table(rows=rows, cols=4)
    fix_widths(t, [0.42, 2.35, 0.34, TEXT_W_IN - 3.41])
    for r in range(rows):
        blank, term, letter, desc = t.rows[r].cells
        no_split(t.rows[r])
        if r < len(terms):
            borders(blank, pal.ink, sz=4, edges=("bottom",))
            cell_margins(blank, top=40, bottom=40, left=0, right=90)
            para(blank, "", 10, first=True)
            para(term, terms[r], 9.5, color=pal.ink, first=True)
        if r < len(descs):
            # Just the letter. No box, no circle. SHULL-CHG-0015.
            para(letter, f"{chr(97 + r)}.", 9.5, bold=True, color=pal.accent, first=True)
            para(desc, descs[r], 9.5, color=pal.ink, first=True)
    unpad_cell(c)
    return c


def reflection_block(doc, block, pal):
    """The written close on his Geology sheet: italic prompts, a full-width rule under
    each. Ruled lines are for prose - this is the one place a Geology sheet writes."""
    c = one_cell(doc)
    for i, q in enumerate(block["prompts"], 1):
        # A prompt and the lines it is answered on are one thing. Left to flow, the
        # question ended a page and its answer lines started the next one - the student
        # turns over to two rules belonging to nothing.
        t = c.add_table(rows=1, cols=1)
        fix_widths(t, [TEXT_W_IN - 0.30])
        no_split(t.rows[0])
        cc = t.rows[0].cells[0]
        cell_margins(cc, top=0, bottom=0, left=0, right=0)
        if i == 1:
            # The heading goes INSIDE the first non-splitting row. Kept outside it, it
            # ended a page on its own with the question it introduces on the next one -
            # the same stranding the diagram block hit, one level up.
            label(cc, block.get("label", "REFLECTION"), pal, first=True)
            p = cc.add_paragraph()
        else:
            p = cc.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        run(p, f"{i}. ", 9.5, bold=True, color=pal.ink)
        run(p, q, 9.5, italic=True, color=pal.ink)
        rule_lines(cc, int(block.get("linesEach", 2)), pal.ink)
    unpad_cell(c)
    return c


# ------------------------------------------------------------- project activities
# The blocks a project handout needs and a question sheet does not: the layout outline a
# student builds from, the requirement list and option table, the format choice, the plan,
# and the rubric. Same primitives, same palette, same ink rule (borders, never fills).

FACTS_SZ = 18                  # solid course-colour border: a fact zone
PICS_SZ = 12                   # dashed ink border: a picture zone


def pin_after(cell, pt):
    """python-docx follows every table it nests in a cell with a full-height empty
    paragraph. Between two zones of an outline that is a quarter inch of nothing, so the
    paragraph is pinned to an exact `pt` instead - the gap it was standing in for."""
    p = cell.paragraphs[-1]
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "0")
    sp.set(qn("w:line"), str(int(pt * 20))); sp.set(qn("w:lineRule"), "exact")
    pPr.append(sp)
    rPr = OxmlElement("w:rPr")
    for name in ("w:sz", "w:szCs"):
        e = OxmlElement(name); e.set(qn("w:val"), str(max(2, int(pt * 2) - 2))); rPr.append(e)
    pPr.append(rPr)


def zone_border(cell, kind, pal):
    """Geology's two colours photocopy to the same grey, so an outline's zone kinds
    differ by border STYLE and by label, not by hue (courses/geology/DECISIONS.md)."""
    if kind == "facts":
        borders(cell, pal.display, sz=FACTS_SZ)
    else:
        borders(cell, pal.ink, sz=PICS_SZ, val="dashed")


def zone_label(cell, lab, hint, pal, kind, first=True, hint_below=False):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = run(p, lab, 8, bold=True, color=pal.accent if kind == "facts" else pal.ink)
    el = OxmlElement("w:spacing"); el.set(qn("w:val"), "26")
    r._element.get_or_add_rPr().append(el)
    if hint:
        if hint_below:
            p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(1)
        else:
            run(p, "   ", 8, color=pal.label)
        run(p, hint, 8, color=pal.label)
    return p


def outline_block(doc, block, pal):
    """A labelled blank layout: where each part of the product goes. One non-splitting
    row, so the outline never breaks across a page."""
    t = doc.add_table(rows=1, cols=1)
    fix_widths(t, [TEXT_W_IN])
    no_split(t.rows[0])
    c = t.rows[0].cells[0]
    cell_margins(c, top=0, bottom=0, left=0, right=0)
    label(c, block.get("label", "OUTLINE"), pal, first=True)
    w = TEXT_W_IN
    for z in block["zones"]:
        if z.get("row"):
            parts = z["row"]
            gut = 0.12
            ws = []
            for k, zz in enumerate(parts):
                ws.append(round((w - gut * (len(parts) - 1)) * zz["share"], 3))
                if k < len(parts) - 1:
                    ws.append(gut)
            zt = c.add_table(rows=1, cols=len(ws))
            fix_widths(zt, ws)
            no_split(zt.rows[0], float(z.get("heightIn", 1.5)))
            for k, zz in enumerate(parts):
                zc = zt.rows[0].cells[k * 2]
                zone_border(zc, zz["kind"], pal)
                cell_margins(zc, top=60, bottom=60, left=110, right=110)
                zone_label(zc, zz["label"], zz.get("hint"), pal, zz["kind"], hint_below=True)
                if zz.get("legend"):
                    # The fossil-site dot, pushed to the bottom of the zone. U+25CF is not
                    # in Archivo, so the dot is a bullet set large.
                    used = 8 * 1.3 + 8 * 1.3 * 2
                    p = zc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    p.paragraph_format.space_before = Pt(max(0, float(z.get("heightIn", 1.5)) * 72
                                                            - used - 16 - 14))
                    run(p, "•", 14, bold=True, color=pal.ink)
                    run(p, "  " + zz["legend"], 8.5, color=pal.ink)
        else:
            zt = c.add_table(rows=1, cols=1)
            fix_widths(zt, [w])
            no_split(zt.rows[0], z.get("heightIn"))
            zc = zt.rows[0].cells[0]
            zone_border(zc, z["kind"], pal)
            cell_margins(zc, top=60, bottom=60, left=110, right=110)
            zone_label(zc, z["label"], z.get("hint"), pal, z["kind"])
            if z.get("periods"):
                ps = z["periods"]
                span = sum(p["start"] - p["end"] for p in ps)
                inner = w - 0.16
                pt = zc.add_table(rows=1, cols=len(ps))
                # Cell widths follow how long each period lasted.
                fix_widths(pt, [round(inner * (p["start"] - p["end"]) / span, 3) for p in ps])
                for pc, p in zip(pt.rows[0].cells, ps):
                    borders(pc, pal.ink, sz=6)
                    cell_margins(pc, top=40, bottom=40, left=90, right=60)
                    pp = pc.paragraphs[0]
                    r = run(pp, p["name"].upper(), 8, bold=True, color=pal.ink)
                    el = OxmlElement("w:spacing"); el.set(qn("w:val"), "20")
                    r._element.get_or_add_rPr().append(el)
                    run(pp, f"   {p['start']} to {p['end']} mya", 8, color=pal.label)
                unpad_cell(zc)
        pin_after(c, 6)
    leg = block.get("legend")
    if leg:
        lt = c.add_table(rows=1, cols=5)
        fix_widths(lt, [0.30, 1.75, 0.30, 2.05, w - 4.40])
        trPr = lt.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:trHeight"); th.set(qn("w:val"), str(int(0.16 * 1440)))
        th.set(qn("w:hRule"), "exact"); trPr.append(th)
        for k, kind in ((0, "facts"), (2, "pics")):
            sw = lt.rows[0].cells[k]
            zone_border(sw, kind, pal)
            cell_margins(sw, top=0, bottom=0, left=0, right=0)
            para(sw, "", 2, first=True)
        for k, txt in ((1, leg[0]), (3, leg[1]), (4, leg[2])):
            tc = lt.rows[0].cells[k]
            cell_margins(tc, top=0, bottom=0, left=90, right=40)
            para(tc, txt, 8.5, color=pal.ink, first=True)
    unpad_cell(c)
    return c


def list_block(doc, block, pal):
    """A numbered requirement list, `across` columns. Just the number, no box
    (SHULL-CHG-0015); a bold lead-in, then the requirement."""
    c = one_cell(doc, protect=True)
    cell_margins(c, top=0, bottom=0, left=0, right=0)
    label(c, block["label"], pal, first=True)
    items = block["items"]
    across = int(block.get("across", 1))
    per = (len(items) + across - 1) // across
    gut = 0.25
    col = round((TEXT_W_IN - gut * (across - 1)) / across, 3)
    ws = []
    for k in range(across):
        ws.append(col)
        if k < across - 1:
            ws.append(gut)
    t = c.add_table(rows=1, cols=len(ws))
    fix_widths(t, ws)
    for k in range(across):
        cc = t.rows[0].cells[k * 2]
        cell_margins(cc, top=0, bottom=0, left=0, right=0)
        for j, (lead, txt) in enumerate(items[k * per:(k + 1) * per]):
            n = k * per + j + 1
            p = cc.paragraphs[0] if j == 0 else cc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.22))
            run(p, f"{n}\t", 9.5, bold=True, color=pal.accent)
            rich(p, f"<b>{lead}</b> {txt}", 9.5, color=pal.ink)
    unpad_cell(c)
    return c


def table_block(doc, block, pal, width=TEXT_W_IN):
    """A reference table: caps header, hairline rows, first column bold. Rows may be
    `{"group": "..."}` to open a labelled group. Text may carry <b> and <i>."""
    c = one_cell(doc, width)
    cell_margins(c, top=0, bottom=0, left=0, right=0)
    first = True
    if block.get("label"):
        label(c, block["label"], pal, first=True); first = False
    if block.get("lead"):
        p = c.paragraphs[0] if first else c.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        rich(p, block["lead"], float(block.get("size", 9.5)), color=pal.ink); first = False
    heads = block["headers"]
    ws = [round(width * f, 3) for f in block["widths"]]
    size = float(block.get("size", 9.5))
    rows = block["rows"]
    t = c.add_table(rows=1 + len(rows), cols=len(heads))
    fix_widths(t, ws)
    hr = t.rows[0]
    hr._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))   # repeats on a new page
    for hc, h in zip(hr.cells, heads):
        borders(hc, pal.ink, sz=8, edges=("bottom",))
        cell_margins(hc, top=20, bottom=30, left=60, right=60)
        para(hc, h, 8, bold=True, color=pal.label, caps_track=True, first=True)
    for ri, row in enumerate(rows, 1):
        tr = t.rows[ri]
        no_split(tr)
        if isinstance(row, dict):
            merged = tr.cells[0].merge(tr.cells[-1])
            borders(merged, pal.display, sz=8, edges=("top",))
            borders(merged, pal.hair, sz=4, edges=("bottom",))
            cell_margins(merged, top=60, bottom=30, left=60, right=60)
            para(merged, row["group"], 8, bold=True, color=pal.accent, caps_track=True,
                 first=True)
            continue
        for k, (tc, val) in enumerate(zip(tr.cells, row)):
            borders(tc, pal.hair, sz=4, edges=("bottom",))
            cell_margins(tc, top=30, bottom=30, left=60, right=60)
            p = tc.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            rich(p, val, size, color=pal.ink, bold=(k == 0 and block.get("boldFirst", True)))
    unpad_cell(c)
    return c


def notes_block(doc, block, pal, width=TEXT_W_IN):
    """Prose lines: a label and bulleted items, or one plain line (`"bullets": false`)."""
    c = one_cell(doc, width)
    cell_margins(c, top=0, bottom=0, left=0, right=0)
    first = True
    if block.get("label"):
        label(c, block["label"], pal, first=True); first = False
    bullets = block.get("bullets", True)
    size = float(block.get("size", 9.5))
    for x in block["items"]:
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(3)
        if bullets:
            p.paragraph_format.left_indent = Inches(0.16)
            p.paragraph_format.first_line_indent = Inches(-0.16)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.16))
            run(p, "•\t", size, bold=True, color=pal.display)
        rich(p, x, size, color=pal.ink)
    return c


def formats_block(doc, block, pal):
    """Two ways to do the same project, side by side, graded the same."""
    c = one_cell(doc, protect=True)
    cell_margins(c, top=0, bottom=0, left=0, right=0)
    label(c, block["label"], pal, first=True)
    if block.get("note"):
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        rich(p, block["note"], 9.5, color=pal.ink)
    boxes = block["boxes"]
    gut = 0.15
    ws = []
    for k, b in enumerate(boxes):
        ws.append(round((TEXT_W_IN - gut * (len(boxes) - 1)) * b["share"], 3))
        if k < len(boxes) - 1:
            ws.append(gut)
    t = c.add_table(rows=1, cols=len(ws))
    fix_widths(t, ws)
    for k, b in enumerate(boxes):
        bc = t.rows[0].cells[k * 2]
        borders(bc, pal.label, sz=6)
        cell_margins(bc, top=60, bottom=60, left=120, right=120)
        p = bc.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
        r = run(p, b["title"], 8, bold=True, color=pal.accent)
        el = OxmlElement("w:spacing"); el.set(qn("w:val"), "26")
        r._element.get_or_add_rPr().append(el)
        if b.get("lead"):
            run(p, "   ", 9)
            rich(p, b["lead"], 9, color=pal.ink)
        for x in b.get("lines", []):
            q = bc.add_paragraph(); q.paragraph_format.space_after = Pt(3)
            rich(q, x, 9, color=pal.ink)
        for x in b.get("bullets", []):
            q = bc.add_paragraph(); q.paragraph_format.space_after = Pt(2)
            q.paragraph_format.left_indent = Inches(0.16)
            q.paragraph_format.first_line_indent = Inches(-0.16)
            q.paragraph_format.tab_stops.add_tab_stop(Inches(0.16))
            run(q, "•\t", 9, bold=True, color=pal.display)
            rich(q, x, 9, color=pal.ink)
    unpad_cell(c)
    return c


def plan_block(doc, block, pal):
    """Labelled write-on lines, one or two fields per line. Ruled lines are for prose,
    and a plan is prose. A field with flex 0 is a label only (a circle-one choice)."""
    c = one_cell(doc, protect=True)
    cell_margins(c, top=0, bottom=0, left=0, right=0)
    label(c, block["label"], pal, first=True)
    for row in block["rows"]:
        lab_ws = [round(text_width_in(k.replace("&nbsp;", " "), 9.5, bold=True) + 0.14, 2)
                  for k, _ in row]
        free = TEXT_W_IN - sum(lab_ws) - 0.2 * (len(row) - 1)
        flex = sum(n for _, n in row) or 1
        ws, kinds = [], []
        for (k, n), lw in zip(row, lab_ws):
            ws.append(lw); kinds.append(("lab", k))
            if n:
                ws.append(round(free * n / flex, 3)); kinds.append(("line", None))
            if (k, n) != row[-1]:
                ws.append(0.2); kinds.append(("gap", None))
        if sum(ws) < TEXT_W_IN - 0.01:
            ws[-1] = round(ws[-1] + TEXT_W_IN - sum(ws), 3)
        t = c.add_table(rows=1, cols=len(ws))
        fix_widths(t, ws)
        no_split(t.rows[0], 0.28)
        for tc, (kind, k) in zip(t.rows[0].cells, kinds):
            cell_margins(tc, top=0, bottom=20, left=0, right=0)
            tc.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
            if kind == "lab":
                rich(tc.paragraphs[0], k, 9.5, bold=True, color=pal.ink)
            elif kind == "line":
                borders(tc, pal.ink, sz=4, edges=("bottom",))
            tc.paragraphs[0].paragraph_format.space_after = Pt(0)
        pin_after(c, 3)
    unpad_cell(c)
    return c


def rubric_total(spec):
    return sum(int(r["points"]) for r in spec["rubric"]["rows"])


def score_box(cell, pal, w=0.42, h=0.26):
    bt = cell.add_table(rows=1, cols=1)
    fix_widths(bt, [w])
    no_split(bt.rows[0], h)
    borders(bt.rows[0].cells[0], pal.ink, sz=6)
    unpad_cell(cell)


def total_line(doc, spec, pal):
    t = doc.add_table(rows=1, cols=3)
    fix_widths(t, [TEXT_W_IN - 1.40, 0.80, 0.60])
    a, b, d = t.rows[0].cells
    for x in (a, b, d):
        cell_margins(x, top=40, bottom=0, left=60, right=60)
        x.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = para(a, "TOTAL", 9, bold=True, color=pal.ink, caps_track=True, first=True)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    score_box(b, pal, w=0.70, h=0.30)
    para(d, f"/ {rubric_total(spec)}", 10, bold=True, color=pal.ink, first=True)


def rubric_block(doc, block, pal, spec, sheet=False):
    """The rubric, written once in the spec and drawn two ways: full level descriptors on
    the handout, and the grading sheet's one-line "look for" (the full-credit text, word
    for word) with a comment column. The handout and the sheet cannot drift."""
    rub = spec["rubric"]
    c = one_cell(doc)
    cell_margins(c, top=0, bottom=0, left=0, right=0)
    label(c, block.get("label", f"RUBRIC · {rubric_total(spec)} POINTS"), pal, first=True)
    if sheet:
        heads = ["PART", "LOOK FOR", "FULL / PARTIAL / NONE", "SCORE", "COMMENT"]
        ws = [1.15, 2.85, 1.10, 0.62, 1.78]
    else:
        heads = ["PART", "FULL CREDIT", "PARTIAL", "LITTLE OR NONE", "SCORE"]
        ws = [1.15, 2.75, 1.60, 1.38, 0.62]
    t = c.add_table(rows=1 + len(rub["rows"]), cols=5)
    fix_widths(t, ws)
    for hc, h in zip(t.rows[0].cells, heads):
        borders(hc, pal.ink, sz=8, edges=("bottom",))
        cell_margins(hc, top=20, bottom=30, left=60, right=60)
        para(hc, h, 8, bold=True, color=pal.label, caps_track=True, first=True)
    for ri, r in enumerate(rub["rows"], 1):
        tr = t.rows[ri]
        no_split(tr, 0.62 if sheet else None)
        cells = tr.cells
        for tc in cells:
            borders(tc, pal.hair, sz=4, edges=("bottom",))
            cell_margins(tc, top=50, bottom=50, left=60, right=60)
        p = cells[0].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        run(p, r["part"], 9.5, bold=True, color=pal.ink)
        para(cells[0], f"{r['points']} points", 8, color=pal.label)
        lv = r["levels"]
        if sheet:
            rich(cells[1].paragraphs[0], r["full"], 9, color=pal.ink)
            para(cells[2], " · ".join(lv), 9, color=pal.ink, first=True)
        else:
            for tc, lev, txt in ((cells[1], lv[0], r["full"]), (cells[2], lv[1], r["partial"]),
                                 (cells[3], lv[2], r["none"])):
                rich(tc.paragraphs[0], f"{lev}. {txt}", 9, color=pal.ink)
        score_box(cells[3 if sheet else 4], pal)
    unpad_cell(c)
    total_line(doc, spec, pal)
    return c


# ------------------------------------------------------------------------ questions

def question_card(cell, q, n, pal, course, base_dir, width):
    """One question, in its own bordered card.

    From his Master_Physics layout: a card per question, the number top left and the
    tier top right. Both live in ONE paragraph with a right tab stop rather than in a
    nested table - a table here drags a blank paragraph in with it and the card gains
    a line of nothing at the top.
    """
    borders(cell, pal.hair, sz=6)
    cell_margins(cell, top=55, bottom=55, left=130, right=130)
    inner = round(width - CARD_PAD_IN, 2)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(inner), WD_TAB_ALIGNMENT.RIGHT)
    # Just the number. No box, no square, no circle. SHULL-CHG-0015 - his own sheets
    # set these in solid squares, which is the withdrawn number box and a fill besides.
    run(p, f"{n}.", 10.5, bold=True, color=pal.ink)
    if q.get("tier"):
        colour = tier_colour(q["tier"], pal)
        run(p, "\t", TIER_PT)
        run(p, "•  ", TIER_PT + 1, bold=True, color=colour)
        r = run(p, q["tier"].upper(), TIER_PT, bold=True, color=colour)
        el = OxmlElement("w:spacing"); el.set(qn("w:val"), "26")
        r._element.get_or_add_rPr().append(el)

    para(cell, q["prompt"], 10, color=pal.ink)
    for part in q.get("parts", []):
        # Multi-part items keep full body size on every part.
        para(cell, "   " + debullet(part), 10, color=pal.ink)

    if q.get("diagram"):
        diagram_block(cell, q["diagram"], pal, inner, base_dir)
    if q.get("draw"):
        draw_block(cell, q["draw"], pal, inner - 0.2)
    if q.get("given") or q.get("need"):
        given_need(cell, q.get("given", ""), q.get("need", ""), pal, inner)

    if q.get("math"):
        # The self-check answer sits in the box's bottom-right corner, not on a line of
        # its own below it (SHULL-CHG-0039). Which questions carry one is 0021's call.
        work_box(cell, q.get("workLabel", ""), pal,
                 float(q.get("workHeightIn", WORK_MIN_IN)), inner,
                 watermark=q.get("watermark", WATERMARK),
                 corner=q.get("selfCheck"))
    elif q.get("answerLines"):
        rule_lines(cell, int(q["answerLines"]), pal.hair)

    if q.get("selfCheck") and not q.get("math"):
        # Bracketed self-check answers for NUMERIC results only - never for an
        # explanation, a vocabulary term, or a graph reading, where the bracket hands
        # over the whole answer.
        pp = para(cell, q["selfCheck"], 8, color=pal.label)
        pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    unpad_cell(cell)
    return cell


def render_questions(doc, questions, pal, course, base_dir):
    """Cards down the page, two across where both questions are short.

    His layout pairs the two warm-ups on one row and gives the longer questions the
    full width. Pairing is decided by `is_short`, so it follows the content rather
    than a hand-placed break.
    """
    gutter = 0.18
    half = round((TEXT_W_IN - gutter) / 2, 2)
    i = 0
    while i < len(questions):
        q = questions[i]
        nxt = questions[i + 1] if i + 1 < len(questions) else None
        pair = nxt is not None and is_short(q) and is_short(nxt)
        if pair:
            t = doc.add_table(rows=1, cols=3)
            fix_widths(t, [half, gutter, half])
            no_split(t.rows[0])
            question_card(t.rows[0].cells[0], q, i + 1, pal, course, base_dir, half)
            question_card(t.rows[0].cells[2], nxt, i + 2, pal, course, base_dir, half)
            i += 2
        else:
            t = doc.add_table(rows=1, cols=1)
            fix_widths(t, [TEXT_W_IN])
            # A question stays whole. Split across a page break, the parts and the work
            # box land on the next sheet with no number above them.
            no_split(t.rows[0])
            question_card(t.rows[0].cells[0], q, i + 1, pal, course, base_dir, TEXT_W_IN)
            i += 1
        gap(doc, 3)


# ------------------------------------------------------------------------- profiles

def refuse(msg, *more):
    print("build_worksheet_docx: " + msg, file=sys.stderr)
    for m in more:
        print("   " + m, file=sys.stderr)
    return 1


def check_profile(spec, course, sec):
    """The course profile, enforced. Returns an error string or None."""
    qs = sec.get("questions", [])
    eqs = sec.get("equations", spec.get("equations")) or []
    blocks = sec.get("blocks", [])
    math_qs = [q for q in qs if q.get("math")]
    work_qs = [q for q in qs if q.get("math") or q.get("answerLines")]
    visual = ([q for q in qs if q.get("diagram") or q.get("draw")]
              + [b for b in blocks if b["kind"] in
                 ("cards", "slots", "draw", "diagram", "sort", "match", "outline")])
    where = f"section {sec['code']}"

    if course == "geology":
        if eqs:
            return (f"{where}: Geology has no math, so it has no equation bar. "
                    'Remove "equations". SHULL-CHG-0017.')
        if math_qs:
            return (f"{where}: Geology has no math. Remove \"math\" from "
                    f"{len(math_qs)} question(s). SHULL-CHG-0017.")
        if not visual and not sec.get("proseOnly"):
            return (f"{where}: nothing to look at, label, order or match - the whole "
                    "sheet is prose. Geology is the visual class, so add a diagram, "
                    "draw, sort, match, cards or slots block.\n   The light ones are "
                    "sort and match: one page, no scissors. Cut-and-glue is worth two "
                    "pages sometimes and should not be the default.\n   If this sheet "
                    'genuinely is a reading response, set "proseOnly": true. '
                    "SHULL-CHG-0022.")

    if course == "physics":
        if work_qs:
            return (f"{where}: Physics carries no work areas at all - no ruled lines, "
                    "no work boxes. Students work in their Hayden-McNeil carbonless "
                    'lab notebooks. Remove "math" and "answerLines" from '
                    f"{len(work_qs)} question(s). SHULL-CHG-0019.")
        if not (sec.get("priorKnowledge") or spec.get("priorKnowledge")):
            return (f"{where}: no prior-knowledge recall block. Students arrive at a "
                    "friction problem having last seen free-body diagrams two weeks "
                    'ago. Add "priorKnowledge".')
        if not eqs:
            return (f"{where}: no equations. They go at the top of the page so "
                    "students know what they may reach for.")

    if course == "chemistry":
        if math_qs and not eqs and not spec.get("noEquationBar"):
            return (f"{where}: math questions and no equation bar. Add \"equations\", "
                    'or "noEquationBar": true if the math genuinely needs none.')

    # Self-check answers. His rule: "on calculations, include answers so the students
    # can self check, add this on all but the final challenge problem." The last
    # question is the one they have to commit to without a safety net, so it is not
    # merely allowed to lack an answer - it is refused if it carries one.
    #
    # Numeric results ONLY. A conceptual question keeps no bracket, because there the
    # bracket would hand over the whole answer rather than confirm arithmetic - which
    # is why a question is marked `calculation` in the spec rather than guessed at from
    # its wording. "Name the kinematic equation you would use" reads like a physics
    # problem and has no number in it.
    for k, q in enumerate(qs):
        last = k == len(qs) - 1
        calc = q.get("calculation") or q.get("math")
        if calc and not last and not q.get("selfCheck"):
            return (f"{where} question {k + 1} is a calculation with no self-check "
                    'answer. Add "selfCheck", or drop "calculation" if the answer is '
                    "not a number. SHULL-CHG-0021.")
        if last and q.get("selfCheck"):
            return (f"{where}: the last question carries a self-check answer. That one "
                    "they finish without a net. SHULL-CHG-0021.")

    ramp = RAMP[course]
    if ramp and qs:
        shape = {}
        for q in qs:
            shape[q.get("tier")] = shape.get(q.get("tier"), 0) + 1
        if shape != ramp:
            want = ", ".join(f"{k} {ramp[k]}" for k in TIERS)
            got = ", ".join(f"{k} {shape.get(k, 0)}" for k in TIERS)
            return (f"{where}: the ramp is {want}. This section has {got}. The ramp is "
                    "what makes a sheet start easy and finish on a full problem.")
    if qs and course != "geology":
        # The order is the ramp, whatever the counts are: difficulty never goes
        # backwards, and the multi-topic problem is last because it is the one that
        # reaches back into earlier units.
        seen = [TIERS.index(q.get("tier")) for q in qs if q.get("tier") in TIERS]
        if len(seen) != len(qs):
            loose = [q.get("tier") for q in qs if q.get("tier") not in TIERS]
            return f"{where}: unknown tier(s) {loose}. Tiers are {list(TIERS)}."
        if seen != sorted(seen):
            return (f"{where}: the questions do not ramp - "
                    f"{[q['tier'] for q in qs]}. Easy to start, harder as it goes.")
        if qs[-1].get("tier") != "multi-topic":
            return (f"{where}: the last question is {qs[-1].get('tier')!r}. It should be "
                    "the multi-topic problem that reaches back into earlier units.")
    return None


# ------------------------------------------------------------------------------ main

def main():
    # --rubric-sheet draws the spec's rubric as the teacher's one-page grading sheet,
    # from the same rows the handout prints, so the two cannot drift.
    sheet_mode = "--rubric-sheet" in sys.argv[1:]
    argv = [sys.argv[0]] + [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(argv) < 2:
        return refuse("usage: build_worksheet_docx.py <spec.json> [out.docx] [--rubric-sheet]")
    spec = json.load(open(argv[1]))
    course = spec["course"]
    if course not in COURSE_CODE:
        return refuse(f"unknown course {course!r}.")
    pal = Palette(course)
    sections = spec["sectionsContent"]

    valid = known_sections(course)
    bad = [s for s in spec["sections"] if s not in valid]
    if bad:
        return refuse(f"section(s) {', '.join(bad)} are not in "
                      f"courses/{course}/DECISIONS.md.",
                      "Nothing is built against a code not in the decisions file.")
    declared = [s["code"] for s in sections]
    if declared != list(spec["sections"]):
        return refuse(f"sections {spec['sections']} but content for {declared}. "
                      "The two lists are the same fact and must agree.")

    for sec in sections:
        err = check_profile(spec, course, sec)
        if err:
            return refuse(err)

    code = COURSE_CODE[course]
    unit = f"U{int(spec['unit']):02d}"
    utitle = unit_title(course, spec["unit"])       # a course fact, read not typed
    span = (f"S{spec['sections'][0]}-S{spec['sections'][-1]}"
            if len(spec["sections"]) > 1 else f"S{spec['sections'][0]}")
    kind = spec.get("docType", "Practice_Set")
    out = argv[2] if len(argv) > 2 else os.path.join(HERE, document_name(
        "Rubric" if sheet_mode else kind, spec["sections"], spec["unit"],
        unit_title=short_title(course, f"U{spec['unit']}"),
        topic=short_title(course, spec["sections"][0])))

    doc = Document()
    s = page_setup(doc)

    def header(sec, title, kind_label):
        # ---- Header. His Master_Physics banner is a solid navy slab across the top
        # of the page - handsome, and 28.46% marked with a 7.61in solid band, which
        # fails both halves of the ink gate on its own. "Make the header more ink
        # saver" was his instruction. The shape survives: the eyebrow, the big title,
        # the code on the right. Only the slab goes, replaced by a heavy accent bar on
        # the left edge, which is the one place a solid mark earns its ink.
        t = doc.add_table(rows=1, cols=2)
        fix_widths(t, [5.55, 1.95])
        h, hr = t.rows[0].cells
        for cc in (h, hr):
            borders(cc, pal.hair, sz=4, edges=("top", "bottom"))
        borders(h, pal.display, sz=30, edges=("left",))
        cell_margins(h, top=80, bottom=80, left=160, right=80)
        cell_margins(hr, top=80, bottom=80, left=80, right=60)
        para(h, f"UNIT {int(spec['unit'])} — {utitle.upper()}"
             + (f" · {sec['eyebrowNote']}" if sec.get("eyebrowNote") else ""), 7.5, bold=True,
             color=pal.accent, caps_track=True, first=True)
        para(h, title, 15, bold=True, color=pal.ink)
        pp = para(hr, f"{unit} / S{sec['code']}", 11, bold=True, color=pal.accent,
                  caps_track=True, first=True)
        pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pp = para(hr, kind_label.replace("_", " ").upper()
                  + (f"  ·  ~{spec['timeTargetMin']} MIN"
                     if spec.get("timeTargetMin") else ""),
                  6.5, color=pal.label, caps_track=True)
        pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def id_row(labels, widths, total):
        t = doc.add_table(rows=1, cols=len(labels))
        fix_widths(t, widths)
        for cc, lab in zip(t.rows[0].cells, labels):
            borders(cc, pal.ink, sz=4, edges=("bottom",))
            cell_margins(cc, top=0, bottom=40, left=0, right=120)
            p = para(cc, lab, 6.5, bold=True, color=pal.label, caps_track=True,
                     first=True)
            if lab == "SCORE":
                p.paragraph_format.tab_stops.add_tab_stop(
                    Inches(widths[-1] - 0.10), WD_TAB_ALIGNMENT.RIGHT)
                run(p, "\t", 9)
                run(p, f"/  {total}", 10, bold=True, color=pal.ink)

    if sheet_mode:
        # ---- The grading sheet: the same rubric rows, drawn as the teacher's sheet.
        rs = spec["rubricSheet"]
        sec = sections[0]
        header(sec, rs["title"], "Rubric")
        gap(doc, 5)
        id_row(rs["fields"], rs["widths"], rubric_total(spec))
        gap(doc, 8)
        rubric_block(doc, {"label": rs["scoreLabel"]}, pal, spec, sheet=True)
        gap(doc, 4)
        c = one_cell(doc)
        para(c, rs["pointer"], 9, color=pal.label, first=True)
        gap(doc, 6)
        c = one_cell(doc)
        cell_margins(c, top=0, bottom=0, left=0, right=0)
        label(c, rs.get("notesLabel", "NOTES"), pal, first=True)
        bt = c.add_table(rows=1, cols=1)
        fix_widths(bt, [TEXT_W_IN])
        no_split(bt.rows[0], float(rs.get("notesHeightIn", 2.5)))
        borders(bt.rows[0].cells[0], pal.label, sz=6)
        unpad_cell(c)
        sections_to_build = []
    else:
        sections_to_build = sections

    for si, sec in enumerate(sections_to_build):
        if si:
            page_break(doc)
        header(sec, sec["title"], kind)

        # ---- Name / Date / Period / Score. Four ruled fields on one line, the label
        # sitting on the rule - his layout, and lighter than a box around each. The
        # score total is summed from the questions, never typed: a header reading "/ 20"
        # over questions adding to 18 is one fact stored in two places.
        total = (sum(int(q.get("points", 1)) for q in sec.get("questions", []))
                 + sum(rubric_total(spec) if b["kind"] == "rubric" else int(b.get("points", 0))
                       for b in sec.get("blocks", [])))
        gap(doc, 5)
        id_row(("NAME", "DATE", "PERIOD", "SCORE"), [3.05, 1.75, 1.10, 1.60], total)

        # ---- What this is about, before anything is asked of them.
        # A project handout has no concept to recall; its job line is the directions.
        if sec.get("concept"):
            gap(doc, 5)
            c = one_cell(doc)
            borders(c, pal.hair, sz=4)
            borders(c, pal.display, sz=24, edges=("left",))
            shade(c, pal.surface)
            cell_margins(c, top=70, bottom=70, left=130, right=130)
            label(c, sec.get("conceptLabel", "REMEMBER"), pal, first=True)
            para(c, sec["concept"]["summary"], 9.5)
            for x in sec["concept"].get("keyIdeas", []):
                para(c, "•  " + debullet(x), 9.5)

        pk = sec.get("priorKnowledge", spec.get("priorKnowledge"))
        if pk:
            gap(doc, 4)
            c = one_cell(doc); borders(c, pal.hair, sz=4)
            cell_margins(c, top=70, bottom=70, left=130, right=130)
            label(c, pk.get("label", "BEFORE YOU START — YOU ALREADY KNOW THIS"),
                  pal, first=True)
            for x in pk["items"]:
                para(c, "•  " + debullet(x), 9.5)

        eqs = sec.get("equations", spec.get("equations")) or []
        if eqs:
            gap(doc, 2)
            equation_bar(doc, eqs, pal,
                         spec.get("equationLabel", "EQUATIONS YOU MAY USE"))

        if sec.get("directions"):
            gap(doc, 2)
            c = one_cell(doc); borders(c, pal.ink, sz=8, edges=("top",))
            label(c, sec.get("directionsLabel", "DIRECTIONS"), pal, first=True)
            para(c, sec["directions"], 9.5)

        gap(doc, 4)
        render_questions(doc, sec.get("questions", []), pal, course, HERE)

        for bi, block in enumerate(sec.get("blocks", [])):
            kindb = block["kind"]
            if block.get("gapBefore"):
                gap(doc, float(block["gapBefore"]))
            if block.get("pageBreakBefore"):
                page_break(doc)
                c = one_cell(doc); borders(c, pal.display, sz=18, edges=("bottom",))
                para(c, block.get("heading", sec["title"]), 15, bold=True,
                     color=pal.ink, first=True)
            if block.get("instruction") and kindb in ("cards", "slots"):
                c = one_cell(doc)
                para(c, block["instruction"], 9.5, color=pal.ink, first=True)
                gap(doc, 2)
            if kindb == "cards":
                cards_block(doc, block, pal)
            elif kindb == "slots":
                slots_block(doc, block, pal)
            elif kindb == "sort":
                sort_block(doc, block, pal)
            elif kindb == "match":
                match_block(doc, block, pal)
            elif kindb == "reflection":
                reflection_block(doc, block, pal)
            elif kindb == "draw":
                draw_block(one_cell(doc), block, pal, TEXT_W_IN - 0.3)
            elif kindb == "diagram":
                diagram_block(one_cell(doc), block, pal, TEXT_W_IN - 0.3, HERE)
            elif kindb == "outline":
                outline_block(doc, block, pal)
            elif kindb == "list":
                list_block(doc, block, pal)
            elif kindb == "table":
                table_block(doc, block, pal)
            elif kindb == "notes":
                notes_block(doc, block, pal)
            elif kindb == "formats":
                formats_block(doc, block, pal)
            elif kindb == "plan":
                plan_block(doc, block, pal)
            elif kindb == "rubric":
                rubric_block(doc, block, pal, spec)
            else:
                return refuse(f"unknown block kind {kindb!r} in section {sec['code']}.")

        if sec.get("close"):
            # Banner and checklist in ONE non-splitting row. As two tables they broke
            # apart at the page edge and a section ended on a banner with its checklist
            # stranded alone on the next page.
            t = doc.add_table(rows=1, cols=1)
            fix_widths(t, [TEXT_W_IN])
            no_split(t.rows[0])
            c = t.rows[0].cells[0]
            borders(c, pal.ink, sz=12, edges=("top",))
            borders(c, pal.display, sz=18, edges=("bottom",))
            para(c, sec["close"]["banner"], 9, bold=True, color=pal.ink,
                 caps_track=True, first=True)
            for x in sec["close"].get("checklist", []):
                check_item(c, x, pal)

    running_footer(s, f"SHULL SCIENCE          {unit} · {span}", pal,
                   with_page_numbers=True)
    trim_tail(doc)
    doc.save(out)
    nq = sum(len(x.get("questions", [])) for x in sections)
    print(f"wrote {out}  —  {code} {unit} {span}, "
          f"{len(sections)} section(s), {nq} question(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
