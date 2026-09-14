#!/usr/bin/env python3
"""Build the separate teacher answer key for a SHULL worksheet / practice set.

    python3 templates/worksheet/build_worksheet_key_docx.py specs/<spec>.json out.docx

Matt's rule, stated plainly: produce the answer key whenever a Physics problem set like this
is created (2026-09-14, courses/physics/DECISIONS.md). The student handout uses inline
self-check brackets on every calculation but the last (SHULL-CHG-0021) - deliberately withheld
so the last question is finished without a net. The key is a separate file precisely so that
withheld answer, and full worked steps for every question, live somewhere: never on the
student page.

This is a companion to build_worksheet_docx.py, not a course-profile-enforcing builder in its
own right - the key is teacher-only, so the no-work-area / no-answer-blank rules that govern
the student handout do not apply to it. It shares the same masthead, palette, fonts, and
footer grammar as the student handout so the two read as one document set, not two different
templates that happened to be filed in the same folder.

Each question in the spec must carry a "solution" field: the full worked steps and final
answer, independent of any "selfCheck" bracket. A question with no "solution" is refused -
a key that silently skips a question is worse than no key.
"""
import json, os, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from _shull_docx import (          # noqa: E402
    COURSE_CODE, Palette, known_sections, unit_title, borders, para, run,
    fix_widths, no_split, page_setup, running_footer, cell_margins, gap, trim_tail,
)

HERE = os.path.dirname(os.path.abspath(__file__))
TEXT_W_IN = 7.50


def refuse(msg, *more):
    print("build_worksheet_key_docx: " + msg, file=sys.stderr)
    for m in more:
        print("   " + m, file=sys.stderr)
    return 1


def page_break(doc):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    p = doc.add_paragraph()
    r = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r._element.append(br)


def answer_block(doc, n, q, pal):
    """One question's key entry: number + tier, the prompt restated small, the full
    worked solution, and the final answer set apart so it can be read at a glance
    while grading - never just the bracket the student sheet withholds."""
    t = doc.add_table(rows=1, cols=1)
    fix_widths(t, [TEXT_W_IN])
    no_split(t.rows[0])          # a question's key entry stays whole across a page break,
    c = t.rows[0].cells[0]       # same rule the student handout's cards follow.
    borders(c, pal.hair, sz=4)
    cell_margins(c, top=70, bottom=70, left=130, right=130)

    head = c.paragraphs[0]
    r = run(head, f"{n}.  ", 11, bold=True, color=pal.ink)
    if q.get("tier"):
        run(head, q["tier"].upper(), 7, bold=True, color=pal.accent)

    p = para(c, q["prompt"], 9, color=pal.label)
    p.runs[0].italic = True
    for part in q.get("parts", []):
        p = para(c, "   " + str(part), 9, color=pal.label)
        p.runs[0].italic = True

    gap(c, 2)
    for line in str(q["solution"]).split("\n"):
        if line.strip():
            para(c, line.strip(), 10, color=pal.ink)

    if q.get("finalAnswer"):
        para(c, "ANSWER:  " + str(q["finalAnswer"]), 10, bold=True, color=pal.accent)
    return c


def main():
    if len(sys.argv) < 2:
        return refuse("usage: build_worksheet_key_docx.py <spec.json> [out.docx]")
    spec = json.load(open(sys.argv[1]))
    course = spec["course"]
    if course not in COURSE_CODE:
        return refuse(f"unknown course {course!r}.")
    pal = Palette(course)
    sections = spec["sectionsContent"]

    valid = known_sections(course)
    bad = [s for s in spec["sections"] if s not in valid]
    if bad:
        return refuse(f"section(s) {', '.join(bad)} are not in courses/{course}/DECISIONS.md.")

    missing = []
    for sec in sections:
        for i, q in enumerate(sec.get("questions", []), start=1):
            if not q.get("solution"):
                missing.append(f"{sec['code']} Q{i}")
    if missing:
        return refuse("no \"solution\" field for: " + ", ".join(missing),
                       "A key that silently skips a question is worse than no key.")

    code = COURSE_CODE[course]
    unit = f"U{int(spec['unit']):02d}"
    utitle = unit_title(course, spec["unit"])
    span = (f"S{spec['sections'][0]}-S{spec['sections'][-1]}"
            if len(spec["sections"]) > 1 else f"S{spec['sections'][0]}")
    kind = spec.get("docType", "Practice_Set")
    out = sys.argv[2] if len(sys.argv) > 2 else \
        os.path.join(HERE, f"SHULL_{code}_{kind}_{unit}_{span}_Key.docx")

    doc = Document()
    s = page_setup(doc)

    for si, sec in enumerate(sections):
        if si:
            page_break(doc)

        t = doc.add_table(rows=1, cols=2)
        fix_widths(t, [5.55, 1.95])
        h, hr = t.rows[0].cells
        for cc in (h, hr):
            borders(cc, pal.hair, sz=4, edges=("top", "bottom"))
        borders(h, pal.display, sz=30, edges=("left",))
        cell_margins(h, top=80, bottom=80, left=160, right=80)
        cell_margins(hr, top=80, bottom=80, left=80, right=60)
        para(h, f"UNIT {int(spec['unit'])} — {utitle.upper()} · TEACHER KEY", 7.5,
             bold=True, color=pal.accent, caps_track=True, first=True)
        para(h, sec["title"], 15, bold=True, color=pal.ink)
        pp = para(hr, f"{unit} / S{sec['code']}", 11, bold=True, color=pal.accent,
                  caps_track=True, first=True)
        pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pp = para(hr, kind.replace("_", " ").upper() + "  ·  ANSWER KEY",
                  6.5, color=pal.label, caps_track=True)
        pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        gap(doc, 6)
        for i, q in enumerate(sec.get("questions", []), start=1):
            answer_block(doc, i, q, pal)
            gap(doc, 4)

    running_footer(s, f"SHULL SCIENCE          {unit} · {span} · KEY", pal,
                   with_page_numbers=True)
    trim_tail(doc)
    doc.save(out)
    nq = sum(len(x.get("questions", [])) for x in sections)
    print(f"wrote {out}  —  {code} {unit} {span} KEY, "
          f"{len(sections)} section(s), {nq} answer(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
