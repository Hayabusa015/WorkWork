#!/usr/bin/env python3
"""Build a snow-day assignment (student copy + teacher key) from a spec.

    python3 templates/snow_day/build_snow_day_docx.py templates/snow_day/specs/geology.json OUTDIR

Snow-day work is posted to Google Classroom and done on a Chromebook, not printed,
so the answer areas are typing boxes rather than ruled lines, and the readings are
live links. It carries no U#/S#.# code: it is written to be usable on any day of the
year, so it belongs to no section. Both files come from one spec so the key cannot
drift from the questions.
"""
import json, os, sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from _shull_docx import Palette, hexof, page_setup, running_footer, one_cell, borders, \
    cell_margins, no_split, gap

# Google Docs has Archivo and not Trade Gothic Next; naming the face it actually has
# keeps the converted Doc in the brand face instead of falling back to Arial.
FACE = "Archivo"
COURSE_NAME = {"chemistry": "Chemistry", "physics": "Physics", "geology": "Geology"}


def r(p, text, size, *, bold=False, color=None, italic=False, track=None):
    x = p.add_run(text)
    x.font.name = FACE
    x.font.size = Pt(size)
    x.bold, x.italic = bold, italic
    if color:
        x.font.color.rgb = RGBColor.from_string(hexof(color))
    if track:
        el = OxmlElement("w:spacing"); el.set(qn("w:val"), str(track))
        x._element.get_or_add_rPr().append(el)
    return x


def p_(container, after=2, before=0):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    return p


def rule(container, hexval, sz=12):
    p = p_(container, after=6)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr"); x = OxmlElement("w:bottom")
    x.set(qn("w:val"), "single"); x.set(qn("w:sz"), str(sz)); x.set(qn("w:color"), hexof(hexval))
    b.append(x); pPr.append(b)


def link(p, url, text, size, color):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    run = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts"); f.set(qn("w:ascii"), FACE); f.set(qn("w:hAnsi"), FACE)
    c = OxmlElement("w:color"); c.set(qn("w:val"), hexof(color))
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single")
    b = OxmlElement("w:b")
    s = OxmlElement("w:sz"); s.set(qn("w:val"), str(int(size * 2)))
    for e in (f, b, c, s, u):
        rPr.append(e)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    run.append(rPr); run.append(t); h.append(run); p._p.append(h)


def head(doc, spec, pal, key):
    kicker = f"GARFIELD {COURSE_NAME[spec['course']].upper()}  ·  SNOW DAY ASSIGNMENT"
    if key:
        kicker += "  ·  TEACHER KEY"
    r(p_(doc, after=0), kicker, 8.5, bold=True, color=pal.accent, track=26)
    p = p_(doc, after=0); r(p, spec["title"], 22, bold=True, color=pal.ink)
    p = p_(doc, after=0); r(p, spec["subtitle"], 11, color=pal.label)
    rule(doc, pal.display)
    if not key:
        p = p_(doc, after=8)
        r(p, "Name: ", 10, bold=True, color=pal.ink)
        r(p, "_" * 40, 10, color=pal.hair)


def section_label(doc, text, pal):
    p = p_(doc, after=3, before=10)
    r(p, text.upper(), 9, bold=True, color=pal.accent, track=26)


def answer_box(doc, pal, lines):
    cell = one_cell(doc, width=7.5)
    borders(cell, pal.hair, sz=6)
    cell_margins(cell, top=80, bottom=80, left=120, right=120)
    no_split(doc.tables[-1].rows[0], height_in=0.2 * lines + 0.15)
    p = cell.paragraphs[0]
    r(p, "", 10.5, color=pal.ink)
    return cell


def build(spec, key, out):
    pal = Palette(spec["course"])
    doc = Document()
    sec = page_setup(doc, margin_lr=0.5, margin_tb=0.5, base_pt=10.5)
    doc.styles["Normal"].font.name = FACE
    doc.styles["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), FACE)
    head(doc, spec, pal, key)

    if not key:
        section_label(doc, "What to do", pal)
        steps = [
            "Open both readings below. The titles are links.",
            f"Answer questions 1 to {sum(1 for q in spec['questions'] if not q.get('optional'))}. "
            "Type right inside each box. A couple of sentences is usually enough.",
            "Write the paragraph in Part 3.",
            "Click Turn In on Google Classroom.",
        ]
        for i, s in enumerate(steps, 1):
            p = p_(doc, after=1); r(p, f"{i}.  ", 10.5, bold=True, color=pal.accent); r(p, s, 10.5, color=pal.ink)
        p = p_(doc, after=0, before=4)
        r(p, "Time: ", 10.5, bold=True, color=pal.ink); r(p, spec["time"] + ".", 10.5, color=pal.ink)
        r(p, "  Answer in your own words. Copying sentences straight from the reading doesn't count.",
          10.5, color=pal.ink)

    section_label(doc, "Part 1 · The readings", pal)
    for rd in spec["readings"]:
        p = p_(doc, after=0, before=4)
        r(p, rd["label"] + ":  ", 10.5, bold=True, color=pal.ink)
        link(p, rd["url"], rd["title"], 10.5, pal.accent)
        p = p_(doc, after=0); r(p, rd["source"], 9.5, italic=True, color=pal.label)
        p = p_(doc, after=2); r(p, rd["note"], 10, color=pal.ink)
        if key:
            p = p_(doc, after=2); r(p, rd["url"], 8.5, color=pal.label)

    section_label(doc, "Part 2 · Questions", pal)
    n = 0
    for q in spec["questions"]:
        p = p_(doc, after=3, before=6)
        if q.get("optional"):
            r(p, "★  ", 10.5, bold=True, color=pal.accent)
        else:
            n += 1
            r(p, f"{n}.  ", 10.5, bold=True, color=pal.accent)
        r(p, q["q"], 10.5, color=pal.ink)
        if key:
            p = p_(doc, after=2); r(p, q["key"], 10.5, bold=True, color=pal.accent)
        else:
            answer_box(doc, pal, q["lines"])

    section_label(doc, "Part 3 · Write", pal)
    p = p_(doc, after=3); r(p, spec["write"]["q"], 10.5, color=pal.ink)
    if key:
        p = p_(doc, after=2); r(p, spec["write"]["key"], 10.5, bold=True, color=pal.accent)
        section_label(doc, "Teacher notes", pal)
        for t in spec["teacher_notes"]:
            p = p_(doc, after=2); r(p, "·  " + t, 10, color=pal.ink)
    else:
        answer_box(doc, pal, spec["write"]["lines"])

    footer = f"Garfield {COURSE_NAME[spec['course']]}  ·  Snow Day Assignment  ·  {spec['title']}"
    running_footer(sec, footer + ("  ·  Key" if key else ""), pal)
    for run in sec.footer.paragraphs[0].runs:
        run.font.name = FACE
    doc.save(out)


def names(spec):
    t = spec["title"]
    return f"Snow Day Activity - {t}.docx", f"Snow Day Activity Key - {t}.docx"


if __name__ == "__main__":
    spec = json.load(open(sys.argv[1]))
    outdir = sys.argv[2]
    os.makedirs(outdir, exist_ok=True)
    stu, key = names(spec)
    build(spec, False, os.path.join(outdir, stu))
    build(spec, True, os.path.join(outdir, key))
    print(os.path.join(outdir, stu)); print(os.path.join(outdir, key))
